#!/usr/bin/env python3
"""Ajoute la section home services à `livrables/regroupement-textes-us.docx`.

    PYTHONPATH=.venv/Lib/site-packages <python uv> etudes-ponctuelles/2026-09-29-home-services-us/note.py

À lancer après `etudes-ponctuelles/2026-09-29-hdbscan-residus-us/note.py`, qui
réécrit le document entier sans cette section. La section est insérée avant
« Suite proposée », renumérotée après la dernière section ; si elle existe déjà,
elle est remplacée. Relancer ensuite logistic-regression-study/python/07D_note.py
(section 5).

Lit `sorties/A1-decoupages.csv`, `A2-groupes.csv`, `B1-voisins.csv`,
`C1-longueur.csv`. Aucun chiffre n'est calculé ici. Sans pandas, bloqué par la
règle Windows de contrôle des applications depuis le 2026-09-29.

Étiquettes des groupes de K = 4 (sans les enseignes, toutes notes) : établies
d'après leur longueur, leur note et leurs mots les plus caractéristiques ; les
trois premiers partagent le même vocabulaire d'éloge (lecture à l'écran,
2026-09-29). Le script s'arrête si les effectifs ne correspondent plus.
"""
from __future__ import annotations

# Remplacé le 2026-09-29 par etudes-ponctuelles/2026-09-29-synthese-textes/note.py,
# qui écrit livrables/regroupement-textes-us.docx en entier. Relancer ce script
# réécrirait l'ancienne structure du document.
raise SystemExit("remplacé par etudes-ponctuelles/2026-09-29-synthese-textes/note.py")


import csv
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402

DOSSIER = Path(__file__).resolve().parent
RACINE = DOSSIER.parents[1]
SORTIES = DOSSIER / "sorties"
FIGURE = SORTIES / "figure-k4-sans-enseignes.png"
LIVRABLE = RACINE / "livrables" / "regroupement-textes-us.docx"
TITRE = "4. Home services US : les suppressions s'isolent-elles dans un groupe ?"
SUITE = "Suite proposée"
SEUIL = 97.5
SE, SE5, TF, TF5 = ("sans_enseignes", "sans_enseignes_5_etoiles", "toutes_fiches",
                    "toutes_fiches_5_etoiles")

# groupe de K = 4 sans les enseignes, toutes notes -> (étiquette, nombre d'avis attendu)
K4 = {0: ("Éloges 5 étoiles courts", 361), 2: ("Éloges 5 étoiles moyens", 693),
      1: ("Éloges 5 étoiles longs", 1134), 3: ("Plaintes, surtout 1 étoile", 140)}


def lire(nom):
    with (SORTIES / nom).open(encoding="utf-8") as f:
        return list(csv.DictReader(f))


def virgule(x, d=1):
    return f"{float(x):.{d}f}".replace(".", ",")


def espace(n):
    return f"{int(float(n)):,}".replace(",", " ")


# ---------------------------------------------------------------------------
# Figure
# ---------------------------------------------------------------------------
SURFACE = "#fcfcfb"; ENCRE = "#0b0b0b"; ENCRE_DOUCE = "#52514e"; GRILLE = "#e4e3df"
BLEU = "#2a78d6"; GRIS = "#bdbcb8"


def figure(lignes):
    lignes = lignes[::-1]
    y = np.arange(len(lignes))
    fig, ax = plt.subplots(figsize=(7.2, 2.7), dpi=150)
    fig.patch.set_facecolor(SURFACE); ax.set_facecolor(SURFACE)
    fig.subplots_adjust(left=0.32, right=0.96, top=0.70, bottom=0.15)
    h = 0.36
    obs = [int(r["n_supprimes"]) for r in lignes]
    att = [float(r["attendus"]) for r in lignes]
    seuil = [int(r["hasard_fiche_haut"]) for r in lignes]
    ax.barh(y + h / 2, obs, h, color=BLEU, label="supprimés")
    ax.barh(y - h / 2, att, h, color=GRIS, label="attendus d'après la note")
    ax.scatter(seuil, y + h / 2, marker="|", s=140, color=ENCRE, zorder=3,
               label="seuil du hasard à fiche égale")
    for yi, o in zip(y, obs):
        ax.text(o + 0.8, yi + h / 2, str(o), va="center", fontsize=7, color=ENCRE)
    ax.set_yticks(y, [f"{K4[int(r['groupe'])][0]} ({espace(r['n_avis'])} avis)"
                      for r in lignes], fontsize=7.5)
    for c in ("top", "right", "left"):
        ax.spines[c].set_visible(False)
    ax.spines["bottom"].set_color(GRILLE)
    ax.tick_params(colors=ENCRE_DOUCE, labelsize=7.5, length=0)
    ax.set_xlabel("avis supprimés", fontsize=7.5, color=ENCRE_DOUCE)
    ax.legend(frameon=False, fontsize=7, loc="lower left", bbox_to_anchor=(-0.02, 1.0), ncol=3)
    fig.text(0.012, 0.97, "Home services US, sans les six enseignes, toutes notes : KMeans à "
             "4 groupes", fontsize=10.5, color=ENCRE, va="top")
    fig.savefig(FIGURE, dpi=150, facecolor=SURFACE)
    plt.close(fig)


# ---------------------------------------------------------------------------
# Document
# ---------------------------------------------------------------------------
def P(doc, texte, style=None):
    p = doc.add_paragraph(style=style)
    for i, m in enumerate(texte.split("**")):
        if m:
            p.add_run(m).bold = i % 2 == 1
    return p


def puces(doc, textes):
    for t in textes:
        P(doc, t, "List Bullet")


def fond(cell, couleur):
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), couleur)
    cell._tc.get_or_add_tcPr().append(sh)


def tableau(doc, titre, entetes, lignes, largeurs, note=None):
    doc.add_paragraph(titre, style="TitreTableau")
    t = doc.add_table(rows=1 + len(lignes), cols=len(entetes))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    bords = OxmlElement("w:tblBorders")
    for cote in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{cote}")
        if cote in ("left", "right", "insideV"):
            e.set(qn("w:val"), "nil")
        else:
            e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), "D9D8DE")
        bords.append(e)
    t._tbl.tblPr.append(bords)
    for i, rangee in enumerate([entetes] + lignes):
        for j, val in enumerate(rangee):
            c = t.rows[i].cells[j]
            c.width = Cm(largeurs[j])
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.RIGHT
            texte, gras = val if isinstance(val, tuple) else (val, False)
            r = p.add_run(str(texte)); r.font.size = Pt(9); r.bold = gras
            if i == 0:
                r.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); fond(c, "322954")
            elif i % 2 == 0:
                fond(c, "F4F4F7")
    for j, gc in enumerate(t._tbl.tblGrid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(int(largeurs[j] * 567)))
    if note:
        doc.add_paragraph(note, style="NoteTableau")


def texte_titre(e) -> str:
    return "".join(t.text or "" for t in e.iter(qn("w:t"))).strip()


def preparer(doc):
    """Retire l'ancienne version de la section, renumérote la suite, et rend le
    titre « Suite proposée » devant lequel insérer."""
    corps = doc.element.body
    enfants = list(corps)
    debut = next((i for i, e in enumerate(enfants) if texte_titre(e) == TITRE), None)
    if debut is not None:
        fin = next(i for i, e in enumerate(enfants) if i > debut and texte_titre(e).endswith(SUITE))
        for e in enfants[debut:fin]:
            corps.remove(e)
    suite = next(p for p in doc.paragraphs if p.text.strip().endswith(SUITE))
    numeros = [int(p.text.split(".")[0]) for p in doc.paragraphs
               if p.style.name == "Heading 1" and p.text[:1].isdigit()
               and not p.text.strip().endswith(SUITE)]
    numero = max(numeros + [int(TITRE.split(".")[0])]) + 1
    for r in suite.runs:
        r.text = ""
    suite.runs[0].text = f"{numero}. {SUITE}"
    return suite._p


def main() -> int:
    dec = lire("A1-decoupages.csv")
    grp = lire("A2-groupes.csv")
    vois = {r["population"]: r for r in lire("B1-voisins.csv")}
    lon = {(r["population"], r["textes"]): r for r in lire("C1-longueur.csv")}

    def groupes(pop, k):
        return [r for r in grp if r["population"] == pop and r["methode"] == "kmeans"
                and r["k"] == str(k)]

    def plus_supprime(pop, k):
        return max(groupes(pop, k), key=lambda r: float(r["n_supprimes"]) - float(r["attendus"]))

    def decoupage(pop, methode, k):
        return next(r for r in dec if r["population"] == pop and r["methode"] == methode
                    and r["k"] == str(k))

    def rang(pop, methode, k):
        return float(decoupage(pop, methode, k)["exces_du_pire_rang_fiche_pct"])

    def ks_au_dessus(pop):
        return [str(k) for k in range(2, 11) if rang(pop, "kmeans", k) > SEUIL]

    k4 = groupes(SE, 4)
    for r in k4:
        if int(r["n_avis"]) != K4[int(r["groupe"])][1]:
            raise SystemExit(f"groupe {r['groupe']} : effectif changé, étiquettes à refaire")
    k4.sort(key=lambda r: float(r["n_supprimes"]) - float(r["attendus"]), reverse=True)
    court, court5 = k4[0], plus_supprime(SE5, 4)
    if court["groupe"] != "0" or int(court5["longueur_mediane"]) > 60:
        raise SystemExit("le groupe le plus supprimé a changé : texte à réécrire")
    figure(k4)

    n_supp_se = sum(int(r["n_supprimes"]) for r in k4)
    n_avis_se = sum(int(r["n_avis"]) for r in k4)
    n_avis_tf = sum(int(r["n_avis"]) for r in groupes(TF, 4))
    n_supp_tf = sum(int(r["n_supprimes"]) for r in groupes(TF, 4))
    n_supp_se5 = sum(int(r["n_supprimes"]) for r in groupes(SE5, 4))
    ks_se, ks_se5 = ks_au_dessus(SE), ks_au_dessus(SE5)
    rangs_tf = [rang(p, "kmeans", k) for p in (TF, TF5) for k in range(2, 11)]
    c_in, c_out = lon[(SE, "courts du groupe")], lon[(SE, "courts des autres groupes")]
    c5_in, c5_out = lon[(SE5, "courts du groupe")], lon[(SE5, "courts des autres groupes")]
    vs, vs5 = vois[SE], vois[SE5]
    n4 = lon[(SE, "note 4 du groupe")]
    autres_notes = int(court["n_supprimes"]) - int(lon[(SE, "note 5 du groupe")]["n_supprimes"])
    if autres_notes != int(n4["n_supprimes"]):
        raise SystemExit("suppressions hors 5 étoiles : pas toutes en 4 étoiles, texte à réécrire")

    doc = Document(LIVRABLE)
    avant = preparer(doc)
    n_avant = len(doc.element.body)

    doc.add_heading(TITRE, level=1)
    doc.add_heading("Ce qu'on a fait", level=2)
    puces(doc, [
        f"**Sous-ensemble** : avis home services US en anglais, {espace(n_avis_tf)} avis dont "
        f"{espace(n_supp_tf)} supprimés ; {espace(n_avis_se)} et {n_supp_se} sans les six "
        "enseignes. Les quatre chaînes antiparasitaires sont en home services.",
        "**Vecteurs e5 bruts** : un seul secteur, rien à retirer.",
        "**Catégories** : KMeans de 2 à 10 groupes, chaque avis rangé dans un groupe ; "
        "HDBSCAN pour comparer.",
        "**Question** : les avis supprimés se concentrent-ils dans un groupe plus que le hasard ?",
        "**Attendus** : suppressions qu'aurait un groupe si chaque avis était supprimé au taux "
        "de sa note, note par note (1 à 5).",
        "**Hasard à fiche égale** : les étiquettes « supprimé » sont redistribuées entre avis "
        "d'une même fiche et d'une même note ; chaque fiche garde son nombre de suppressions. "
        "Les avis d'une fiche se ressemblent (même métier, mêmes prénoms) : ce tirage écarte "
        "l'effet fiche.",
        "**Contrôle à note unique** : tout le calcul refait sur les seuls avis 5 étoiles.",
        "**Voisins** : pour chaque avis, les 10 avis au texte le plus proche, pris dans d'autres "
        "fiches ; part de supprimés parmi eux.",
    ])

    doc.add_heading("Résultats", level=2)
    puces(doc, [
        "**KMeans découpe surtout selon la longueur et le ton.** À K = 4, trois groupes sont le "
        "même éloge 5 étoiles, au même vocabulaire, coupé en trois longueurs ; le quatrième "
        "réunit les plaintes.",
        f"**Toutes notes, sans les enseignes, les éloges courts sont plus supprimés** : "
        f"{int(court['n_supprimes'])} suppressions pour {virgule(court['attendus'])} attendues, "
        f"soit {virgule(100 * int(court['n_supprimes']) / n_supp_se, 0)} % des suppressions pour "
        f"{virgule(100 * int(court['n_avis']) / n_avis_se, 0)} % des avis, sur "
        f"{court['fiches_touchees']} fiches. L'écart dépasse le hasard à fiche égale pour "
        f"K = {', '.join(ks_se)}.",
        f"**Sur les seuls 5 étoiles, l'écart s'affaiblit** : {int(court5['n_supprimes'])} pour "
        f"{virgule(court5['attendus'])} attendues. Il ne dépasse le hasard à fiche égale que pour "
        f"K = {', '.join(ks_se5) or 'aucune valeur'}. Dans le groupe toutes notes, "
        f"{autres_notes} des {int(court['n_supprimes'])} suppressions sont des avis 4 étoiles : "
        f"{n4['n_supprimes']} sur les {n4['n_avis']} avis 4 étoiles du groupe.",
        f"**La longueur seule ne l'explique pas** : parmi les textes de moins de 60 caractères, "
        f"ceux du groupe sont supprimés à {virgule(c_in['part_supprimee_pct'])} %, ceux des "
        f"autres groupes à {virgule(c_out['part_supprimee_pct'])} % "
        f"({virgule(c5_in['part_supprimee_pct'])} % et {virgule(c5_out['part_supprimee_pct'])} % "
        "en 5 étoiles seules).",
        "**À la lecture**, les supprimés du groupe sont des éloges sans détail, souvent sans nom "
        "(« Great service », « Best around ») ; les restés citent plus souvent le technicien. "
        "Observation non mesurée.",
        "**Avec les enseignes, rien au-delà de l'effet fiche** : aucun découpage KMeans ne dépasse "
        f"le hasard à fiche égale (de {virgule(min(rangs_tf), 0)} à {virgule(max(rangs_tf), 0)} % "
        "des tirages). Les avis des chaînes se ressemblent, et c'est leur fiche qui porte les "
        "suppressions.",
        f"**Voisins** : sans les enseignes, un avis supprimé a {virgule(vs['ecart_points'])} points "
        "de voisins supprimés en plus qu'un avis resté ; le hasard à fiche égale en donne "
        f"{virgule(vs['hasard_fiche_moyen_points'])}. En 5 étoiles seules, l'écart tombe à "
        f"{virgule(vs5['ecart_points'])} point.",
        "**Conclusion** : une tendance des éloges courts et génériques à être supprimés, qui "
        "rejoint la section 1, mais qui ne résiste pas à tous les découpages. Aucun groupe "
        "n'isole les suppressions.",
    ])

    tableau(doc, "Tableau 3. KMeans à 4 groupes, sans les six enseignes, toutes notes",
            ["Groupe de textes", "Avis", "Supprimés", "Attendus", "Seuil hasard (fiche)",
             "Fiches touchées", "Longueur"],
            [[K4[int(r["groupe"])][0], espace(r["n_avis"]),
              (espace(r["n_supprimes"]), r["au_dessus_fiche"] == "True"),
              virgule(r["attendus"]), r["hasard_fiche_haut"], r["fiches_touchees"],
              espace(r["longueur_mediane"])] for r in k4],
            [5.0, 1.5, 1.8, 1.8, 2.6, 2.2, 1.7],
            "Attendus : d'après le taux de suppression de chaque note. Seuil hasard (fiche) : le "
            "groupe ne dépasse ce nombre de suppressions que dans 2,5 % des tirages à fiche "
            "égale. En gras : au-dessus du seuil. Longueur : médiane, en caractères.")

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(FIGURE), width=Cm(15.5))
    doc.add_paragraph("Figure 2. Suppressions observées, attendues, et seuil du hasard à fiche "
                      "égale, par groupe.", style="Legende")

    def cellule(pop, methode, k):
        v = rang(pop, methode, k)
        return (f"{virgule(v, 1)} %", v > SEUIL)

    tableau(doc, "Tableau 4. Le groupe le plus supprimé dépasse-t-il le hasard à fiche égale ?",
            ["Découpage", "Sans enseignes, toutes notes", "Sans enseignes, 5 étoiles",
             "Toutes fiches, toutes notes", "Toutes fiches, 5 étoiles"],
            [[f"KMeans, K = {k}"] + [cellule(p, "kmeans", k) for p in (SE, SE5, TF, TF5)]
             for k in range(2, 11)]
            + [["HDBSCAN"] + [cellule(p, "hdbscan", "") for p in (SE, SE5, TF, TF5)]],
            [3.2, 3.3, 3.1, 3.3, 3.1],
            "Part des tirages à fiche égale où le groupe le plus supprimé dépasse moins ses "
            "attendus que dans l'observé. En gras : plus de 97,5 %, le découpage isole les "
            "suppressions au-delà de l'effet fiche.")

    doc.add_heading("Limites", level=2)
    puces(doc, [
        f"{n_supp_se} suppressions sans les enseignes, {n_supp_se5} en 5 étoiles seules : "
        "l'effet repose sur peu d'avis.",
        "Neuf valeurs de K testées ; les découpages ne sont pas indépendants.",
        "Première version corrigée : elle regroupait les notes 2 à 4 dans le calcul des "
        "attendus.",
    ])

    nouveaux = list(doc.element.body)[n_avant - 1:-1]   # hors sectPr final
    for e in nouveaux:
        avant.addprevious(e)
    doc.save(LIVRABLE)
    print(f"écrit : {LIVRABLE}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
