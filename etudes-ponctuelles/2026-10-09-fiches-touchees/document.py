#!/usr/bin/env python3
"""Word de l'analyse exploratoire : pourquoi certaines fiches perdent leurs nouveaux avis.

    PYTHONPATH=".venv/Lib/site-packages;outils" <python uv> etudes-ponctuelles/2026-10-09-fiches-touchees/document.py [nom.docx]

Demande de Matthieu du 2026-10-09 : les résultats de calcul.py dans un Word, en expliquant
pour chaque graphique ce que c'est, sur quelles données il est fait, et ce qu'on en tire.
Les chiffres sont lus dans les sorties de calcul.py (lancer calcul.py d'abord). Mise en
forme reprise de 2026-10-07-prenom-dans-les-fiches/document.py (Roboto).

Refuse d'écrire si le fichier existe déjà (il a pu être retouché à la main).
"""
import importlib.util
import sys
from pathlib import Path

import pandas as pd

RACINE = Path(__file__).resolve().parents[2]
_spec = importlib.util.spec_from_file_location(
    "mise_en_forme", RACINE / "etudes-ponctuelles" / "2026-10-07-prenom-dans-les-fiches" / "document.py")
mf = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(mf)
from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt  # noqa: E402

DOSSIER = RACINE / "livrables" / "restitution_finale_v2" / "analyse_exploratoire"
SORTIE = DOSSIER / "fiches_touchees_nouveaux_avis.docx"
BILAN_84 = RACINE / "etudes-ponctuelles" / "2026-10-08-bilan-84-fiches" / "sorties" / "bilan_84_fiches.csv"
SECTEURS = {"home_services": "services à domicile", "healthcare": "santé", "travel": "voyage",
            "wellness_fitness": "sport et bien-être", "automotive": "automobile",
            "food_beverage": "restauration", "hospitality": "hôtellerie"}


def nombre(x: float, nd: int = 0) -> str:
    s = f"{x:,.{nd}f}".replace(",", " ").replace(".", ",")
    return s


def pct(x: float) -> str:
    return f"{nombre(100 * x)} %"


def main() -> int:
    global SORTIE
    if len(sys.argv) > 1:
        SORTIE = DOSSIER / sys.argv[1]
    if SORTIE.exists():
        raise SystemExit(f"{SORTIE} existe déjà : il a pu être retouché. Choisir un autre nom.")

    sel = pd.read_csv(DOSSIER / "selection_fiches.csv")
    comp = pd.read_csv(DOSSIER / "comparaison_touchees_temoins.csv").set_index("indicateur")
    mod = pd.read_csv(DOSSIER / "modele_fiche.csv").set_index("parametre")
    fiches = pd.read_csv(DOSSIER / "fiches_indicateurs.csv", dtype={"cid": str})
    touchees = fiches[fiches["groupe"] == "touchée"]
    anciennes = pd.read_csv(BILAN_84, dtype={"cid": str})["cid"]
    dans_84 = int(touchees["cid"].isin(anciennes).sum())

    s = sel.assign(etape=sel["etape"].str.strip()).set_index("etape")
    coh = s.loc["fiches de la cohorte"]
    ens = s.loc["dont six enseignes"]
    etu = s.iloc[2]
    tou = s.loc["touchée"]
    tem = s.loc["témoin"]
    part_supp_touchees = tou["avis_supprimes"] / etu["avis_supprimes"]
    part_fiches_touchees = tou["fiches"] / etu["fiches"]
    n_temoins_apparies = int(comp["fiches_temoins"].max())
    regions = touchees["region"].value_counts()
    secteurs = touchees["industry"].value_counts()

    def c(ind, col):
        return comp.loc[ind, col]

    def rc(p):
        return nombre(mod.loc[p, "rapport_de_cotes"], 2)

    def ic(p):
        return f"de {nombre(mod.loc[p, 'borne_basse'], 2)} à {nombre(mod.loc[p, 'borne_haute'], 2)}"

    m_fiches, m_avis, m_supp = (int(mod["fiches"].iloc[0]), int(mod["avis"].iloc[0]),
                                int(mod["supprimes"].iloc[0]))

    doc = Document()
    for sec in doc.sections:
        sec.page_height, sec.page_width = Cm(29.7), Cm(21)
        sec.top_margin = sec.bottom_margin = Cm(2)
        sec.left_margin = sec.right_margin = Cm(2.2)
    normal = doc.styles["Normal"]
    normal.font.name = mf.POLICE
    normal.font.size = Pt(10.5)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), mf.POLICE)

    t = doc.add_paragraph()
    mf.police(t.add_run("Pourquoi certaines fiches perdent leurs nouveaux avis"), 18, True, mf.ENCRE)
    t.paragraph_format.space_after = Pt(2)
    mf.paragraphe(doc, "Analyse exploratoire – 9 octobre 2026", 10, apres=10, couleur=mf.ENCRE_2)
    mf.paragraphe(doc, [("Question : ", True),
                        ("Axel s'est-il emballé sur la sollicitation d'avis ? Avant de chiffrer ce "
                         "qu'elle rapporte, on regarde quelles fiches perdent beaucoup de leurs "
                         "nouveaux avis, et ce qui les distingue de fiches semblables qui n'en "
                         "perdent pas.", False)])

    # En bref --------------------------------------------------------------------
    mf.titre(doc, "En bref", 1)
    mf.puce(doc, f"Les pertes sont très concentrées : {nombre(tou['fiches'])} fiches, soit "
                 f"{pct(part_fiches_touchees)} des fiches étudiées, portent {pct(part_supp_touchees)} "
                 "des suppressions de nouveaux avis.")
    mf.puce(doc, "À volume habituel égal, ce n'est pas la quantité d'avis sollicités qui fait la "
                 "différence, mais la manière : des avis concentrés sur un même jour, des auteurs "
                 "sans niveau Local Guide, et un propriétaire qui répond moins.")
    mf.puce(doc, "La part d'avis qui citent un prénom ne distingue pas ces fiches.")
    mf.puce(doc, "Pour Axel : solliciter n'est pas un problème en soi ; solliciter en rafale, "
                 "auprès de comptes peu établis, l'est. Ce que la sollicitation rapporte en solde "
                 "d'avis reste à chiffrer.")

    # 1. Données ------------------------------------------------------------------
    mf.titre(doc, "1. Les données et les définitions", 1)
    mf.puce(doc, [("Les avis : ", True),
                  ("les avis 4 et 5 étoiles publiés du 11 au 17 août 2026, sur les fiches "
                   "relevées jusqu'au 24 août. Seuls les nouveaux avis comptent : aucun avis "
                   "publié avant le 11 août.", False)])
    mf.puce(doc, [("La suppression : ", True),
                  (f"un avis est supprimé si sa disparition est notée au plus tard 7 jours après "
                   f"sa publication. Chaque avis a donc le même temps pour être supprimé. Le "
                   f"taux d'une fiche est le nombre d'avis supprimés pour 100 avis publiés.", False)])
    mf.puce(doc, [("Les fiches étudiées : ", True),
                  ("celles qui ont publié au moins 5 avis dans la semaine, pour que leur taux ait "
                   "un sens. Les six enseignes déjà connues (quatre chaînes antiparasitaires "
                   "américaines, deux salles de sport espagnoles) sont mises de côté.", False)])
    mf.puce(doc, [("Les fiches touchées : ", True),
                  ("au moins 2 avis supprimés et au moins 10 % de leurs avis. ", False),
                  ("Les fiches témoins : ", True), ("aucun avis supprimé.", False)])
    mf.donnees(doc, f"corpus 03C (03C_7_17_aout_prenom.parquet, doublons retirés, prénoms annotés "
                    f"par Axel), restreint aux avis publiés du 11 au 17 août : {nombre(coh['avis_publies'])} "
                    f"avis sur {nombre(coh['fiches'])} fiches, dont {nombre(coh['avis_supprimes'])} "
                    "supprimés dans les 7 jours.")

    # 2. Concentration ---------------------------------------------------------------
    mf.titre(doc, "2. Une poignée de fiches concentre les pertes", 1)
    lignes = [
        ["Toute la cohorte", nombre(coh["fiches"]), nombre(coh["avis_publies"]),
         nombre(coh["avis_supprimes"]), nombre(coh["supprimes_pour_100"], 1)],
        ["dont six enseignes", nombre(ens["fiches"]), nombre(ens["avis_publies"]),
         nombre(ens["avis_supprimes"]), nombre(ens["supprimes_pour_100"], 1)],
        ["Fiches étudiées (au moins 5 avis, hors enseignes)", nombre(etu["fiches"]),
         nombre(etu["avis_publies"]), nombre(etu["avis_supprimes"]), nombre(etu["supprimes_pour_100"], 1)],
        ["  dont fiches touchées", nombre(tou["fiches"]), nombre(tou["avis_publies"]),
         nombre(tou["avis_supprimes"]), nombre(tou["supprimes_pour_100"], 1)],
        ["  dont fiches témoins", nombre(tem["fiches"]), nombre(tem["avis_publies"]),
         nombre(tem["avis_supprimes"]), nombre(tem["supprimes_pour_100"], 1)],
    ]
    mf.tableau(doc, ["", "Fiches", "Avis publiés", "Avis supprimés", "Supprimés pour 100"],
               lignes, [6.6, 2.0, 2.6, 2.6, 2.8])
    mf.paragraphe(doc, "", apres=2)
    mf.puce(doc, f"Les {nombre(tou['fiches'])} fiches touchées perdent {nombre(tou['supprimes_pour_100'], 1)} "
                 f"avis pour 100 publiés, contre {nombre(etu['supprimes_pour_100'], 1)} pour l'ensemble "
                 f"des fiches étudiées. Elles portent {pct(part_supp_touchees)} des suppressions.")
    mf.puce(doc, f"{regions.get('US', 0)} sont aux États-Unis et {regions.get('Europe', 0)} en Europe. "
                 "Les secteurs les plus représentés sont "
                 + ", ".join(f"{SECTEURS.get(k, k)} ({v})" for k, v in secteurs.head(3).items()) + ".")
    mf.puce(doc, f"{dans_84} des {nombre(tou['fiches'])} fiches touchées faisaient déjà partie des "
                 "84 fiches de l'étude du 8 octobre, choisies sur le nombre d'avis perdus de tous "
                 "âges. Le choix par le nombre ne trompait donc pas sur l'essentiel.")
    mf.puce(doc, f"Les six enseignes perdent {nombre(ens['supprimes_pour_100'], 1)} avis pour 100 : "
                 f"à elles seules, {pct(ens['avis_supprimes'] / coh['avis_supprimes'])} des "
                 "suppressions de la cohorte.")
    mf.paragraphe(doc, [("Ce qu'on en tire : ", True),
                        ("la perte de nouveaux avis n'est pas un phénomène diffus. Elle touche un petit "
                         "nombre de fiches, surtout américaines, en santé et en services à domicile. "
                         "C'est sur elles qu'il faut chercher l'explication.", False)], avant=4)

    # 3. Comparaison -------------------------------------------------------------------
    mf.titre(doc, "3. Ce qui distingue les fiches touchées de fiches semblables", 1)
    mf.paragraphe(doc, [("Ce que montre le graphique. ", True),
                        ("Pour chaque indicateur, deux points : la valeur moyenne des fiches touchées "
                         "(bleu) et celle des fiches témoins (orange). Plus l'écart entre les deux "
                         "points est grand, plus l'indicateur distingue les fiches touchées.", False)])
    mf.paragraphe(doc, [("Comment les témoins sont choisis. ", True),
                        ("Une fiche touchée n'est comparée qu'à des témoins de même région, même "
                         "secteur et même quart de volume habituel d'avis. Les témoins sont pondérés "
                         "pour avoir la même répartition que les fiches touchées. Ainsi, les deux "
                         "groupes reçoivent habituellement autant d'avis : "
                         f"{nombre(c('volume_semaine', 'touchees'), 1)} contre "
                         f"{nombre(c('volume_semaine', 'temoins'), 1)} par semaine en médiane.", False)])
    mf.paragraphe(doc, [("Les indicateurs. ", True),
                        ("Sollicitation : jour de pic (au moins 3 avis et au moins 4 fois la moyenne "
                         "habituelle de la fiche), part des avis arrivés le jour le plus chargé. "
                         "Auteurs : comptes neufs (2 avis ou moins au total), auteurs sans niveau "
                         "Local Guide, auteur qui publie sur une autre fiche le même jour. Contenu : "
                         "5 étoiles, sans texte, texte court, prénom cité, photo, avis au même texte. "
                         "Propriétaire : taux de réponse et réponses identiques, mesurés sur l'année "
                         "précédant le 11 août pour ne pas dépendre des suppressions étudiées.", False)])
    mf.donnees(doc, f"les {nombre(tou['fiches'])} fiches touchées et {n_temoins_apparies} fiches "
                    "témoins de leurs strates (même région, secteur et quart de volume), sur les avis "
                    "4-5 étoiles publiés du 11 au 17 août 2026.")
    mf.image(doc, DOSSIER / "comparaison_touchees_temoins.png",
             "Graphique 1 – Indicateurs des fiches touchées et des fiches témoins semblables "
             "(moyenne des parts de chaque fiche, ou part des fiches concernées).")
    mf.tableau(doc, ["Indicateur", "Fiches touchées", "Fiches témoins"], [
        ["Au moins un jour de pic dans la semaine", pct(c("au_moins_un_pic", "touchees")),
         pct(c("au_moins_un_pic", "temoins"))],
        ["Taux de réponse du propriétaire (année précédente)", pct(c("taux_reponse", "touchees")),
         pct(c("taux_reponse", "temoins"))],
        ["Auteurs sans niveau Local Guide", pct(c("part_sans_niveau", "touchees")),
         pct(c("part_sans_niveau", "temoins"))],
        ["Prénom cité", pct(c("part_prenom", "touchees")), pct(c("part_prenom", "temoins"))],
        ["Comptes neufs (2 avis ou moins)", pct(c("part_comptes_neufs", "touchees")),
         pct(c("part_comptes_neufs", "temoins"))],
        ["Avis de la semaine, rapportés au rythme habituel (médiane)",
         f"× {nombre(c('intensite', 'touchees'), 2)}", f"× {nombre(c('intensite', 'temoins'), 2)}"],
    ], [9.0, 3.5, 3.5])
    mf.paragraphe(doc, "", apres=2)
    mf.paragraphe(doc, [("Ce qu'on en tire : ", True),
                        ("à volume égal, les fiches touchées ont plus souvent un jour de pic, leurs "
                         "auteurs ont plus souvent un profil sans niveau Local Guide, et leur "
                         "propriétaire répond moins. En revanche, elles ne citent pas plus de prénoms "
                         "et n'ont pas plus de comptes neufs. Ce graphique compare les indicateurs un "
                         "par un : il ne dit pas lesquels comptent quand on les prend ensemble. "
                         "C'est l'objet de la partie 4.", False)], avant=4)

    # 4. Modèle -----------------------------------------------------------------------------
    mf.titre(doc, "4. Tous les indicateurs ensemble : ce qui compte vraiment", 1)
    mf.paragraphe(doc, [("Ce que montre le graphique. ", True),
                        ("Un modèle estime, pour chaque fiche, le risque qu'un de ses nouveaux avis "
                         "soit supprimé, avec tous les indicateurs en même temps, à région, secteur, "
                         "taille et volume habituel égaux. Chaque point est l'effet d'un indicateur, "
                         "les autres étant tenus égaux. 1 veut dire pas d'effet ; 1,91 veut dire un "
                         "risque multiplié par 1,91. Le trait est la marge d'incertitude à 95 % ; un "
                         "point plein est significatif. Entre parenthèses, l'écart auquel se lit "
                         "chaque effet : « +10 points » de part, un doublement, un jour, ou la présence "
                         "(oui / non).", False)])
    mf.donnees(doc, f"les fiches étudiées qui ont au moins 5 avis dans l'année précédente (pour "
                    f"mesurer leurs réponses) : {m_fiches} fiches, {nombre(m_avis)} avis publiés du "
                    f"11 au 17 août 2026, {m_supp} supprimés dans les 7 jours. Hors six enseignes. "
                    "Marges d'incertitude robustes.")
    mf.image(doc, DOSSIER / "modele_fiche.png",
             "Graphique 2 – Effet de chaque indicateur sur le risque qu'un nouvel avis de la fiche "
             "soit supprimé, tous les indicateurs ensemble.")
    mf.paragraphe(doc, [("Ce qui ressort :", True)], apres=3)
    mf.puce(doc, f"Auteurs sans niveau Local Guide : risque × {rc('part_sans_niveau')} pour 10 points "
                 f"de plus ({ic('part_sans_niveau')}). C'est l'effet le plus net.")
    mf.puce(doc, f"Avis concentrés sur un même jour : × {rc('part_jour_max')} pour 10 points de plus "
                 f"({ic('part_jour_max')}). Une fois cette concentration prise en compte, le nombre "
                 f"de jours de pic (× {rc('jours_pic')}) et l'intensité de la semaine "
                 f"(× {rc('log2_intensite')} quand elle double) ne jouent plus.")
    mf.puce(doc, f"Taux de réponse du propriétaire : × {rc('taux_reponse')} pour 10 points de plus "
                 f"({ic('taux_reponse')}). Une fiche qui répond perd moins.")
    mf.puce(doc, f"Prénom cité (× {rc('part_prenom')}), comptes neufs (× {rc('part_comptes_neufs')}), "
                 f"5 étoiles (× {rc('part_5_etoiles')}), photo, réponses identiques : pas d'effet "
                 "démontré.")
    mf.puce(doc, f"Deux résultats à ne pas mettre en avant : les textes courts protègent "
                 f"(× {rc('part_texte_court')}), ce que nous n'expliquons pas ; deux avis au même "
                 f"texte multiplient le risque par {rc('texte_double')}, mais seulement 2 fiches "
                 "touchées sont concernées.")
    mf.paragraphe(doc, [("Ce qu'on en tire : ", True),
                        ("le volume d'avis sollicités ne suffit pas à expliquer les pertes. Ce qui "
                         "compte, c'est la manière : des avis qui arrivent d'un coup, le même jour, "
                         "écrits par des comptes au profil peu établi, sur une fiche dont le "
                         "propriétaire répond peu. C'est le schéma qu'un filtre anti-fraude cherche "
                         "à repérer.", False)], avant=4)

    # 5. Conclusion --------------------------------------------------------------------------
    mf.titre(doc, "5. Ce qu'on peut répondre à Axel", 1)
    mf.puce(doc, "Solliciter des avis ne fait pas perdre d'avis en soi : à volume égal, recevoir plus "
                 "d'avis que d'habitude n'a pas d'effet démontré sur le risque de suppression.")
    mf.puce(doc, "Ce qui fait perdre des avis, c'est la sollicitation en rafale : beaucoup d'avis le "
                 "même jour, par des comptes sans niveau Local Guide. Étaler les demandes dans le "
                 "temps et répondre aux avis va dans le bon sens.")
    mf.puce(doc, "Les prénoms ne sont pas en cause : leur excès dans les fiches les plus touchées "
                 "venait de leur gros volume d'avis.")
    mf.paragraphe(doc, [("Limites", True)], apres=3, avant=6)
    mf.puce(doc, f"{nombre(tou['fiches'])} fiches touchées, c'est peu : les marges d'incertitude sont "
                 "larges, et les écarts modestes peuvent ne pas se confirmer.")
    mf.puce(doc, "Un niveau Local Guide vide peut être un état transitoire du profil, voire une "
                 "conséquence d'un compte signalé. C'est le résultat le plus fort, et le premier à "
                 "vérifier.")
    mf.puce(doc, "Une semaine de publication seulement (11 au 17 août) : les fiches qui sollicitent "
                 "à d'autres moments ne sont pas vues comme telles.")
    mf.puce(doc, "Ce document ne chiffre pas ce que la sollicitation rapporte en solde d'avis "
                 "(avis gagnés moins avis perdus). C'est l'étape suivante.")

    mf.paragraphe(doc, [("Pour refaire les calculs. ", True),
                        ("etudes-ponctuelles/2026-10-09-fiches-touchees/calcul.py (sorties : "
                         "selection_fiches.csv, comparaison_touchees_temoins.csv, modele_fiche.csv, "
                         "fiches_indicateurs.csv, une ligne par fiche avec son nom), puis document.py "
                         "pour ce Word.", False)],
                  taille=8, apres=0, couleur=mf.ENCRE_2, avant=10)

    doc.save(SORTIE)
    print(f"écrit : {SORTIE}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
