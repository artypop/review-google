#!/usr/bin/env python3
"""Ajoute les régressions du 07C au document `livrables/analyses-complémentaires.docx`.

    .venv/Scripts/python.exe logistic-regression-study/python/07C_note.py

Lit `output-study/2026-09-28-sorties-07C/07C_effets_6_passages.csv`. Aucun
modèle n'est ajusté ici. Le document existant est ouvert tel quel : la section
« 4. Régressions logistiques » est retirée si elle existe, puis réécrite à la
fin. Le reste du document n'est pas touché.
"""
from __future__ import annotations

import sys
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402

PY = Path(__file__).resolve().parent
RACINE = PY.parents[1]
SORTIES = PY.parent / "output-study" / "2026-09-28-sorties-07C"
LIVRABLE = RACINE / "livrables" / "analyses-complémentaires.docx"
TITRE_SECTION = "4. Régressions logistiques réduites, panel 03B"
NUM_FIGURE, NUM_TABLEAUX = 3, (5, 6)

LIBELLES = {
    "etoiles_1": "1 étoile", "etoiles_2": "2 étoiles", "etoiles_3": "3 étoiles",
    "etoiles_4": "4 étoiles",
    "log_ratio_pic_journalier_fiche": "Pic d'avis sur la fiche le jour du dépôt",
    "has_photo": "Photo jointe à l'avis",
    "log_photos_auteur": "Photos publiées par l'auteur",
    "guide_4_et_plus": "Local Guide niveau 4 et plus",
    "a_une_reponse": "Réponse du propriétaire",
    "secteur_food_beverage": "Restauration", "secteur_healthcare": "Santé",
    "secteur_home_services": "Services à domicile", "secteur_hospitality": "Hôtellerie",
    "secteur_travel": "Voyage", "secteur_wellness_fitness": "Sport et bien-être",
}
GROUPES = [("Note (référence : 5 étoiles)", ["etoiles_1", "etoiles_2", "etoiles_3", "etoiles_4"]),
           ("L'avis et son auteur", ["log_ratio_pic_journalier_fiche", "has_photo",
                                     "log_photos_auteur", "guide_4_et_plus", "a_une_reponse"]),
           ("Secteur (référence : automobile)",
            ["secteur_food_beverage", "secteur_healthcare", "secteur_home_services",
             "secteur_hospitality", "secteur_travel", "secteur_wellness_fitness"])]
PERIMETRES = [("US", "États-Unis"), ("Europe", "Europe"), ("tous", "Ensemble")]


def virgule(x, d=2):
    return f"{x:.{d}f}".replace(".", ",")


def espace(n):
    return f"{int(n):,}".replace(",", " ")


def cellule(effets, passage, variable):
    ligne = effets[(effets["passage"] == passage) & (effets["variable"] == variable)]
    if ligne.empty:
        return "—", False
    r = ligne.iloc[0]
    net = r["borne_basse"] > 1 or r["borne_haute"] < 1
    return (f"×{virgule(r['risque_relatif'])} [{virgule(r['borne_basse'])} à "
            f"{virgule(r['borne_haute'])}]", net)


# ---------------------------------------------------------------------------
# Figure
# ---------------------------------------------------------------------------
SURFACE = "#fcfcfb"; ENCRE = "#0b0b0b"; ENCRE_DOUCE = "#52514e"; GRILLE = "#e4e3df"
BLEU = "#2a78d6"; ORANGE = "#eb6834"


def figure(effets, perimetre, titre, fichier):
    variables = [v for _, vs in GROUPES for v in vs]
    y = np.arange(len(variables))[::-1].astype(float)
    fig, ax = plt.subplots(figsize=(6.8, 6.0), dpi=150)
    fig.patch.set_facecolor(SURFACE); ax.set_facecolor(SURFACE)
    fig.subplots_adjust(left=0.36, right=0.97, top=0.84, bottom=0.08)
    for c in ("top", "right", "left"):
        ax.spines[c].set_visible(False)
    ax.spines["bottom"].set_color(GRILLE)
    ax.tick_params(colors=ENCRE_DOUCE, labelsize=8.5, length=0)
    ax.set_xscale("log")
    ax.axvline(1, color=ENCRE, lw=1.4, zorder=2)
    ax.grid(True, axis="x", color=GRILLE, lw=1, zorder=0)
    retraits = ("Sans les salles espagnoles" if perimetre == "Europe"
                else "Sans salles espagnoles ni antiparasitaires US")
    for decal, passage, couleur, nom in ((0.17, f"{perimetre}_complet", BLEU, "Toutes les fiches"),
                                         (-0.17, f"{perimetre}_sans_retraits", ORANGE, retraits)):
        e = effets[effets["passage"] == passage].set_index("variable").reindex(variables)
        ax.errorbar(e["risque_relatif"], y + decal,
                    xerr=[e["risque_relatif"] - e["borne_basse"],
                          e["borne_haute"] - e["risque_relatif"]],
                    fmt="o", ms=4, color=couleur, ecolor=couleur, elinewidth=1.2,
                    capsize=0, zorder=3, label=nom)
    ax.set_yticks(y, [LIBELLES[v] for v in variables])
    # Même échelle pour les trois figures : celle qui contient toutes les fourchettes.
    affiches = effets[effets["variable"].isin(variables)]
    gauche = affiches["borne_basse"].min() * 0.85
    droite = max(10.5, affiches["borne_haute"].max() * 1.15)
    graduations = [g for g in (0.1, 0.2, 0.5, 1, 2, 5, 10, 20) if gauche <= g <= droite]
    ax.set_xticks(graduations, [f"×{virgule(g, 1).replace(',0', '')}" for g in graduations])
    ax.set_xlim(gauche, droite)
    ax.set_ylim(-0.7, len(variables) - 0.3)
    # Séparateurs de familles
    i = len(variables)
    for _, vs in GROUPES[:-1]:
        i -= len(vs)
        ax.axhline(i - 0.5, color=GRILLE, lw=1)
    leg = ax.legend(frameon=False, fontsize=8, loc="lower left", ncol=2,
                    bbox_to_anchor=(-0.55, 1.0), borderaxespad=0.3)
    for t in leg.get_texts():
        t.set_color(ENCRE_DOUCE)
    fig.text(0.012, 0.975, titre, fontsize=11, color=ENCRE, va="top")
    fig.text(0.012, 0.93, "Point : risque relatif. Trait : fourchette à 95 %. Barre verticale : "
             "×1, aucun effet. Échelle logarithmique.", fontsize=8.5, color=ENCRE_DOUCE, va="top")
    fig.savefig(fichier, dpi=150, facecolor=SURFACE)
    plt.close(fig)
    return fichier


# ---------------------------------------------------------------------------
# Document
# ---------------------------------------------------------------------------
NAVY = RGBColor(0x32, 0x29, 0x54)


def P(doc, texte, style=None):
    p = doc.add_paragraph(style=style)
    morceaux = texte.split("**")
    for i, m in enumerate(morceaux):
        if m:
            p.add_run(m).bold = i % 2 == 1
    return p


def fond(cell, couleur):
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), couleur)
    cell._tc.get_or_add_tcPr().append(sh)


def tableau(doc, titre, effets, jeu, note, nouvelle_page=False):
    doc.add_paragraph(titre, style="TitreTableau").paragraph_format.page_break_before = nouvelle_page
    lignes = []                           # (valeurs, type) : type = entete, groupe, ligne, effectif
    lignes.append((["", *[n for _, n in PERIMETRES]], "entete"))
    passages = [f"{c}_{jeu}" for c, _ in PERIMETRES]
    for nom, vs in GROUPES:
        lignes.append(([nom, "", "", ""], "groupe"))
        for v in vs:
            cells = [cellule(effets, p, v) for p in passages]
            lignes.append(([LIBELLES[v], *cells], "ligne"))
    for libelle, col in (("Avis", "avis"), ("Suppressions", "suppressions"), ("Fiches", "fiches")):
        valeurs = [espace(effets.loc[effets["passage"] == p, col].iloc[0]) for p in passages]
        lignes.append(([libelle, *valeurs], "effectif"))

    largeurs = [5.2, 3.9, 3.9, 3.9]
    t = doc.add_table(rows=len(lignes), cols=4)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    bords = OxmlElement("w:tblBorders")
    for cote in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{cote}")
        if cote in ("left", "right", "insideV"):
            e.set(qn("w:val"), "nil")
        else:
            e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), "D9D8DE")
        bords.append(e)
    t._tbl.tblPr.append(bords)
    zebre = 0
    for i, (valeurs, genre) in enumerate(lignes):
        if genre == "groupe":
            premiere = t.rows[i].cells[0]
            premiere.merge(t.rows[i].cells[3])
        zebre = zebre + 1 if genre == "ligne" else 0
        for j, val in enumerate(valeurs):
            if genre == "groupe" and j > 0:
                break
            c = t.rows[i].cells[j]
            c.width = Cm(largeurs[j]) if genre != "groupe" else Cm(sum(largeurs))
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.RIGHT
            texte, gras = (val if isinstance(val, tuple) else (val, False))
            r = p.add_run(texte); r.font.size = Pt(9)
            if genre == "entete":
                r.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); fond(c, "322954")
            elif genre == "groupe":
                r.bold = True; r.font.color.rgb = NAVY; fond(c, "E6E5EC")
            elif genre == "effectif":
                r.font.color.rgb = RGBColor(0x6B, 0x6A, 0x66)
            else:
                r.bold = gras
                if zebre % 2 == 0:
                    fond(c, "F4F4F7")
    for j, gc in enumerate(t._tbl.tblGrid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(int(largeurs[j] * 567)))
    doc.add_paragraph(note, style="NoteTableau")


def retirer_section(doc):
    corps = doc.element.body
    debut = next((p._p for p in doc.paragraphs if p.text.strip() == TITRE_SECTION), None)
    if debut is None:
        return
    a_retirer, courant = [], debut
    while courant is not None:
        if courant.tag == qn("w:sectPr"):
            break
        a_retirer.append(courant)
        courant = courant.getnext()
    for e in a_retirer:
        corps.remove(e)


def main() -> int:
    effets = pd.read_csv(SORTIES / "07C_effets_6_passages.csv")
    passages = pd.read_csv(SORTIES / "07C_passages.csv").set_index("passage")
    source = (SORTIES / "07C_source.txt").read_text(encoding="utf-8")
    figures = [
        (figure(effets, perimetre, f"Risque relatif de suppression, {nom_fig} du panel 03B",
                SORTIES / f"07C_figure_{perimetre}.png"), legende)
        for perimetre, nom_fig, legende in (
            ("tous", "ensemble des avis", "ensemble des avis"),
            ("US", "avis des fiches américaines", "États-Unis"),
            ("Europe", "avis des fiches européennes", "Europe"))]

    def rr(passage, v):
        r = effets[(effets["passage"] == passage) & (effets["variable"] == v)].iloc[0]
        return f"×{virgule(r['risque_relatif'])}"

    ph = effets[effets["variable"] == "log_photos_auteur"]
    if not (ph["borne_haute"] < 1).all():
        raise SystemExit("La phrase sur les photos de l'auteur ne tient plus : à réécrire.")
    photos = ph["risque_relatif"]

    doc = Document(LIVRABLE)
    retirer_section(doc)

    h = doc.add_heading(TITRE_SECTION, level=1)
    h.paragraph_format.page_break_before = True
    ps = passages.loc["tous_sans_retraits"]
    P(doc, "Avis publiés du 4 au 17 août, suivis jusqu'au 24 août : "
           f"{espace(passages.loc['tous_complet', 'avis'])} avis, "
           f"{espace(passages.loc['tous_complet', 'suppressions'])} suppressions. Chaque modèle "
           "est ajusté une fois sur toutes les fiches, une fois sans les 2 salles de sport "
           f"espagnoles ni les {int(ps['fiches_antiparasitaires_retirees'])} fiches "
           "antiparasitaires américaines.")
    P(doc, "Lecture : ×2 veut dire que l'avis est supprimé deux fois plus souvent qu'un avis "
           "de référence qui lui ressemble sur le reste. En gras, les effets dont la fourchette "
           "exclut ×1.")

    note = ("Entre crochets : fourchette à 95 %, marges groupées par fiche. Photos de l'auteur "
            "et pic d'avis : risque par cran de log(1 + n), soit de 0 à 2, de 2 à 6, de 6 à 19, "
            "de 19 à 54. Pic d'avis : avis reçus par la fiche le jour du dépôt, rapportés à sa "
            "moyenne journalière de l'année précédente. Local Guide niveau 4 et plus : comparé à "
            "tous les autres auteurs. Réponse : vue ou non au dernier relevé. Contrôles non "
            "affichés : âge de l'avis à sa première observation, jours de suivi, avis vus "
            "tardivement, région pour l'ensemble.")
    tableau(doc, f"Tableau {NUM_TABLEAUX[0]}. Risque relatif de suppression, toutes les fiches",
            effets, "complet", note)
    tableau(doc, f"Tableau {NUM_TABLEAUX[1]}. Même modèle, sans les 2 salles espagnoles ni les "
                 "fiches antiparasitaires US", effets, "sans_retraits",
            f"Retirés : {espace(ps['avis_retires'])} avis et "
            f"{espace(ps['suppressions_retirees'])} suppressions, dont "
            f"{espace(passages.loc['Europe_sans_retraits', 'avis_retires'])} avis en Europe. "
            "Antiparasitaires : repérage de la section 3 du rapport (4 chaînes, leurs "
            "variantes, autres enseignes repérées par mots-clés, ABC Home).", nouvelle_page=True)

    doc.add_paragraph("À retenir", style="TitreTableau")
    for texte in (
            f"**1 étoile** : {rr('tous_complet', 'etoiles_1')} sur l'ensemble, "
            f"{rr('tous_sans_retraits', 'etoiles_1')} sans les fiches retirées.",
            f"**Photos publiées par l'auteur** : de ×{virgule(photos.min())} à "
            f"×{virgule(photos.max())} par cran selon le passage, fourchette sous ×1 dans les six.",
            f"**Services à domicile** : {rr('tous_complet', 'secteur_home_services')} sur "
            f"l'ensemble, {rr('tous_sans_retraits', 'secteur_home_services')} sans les fiches "
            "antiparasitaires.",
            f"**Pic d'avis** : {rr('Europe_complet', 'log_ratio_pic_journalier_fiche')} en Europe, "
            f"{rr('US_complet', 'log_ratio_pic_journalier_fiche')} aux États-Unis.",
            f"**Réponse du propriétaire** : {rr('tous_complet', 'a_une_reponse')} sur "
            f"l'ensemble, {rr('tous_sans_retraits', 'a_une_reponse')} sans les fiches retirées. "
            "À lire avec la première réserve."):
        P(doc, texte, "List Bullet")

    for k, (fichier, legende) in enumerate(figures):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.page_break_before = True
        p.add_run().add_picture(str(fichier), width=Cm(15.5))
        doc.add_paragraph(f"Figure {NUM_FIGURE + k}. Risque relatif de suppression, {legende}.",
                          style="Legende")

    doc.add_heading("Réserves", level=1)
    for texte in (
            "Réponse du propriétaire : état au dernier relevé. Un avis qui reste en ligne a plus "
            "de jours pour recevoir une réponse, ce qui fait paraître la réponse protectrice. "
            "Le chiffre mesure une association, pas une protection ; la mesure à jalon fixe "
            "de la section 6 du rapport reste la référence.",
            "Modèle réduit : les effets ne se comparent pas un à un à ceux de la section 5 du "
            "rapport, qui contrôle aussi le texte, l'auteur, la taille de la fiche et la "
            "rafale d'avis de l'auteur."):
        P(doc, texte, "List Bullet")
    P(doc, "Sources : logistic-regression-study/python/ (03B_features_local.py, "
           "03B_features_local_controle.py, 07C_regression_reduite.py, 07C_note.py), "
           "output-study/2026-09-28-sorties-07C/. " + source.replace("\n", " ").strip(),
      "NoteTableau")

    doc.save(LIVRABLE)
    print(f"écrit : {LIVRABLE}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
