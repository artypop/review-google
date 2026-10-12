#!/usr/bin/env python3
"""Document Word : les avis qui ont un jumeau, fiches antiparasitaires américaines.

    PYTHONPATH=.venv/Lib/site-packages <python uv> etudes-ponctuelles/2026-09-30-avis-antiparasitaires/note.py

Demande de Matthieu du 2026-10-01. Ne calcule rien : relit les CSV de
`sorties/copies/`, écrits par `calcul_copies.py` et `calcul_copies_robustesse.py`
le 2026-09-30. Sortie : `livrables/restitution/jumeaux-antiparasitaires.docx`,
et la figure `sorties/copies/figure-types.png`.
"""
from __future__ import annotations

import csv
import sys
from datetime import date
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402

DOSSIER = Path(__file__).resolve().parent
RACINE = DOSSIER.parents[1]
COPIES = DOSSIER / "sorties" / "copies"
SORTIE = RACINE / "livrables" / "restitution" / "jumeaux-antiparasitaires.docx"

POLICE = "Calibri"
NAVY, GRIS, FOND_FAMILLE = "322954", "6B6A66", "EDEDF2"
# Figure : palette de référence du skill dataviz (emplacement 1) et un gris neutre
# pour la référence « attendu ».
BLEU, GRIS_REF = "#2a78d6", "#b4b2ab"
ENCRE, ENCRE_2, GRILLE, FOND = "#0b0b0b", "#52514e", "#e4e3df", "#fcfcfb"
SEUIL_RANG = 97.5
TYPES = ["un auteur, plusieurs fiches", "gabarit sur une fiche",
         "gabarit sur une enseigne", "rencontre"]


def lire(nom):
    with (COPIES / nom).open(encoding="utf-8") as f:
        return list(csv.DictReader(f))


def nombre(x) -> str:
    return f"{int(float(x)):,}".replace(",", " ")


def virgule(x, d=1) -> str:
    return f"{float(x):.{d}f}".replace(".", ",")


def etabli(r) -> str:
    """Écart établi : le hasard ne l'explique ni à période égale ni à fiche égale."""
    p, f = float(r["rang_periode_pct"]), float(r["rang_fiche_pct"])
    if p >= SEUIL_RANG and f >= SEUIL_RANG:
        return "oui"
    return "non"


def rapport(r) -> str:
    return f"×{virgule(r['rapport_observe_attendu'])}" if r["rapport_observe_attendu"] else "—"


# ---------------------------------------------------------------------------
# Mise en forme, reprise des notes de livrables/
# ---------------------------------------------------------------------------

def styler(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(1.8)
    st = doc.styles
    n = st["Normal"]
    n.font.name, n.font.size = POLICE, Pt(10.5)
    n.font.color.rgb = RGBColor(0x1F, 0x1F, 0x1F)
    n.element.rPr.rFonts.set(qn("w:eastAsia"), POLICE)
    n.paragraph_format.space_after = Pt(6)
    for nom, taille, avant, apres in (("Title", 22, 0, 4), ("Heading 1", 15, 16, 6),
                                      ("Heading 2", 11.5, 10, 4)):
        s = st[nom]
        s.font.name, s.font.size, s.font.bold, s.font.italic = POLICE, Pt(taille), True, False
        s.font.color.rgb = RGBColor.from_string(NAVY)
        rf = s.element.get_or_add_rPr().find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            s.element.get_or_add_rPr().append(rf)
        for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rf.set(qn(a), POLICE)
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            rf.attrib.pop(qn(a), None)
        s.paragraph_format.space_before, s.paragraph_format.space_after = Pt(avant), Pt(apres)
        s.paragraph_format.keep_with_next = True
        for b in s.element.get_or_add_pPr().findall(qn("w:pBdr")):
            s.element.get_or_add_pPr().remove(b)
    tt = st.add_style("TitreTableau", 1)
    tt.base_style = n
    tt.font.size, tt.font.bold = Pt(9.5), True
    tt.font.color.rgb = RGBColor.from_string(NAVY)
    tt.paragraph_format.space_before = Pt(10)
    tt.paragraph_format.keep_with_next = True
    nt = st.add_style("NoteTableau", 1)
    nt.base_style = n
    nt.font.size, nt.font.color.rgb = Pt(9), RGBColor.from_string(GRIS)
    nt.paragraph_format.space_after = Pt(10)


def fond(cellule, couleur):
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), couleur)
    cellule._tc.get_or_add_tcPr().append(sh)


def tableau(doc, titre, entetes, lignes, largeurs, note=None):
    """`lignes` : listes de cellules ; une cellule est un texte ou (texte, gras)."""
    doc.add_paragraph(titre, style="TitreTableau")
    t = doc.add_table(rows=1 + len(lignes), cols=len(entetes))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    bords = OxmlElement("w:tblBorders")
    for cote in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{cote}")
        if cote in ("left", "right", "insideV"):
            e.set(qn("w:val"), "nil")
        else:
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), "4")
            e.set(qn("w:color"), "D9D8DE")
        bords.append(e)
    t._tbl.tblPr.append(bords)
    for i, rangee in enumerate([entetes] + lignes):
        for j, contenu in enumerate(rangee):
            c = t.rows[i].cells[j]
            c.width = Cm(largeurs[j])
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.RIGHT
            texte, gras = contenu if isinstance(contenu, tuple) else (contenu, False)
            r = p.add_run(str(texte))
            r.font.size = Pt(9)
            r.bold = gras
            if i == 0:
                r.bold, r.font.color.rgb = True, RGBColor(0xFF, 0xFF, 0xFF)
                fond(c, NAVY)
    for j, gc in enumerate(t._tbl.tblGrid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(int(largeurs[j] * 567)))
    entete = t.rows[0]._tr.get_or_add_trPr()
    repete = OxmlElement("w:tblHeader")
    repete.set(qn("w:val"), "true")
    entete.append(repete)
    if note:
        doc.add_paragraph(note, style="NoteTableau")


def puce(doc, morceaux):
    p = doc.add_paragraph(style="List Bullet")
    for texte, gras in ([(morceaux, False)] if isinstance(morceaux, str) else morceaux):
        p.add_run(texte).bold = gras
    p.paragraph_format.space_after = Pt(3)


def nouvelle_page(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ---------------------------------------------------------------------------
# Figure : part supprimée observée et attendue, par type de jumeau
# ---------------------------------------------------------------------------

def figure(c1) -> Path:
    lignes = {r["groupe"]: r for r in c1 if r["longueur_min"] == "80"}
    groupes = ["sans jumeau", "toutes copies"] + TYPES
    noms = ["sans jumeau", "avec un jumeau", "un auteur,\nplusieurs fiches",
            "gabarit sur\nune fiche", "gabarit sur\nune enseigne", "rencontre"]
    obs = [float(lignes[g]["part_supprimee_pct"]) for g in groupes]
    att = [100 * float(lignes[g]["attendus_periode_note"]) / float(lignes[g]["avis"])
           for g in groupes]
    fig, ax = plt.subplots(figsize=(8.5, 4.2), dpi=150)
    fig.patch.set_facecolor(FOND)
    ax.set_facecolor(FOND)
    x = range(len(groupes))
    larg = 0.38
    ax.bar([i - 0.2 for i in x], obs, width=larg, color=BLEU, label="observé", zorder=3,
           edgecolor=FOND, linewidth=1)
    ax.bar([i + 0.2 for i in x], att, width=larg, color=GRIS_REF,
           label="attendu, à période et note égales", zorder=3, edgecolor=FOND, linewidth=1)
    for i, (o, a) in enumerate(zip(obs, att)):
        ax.annotate(f"{o:.2f} %".replace(".", ","), (i - 0.2, o), xytext=(0, 3),
                    textcoords="offset points", ha="center", fontsize=7, color=ENCRE)
        ax.annotate(f"{a:.2f} %".replace(".", ","), (i + 0.2, a), xytext=(0, 3),
                    textcoords="offset points", ha="center", fontsize=7, color=ENCRE_2)
    ax.set_xticks(list(x))
    ax.set_xticklabels(noms, fontsize=8, color=ENCRE)
    ax.tick_params(axis="y", colors=ENCRE_2, labelsize=8, length=0)
    ax.tick_params(axis="x", length=0)
    ax.yaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(
        lambda v, _: f"{v:g} %".replace(".", ",")))
    ax.grid(axis="y", color=GRILLE, lw=0.8, zorder=0)
    for cote in ["top", "right", "left"]:
        ax.spines[cote].set_visible(False)
    ax.spines["bottom"].set_color(GRILLE)
    ax.set_ylabel("part des avis supprimés", color=ENCRE_2, fontsize=8)
    h = fig.get_figheight()
    fig.text(0.01, 1 - 0.18 / h, "Part des avis supprimés selon leur jumeau",
             va="top", fontsize=11, color=ENCRE)
    fig.text(0.01, 1 - 0.44 / h, "Fiches antiparasitaires américaines, avis de 80 caractères "
             "ou plus. Jumeau : un autre avis semblable à 0,97 ou plus.",
             va="top", fontsize=8, color=ENCRE_2)
    poignees, textes = ax.get_legend_handles_labels()
    fig.legend(poignees, textes, loc="upper left", bbox_to_anchor=(0.005, 1 - 0.60 / h),
               ncol=2, frameon=False, fontsize=8, labelcolor=ENCRE, handlelength=1.4)
    fig.tight_layout(rect=(0, 0, 1, 1 - 0.95 / h))
    chemin = COPIES / "figure-types.png"
    fig.savefig(chemin, facecolor=FOND)
    plt.close(fig)
    return chemin


# ---------------------------------------------------------------------------
# Document
# ---------------------------------------------------------------------------

def main() -> int:
    c1, c2, c3 = lire("C1-types.csv"), lire("C2-chaines.csv"), lire("C3-periodes.csv")
    c4, c5, c7 = lire("C4-seuils.csv"), lire("C5-ordre.csv"), lire("C7-robustesse.csv")
    g80 = {r["groupe"]: r for r in c1 if r["longueur_min"] == "80"}
    g30 = {r["groupe"]: r for r in c1 if r["longueur_min"] == "30"}
    tous, copies, sans = g80["tous les avis"], g80["toutes copies"], g80["sans jumeau"]
    ch = {r["chaine"]: r for r in c2}
    per = {r["periode"]: r for r in c3}
    rob = {r["population"]: r for r in c7}
    ordre = {r["cas"]: r for r in c5}
    seuil = {(r["longueur_min"], r["similarite_min"]): r for r in c4}
    fig = figure(c1)

    doc = Document()
    styler(doc)
    doc.add_paragraph("Les avis qui ont un jumeau", style="Title")
    doc.add_paragraph(f"Fiches antiparasitaires américaines. Calculs du 30/09/2026, document "
                      f"produit le {date.today():%d/%m/%Y}.", style="NoteTableau")

    # --- ce qui a été fait ---------------------------------------------------
    doc.add_paragraph("Ce que nous avons fait", style="Heading 1")
    puce(doc, [("174 fiches", True),
               (" de services à domicile américaines : celles dont le nom contient « pest », "
                "plus la chaîne Bulwark Exterminating. ", False),
               ("259 901 avis avec texte", True),
               (", publiés depuis 2004, dont 1 086 supprimés pendant le relevé du 11 au 24 "
                "août 2026.", False)])
    puce(doc, "Chaque texte a été traduit en vecteur par un modèle de langue (e5-base), sur une base "
         "construite pour l'occasion. Deux textes proches par le sens ont des vecteurs proches ; "
         "leur ressemblance va de 0 à 1.")
    puce(doc, [("Un avis a un jumeau", True),
          (" quand un autre avis lui ressemble à 0,97 ou plus : même texte, ou presque, "
           "à quelques mots près. Règle reprise du panel 03B, où elle avait fait ressortir les "
           "copies (149 avis, 38 % supprimés).", False)])
    puce(doc, "Les jumeaux sont classés selon le lien entre les deux avis, du plus précis au moins "
         "précis (tableau 1). Un avis qui a plusieurs jumeaux prend le type le plus précis.")
    puce(doc, "Comparaison : la part d'avis supprimés dépend surtout de l'âge de l'avis (11,5 % des "
         "avis d'août 2026, 0,04 % de ceux d'avant 2023) et un peu de sa note. Pour chaque "
         "groupe, on calcule donc les suppressions attendues si chacun de ses avis avait le "
         "taux des avis de même période et de même note, puis on compare aux suppressions "
         "observées.")
    puce(doc, "Contrôle du hasard : les étiquettes « supprimé » sont redistribuées 2 000 fois au "
         "hasard, une fois entre avis de même période et même note, une fois à l'intérieur "
         "de chaque fiche. Un écart est dit établi quand l'observé dépasse 97,5 % des tirages "
         "dans les deux cas.")

    tableau(doc, "Tableau 1. Les types de jumeaux",
            ["Type", "Auteur", "Fiche"],
            [["doublon", "le même", "la même"],
             ["un auteur, plusieurs fiches", "le même", "une autre"],
             ["gabarit sur une fiche", "un autre", "la même"],
             ["gabarit sur une enseigne", "un autre", "une autre, de la même enseigne"],
             ["rencontre", "un autre", "une autre enseigne"]],
            [5.6, 3.5, 7.5],
            "Auteur : repéré par le lien de son profil Google. Enseigne : deux premiers mots du "
            "nom de la fiche. Aucun doublon dans ces données.")

    # --- à retenir ------------------------------------------------------------
    doc.add_paragraph("À retenir", style="Heading 1")
    puce(doc, [(f"Un avis de 80 caractères ou plus qui a un jumeau est supprimé "
                f"{virgule(copies['rapport_observe_attendu'])} fois plus souvent que prévu", True),
               (f" : {nombre(copies['supprimes'])} suppressions sur {nombre(copies['avis'])} "
                f"avis ({virgule(copies['part_supprimee_pct'], 2)} %), pour "
                f"{virgule(copies['attendus_periode_note'])} attendues. Sans jumeau : "
                f"{virgule(sans['part_supprimee_pct'], 2)} %, conforme aux attendus.", False)])
    t_a, t_f = g80["un auteur, plusieurs fiches"], g80["gabarit sur une fiche"]
    puce(doc, [("L'excès vient des jumeaux liés à un auteur ou à une fiche", True),
               (f" : le même auteur sur plusieurs fiches {rapport(t_a)} "
                f"({t_a['supprimes']} supprimés sur {t_a['avis']}), le même texte par plusieurs "
                f"auteurs sur une fiche {rapport(t_f)} ({t_f['supprimes']} sur {t_f['avis']}). "
                "Les jumeaux sans lien, sur d'autres fiches de l'enseigne ou d'autres "
                "enseignes, ne montrent pas d'écart établi.", False)])
    a26, avant = rob["avis publiés en août 2026"], rob["avis publiés avant août 2026"]
    puce(doc, [("L'effet porte sur les avis récents", True),
               (f" : en août 2026, {a26['supprimes']} copies supprimées sur {a26['copies']} "
                f"({virgule(a26['part_supprimee_pct'])} %), "
                f"{virgule(a26['rapport_observe_attendu'])} fois les attendus. Avant "
                f"août 2026, {avant['supprimes']} sur {nombre(avant['copies'])} pour "
                f"{virgule(avant['attendus_periode_note'])} attendues : pas d'écart établi.",
                False)])
    puce(doc, [("Il porte aussi sur les chaînes", True),
               (f" : Bulwark {rapport(ch['Bulwark'])}, Insight {rapport(ch['Insight'])}, "
                f"EcoShield {rapport(ch['EcoShield'])}, Pointe aucune copie supprimée. Sur les "
                f"fiches hors des quatre chaînes, {ch['autres']['supprimes']} copie supprimée "
                f"sur {ch['autres']['copies']}.", False)])
    sf = rob["sans la fiche la plus touchée"]
    puce(doc, f"Une fiche porte {int(copies['supprimes']) - int(sf['supprimes'])} des "
              f"{copies['supprimes']} copies supprimées. Sans elle, l'écart tient : "
              f"{rapport(sf)}, établi.")
    pr, ap = ordre["premier déposé"], ordre["déposé après un jumeau"]
    js = ordre["un jumeau supprimé"]
    puce(doc, f"Le jumeau déposé en second est plus supprimé que le premier : "
              f"{virgule(ap['part_supprimee_pct'])} % contre {virgule(pr['part_supprimee_pct'])} %. "
              f"Quand un jumeau est supprimé, l'autre l'est aussi dans "
              f"{virgule(js['part_supprimee_pct'])} % des cas.")
    s99 = seuil[("80", "0.99")]
    puce(doc, f"Plus la ressemblance est forte, plus l'écart grandit : à 0,99 ou plus, "
              f"{virgule(s99['part_supprimee_pct'])} % des avis supprimés, "
              f"{virgule(s99['rapport_observe_attendu'])} fois les attendus. Avec les textes courts (30 caractères ou plus), l'écart tombe à "
              f"{rapport(g30['toutes copies'])} : les formules courantes s'y ajoutent.")

    nouvelle_page(doc)
    doc.add_paragraph("Le détail", style="Heading 1")
    doc.add_picture(str(fig), width=Cm(16.6))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    entetes = ["Groupe", "Avis", "Supprimés", "Part supprimée", "Attendus", "Rapport",
               "Écart établi"]
    larg = [4.9, 2.0, 1.9, 2.2, 1.9, 1.8, 1.9]

    def lignes_types(g):
        out = []
        for nom, lib in [("tous les avis", "tous les avis"), ("toutes copies", "avec un jumeau")] \
                + [(t, f"  {t}") for t in TYPES] + [("sans jumeau", "sans jumeau")]:
            r = g[nom]
            if nom == "tous les avis":
                out.append([lib, nombre(r["avis"]), nombre(r["supprimes"]),
                            f"{virgule(r['part_supprimee_pct'], 2)} %", "", "", ""])
                continue
            e = etabli(r)
            out.append([lib, nombre(r["avis"]), nombre(r["supprimes"]),
                        f"{virgule(r['part_supprimee_pct'], 2)} %",
                        virgule(r["attendus_periode_note"]), (rapport(r), e == "oui"), e])
        return out

    note_types = ("Attendus : suppressions si chaque avis avait le taux des avis de même "
                  "période de publication et même note. Rapport : observés ÷ attendus. Écart "
                  "établi : les supprimés dépassent 97,5 % de 2 000 tirages au hasard, à "
                  "période et note égales comme à l'intérieur de chaque fiche.")
    tableau(doc, "Tableau 2. Avis de 80 caractères ou plus, selon leur jumeau", entetes,
            lignes_types(g80), larg, note_types)
    tableau(doc, "Tableau 3. Même tableau, avis de 30 caractères ou plus", entetes,
            lignes_types(g30), larg,
            "À 30 caractères, les copies passent de 1 840 à 19 821 avis : les formules courtes "
            "et courantes s'y ajoutent. Parmi les types, seul le même auteur sur plusieurs "
            "fiches garde un écart établi.")

    tableau(doc, "Tableau 4. Copies de 80 caractères ou plus, par chaîne",
            ["Chaîne", "Avis de la chaîne", "Copies", "Copies supprimées", "Attendus",
             "Rapport", "Écart établi"],
            [[r["chaine"] if r["chaine"] != "autres" else "hors des 4 chaînes",
              nombre(r["avis_de_la_chaine"]), nombre(r["copies"]), nombre(r["supprimes"]),
              virgule(r["attendus_periode_note"]), (rapport(r), etabli(r) == "oui"), etabli(r)]
             for r in c2],
            [3.6, 2.6, 1.8, 2.4, 1.9, 1.9, 2.4])

    tableau(doc, "Tableau 5. Copies de 80 caractères ou plus, par période de publication",
            ["Période", "Avis", "Part supprimée", "Copies", "Copies supprimées",
             "Part des copies", "Attendus"],
            [[r["periode"], nombre(r["avis"]), f"{virgule(r['part_supprimee_pct'], 2)} %",
              nombre(r["copies"]), nombre(r["copies_supprimees"]),
              f"{virgule(r['part_copies_supprimee_pct'], 1)} %" if r["part_copies_supprimee_pct"]
              else "—", virgule(r["attendus_copies"])] for r in c3],
            [3.6, 2.2, 2.3, 1.8, 2.5, 2.2, 2.0],
            "Les avis anciens ne peuvent être vus supprimés que pendant les 14 jours du relevé : "
            "leur part supprimée est faible, et celle de leurs copies aussi.")

    tableau(doc, "Tableau 6. Selon le seuil de ressemblance",
            ["Longueur", "Ressemblance", "Avis", "Supprimés", "Part supprimée", "Attendus",
             "Rapport", "Écart établi"],
            [[f"{r['longueur_min']} car. ou plus", f"{virgule(r['similarite_min'], 2)} ou plus",
              nombre(r["avis"]), nombre(r["supprimes"]),
              f"{virgule(r['part_supprimee_pct'], 2)} %", virgule(r["attendus_periode_note"]),
              (rapport(r), etabli(r) == "oui"), etabli(r)] for r in c4],
            [2.4, 2.4, 1.8, 1.8, 2.2, 1.8, 1.7, 2.5],
            "Ici l'avis compte dès qu'il a un autre avis assez ressemblant, sans distinguer le "
            "type de lien.")

    tableau(doc, "Tableau 7. Ordre de dépôt et sort du jumeau",
            ["Cas", "Avis", "Supprimés", "Part supprimée"],
            [[r["cas"], nombre(r["avis"]), nombre(r["supprimes"]),
              f"{virgule(r['part_supprimee_pct'])} %"] for r in c5],
            [7.6, 3.0, 3.0, 3.0],
            "Copies de 80 caractères ou plus. Premier déposé : publié avant tous ses jumeaux, "
            "ou le même jour.")

    tableau(doc, "Tableau 8. L'écart tient-il ?",
            ["Population", "Copies", "Supprimées", "Part supprimée", "Attendus", "Rapport",
             "Écart établi"],
            [[r["population"], nombre(r["copies"]), nombre(r["supprimes"]),
              f"{virgule(r['part_supprimee_pct'], 2)} %", virgule(r["attendus_periode_note"]),
              (rapport(r), etabli(r) == "oui"), etabli(r)] for r in c7],
            [5.0, 1.8, 2.0, 2.2, 1.9, 1.8, 1.9],
            "Copies de 80 caractères ou plus. Fiche la plus touchée : celle qui a le plus de "
            "copies supprimées, retirée entièrement.")

    # --- limites et sources ----------------------------------------------------
    doc.add_paragraph("Limites", style="Heading 1")
    puce(doc, "Suppression : avis dont le robot a noté la disparition pendant le relevé. Cette "
              "analyse n'applique pas la règle « absent deux jours de suite ou plus » de la "
              "régression : une disparition au dernier relevé compte.")
    puce(doc, f"Petits effectifs : {copies['supprimes']} copies supprimées, réparties sur "
              f"{copies['fiches_avec_suppression']} fiches, dont une en porte "
              f"{int(copies['supprimes']) - int(sf['supprimes'])}.")
    puce(doc, "Un auteur est repéré par le lien de son profil : une personne avec deux comptes "
              "compte pour deux, et ses avis jumeaux passent pour des gabarits.")
    puce(doc, "La population se limite aux fiches antiparasitaires américaines. Elle comprend "
              "les 4 chaînes comptées parmi les 6 enseignes signalées de l'étude : le tableau 4 "
              "sépare leurs résultats de ceux des autres fiches.")
    puce(doc, "Le seuil de 0,97 vient du panel 03B. Le tableau 6 montre que les résultats en "
              "dépendent.")
    doc.add_paragraph(
        "Sources : etudes-ponctuelles/2026-09-30-avis-antiparasitaires/ (extraire.py, "
        "calcul_copies.py, calcul_copies_robustesse.py, note.py ; sorties/copies/). Vecteurs : "
        "data/embeddings/pest_e5base.parquet.", style="NoteTableau")

    SORTIE.parent.mkdir(parents=True, exist_ok=True)
    try:
        doc.save(SORTIE)
    except PermissionError:
        sys.exit(f"{SORTIE.name} est ouvert dans Word : le fermer, puis relancer.")
    print(f"écrit : {SORTIE}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
