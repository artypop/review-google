#!/usr/bin/env python3
"""Écrit livrables/restitution_finale_v2/passation.docx, la passation pour le Claude d'Axel.

    PYTHONPATH=".venv/Lib/site-packages" <python uv> etudes-ponctuelles/2026-10-12-dataset-features-axel/passation_docx.py

Demande de Matthieu du 2026-10-12 : même forme que « Description Corpus Claude.docx » de
Romain, qu'elle complète sans la répéter. Axel ne reçoit aucun script Python : le document ne
renvoie qu'aux fichiers transmis et à du SQL. Chiffres vérifiés le 2026-10-12 sur les fichiers
transmis. Police Roboto (Roboto Mono pour le code).
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

RACINE = Path(__file__).resolve().parents[2]
SORTIE = RACINE / "livrables" / "restitution_finale_v2" / "passation.docx"
POLICE, MONO = "Roboto", "Roboto Mono"
ENCRE, CODE = RGBColor(0x1F, 0x1F, 0x1F), RGBColor(0x37, 0x47, 0x4F)

doc = Document()
for s in doc.sections:
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.2)
    s.top_margin = s.bottom_margin = Cm(2)


def polices(element, nom):
    rpr = element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), nom)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rf.attrib.pop(qn(a), None)


normal = doc.styles["Normal"]
normal.font.size = Pt(10.5)
normal.font.color.rgb = ENCRE
polices(normal.element, POLICE)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.15
for nom, taille, avant in [("Heading 1", 18, 0), ("Heading 2", 14, 16), ("Heading 3", 11.5, 10)]:
    st = doc.styles[nom]
    st.font.size, st.font.bold, st.font.color.rgb = Pt(taille), True, ENCRE
    polices(st.element, POLICE)
    st.paragraph_format.space_before, st.paragraph_format.space_after = Pt(avant), Pt(6)
polices(doc.styles["List Bullet"].element, POLICE)
polices(doc.styles["Table Grid"].element, POLICE)

MOTIF = re.compile(r"(\*\*.+?\*\*|`.+?`)")


def remplir(p, texte, taille=None):
    for morceau in MOTIF.split(texte):
        if not morceau:
            continue
        if morceau.startswith("**"):
            r = p.add_run(morceau[2:-2])
            r.bold = True
        elif morceau.startswith("`"):
            r = p.add_run(morceau[1:-1])
            r.font.name = MONO
            polices(r._element, MONO)
            r.font.color.rgb = CODE
            r.font.size = Pt((taille or 10.5) - 1)
        else:
            r = p.add_run(morceau)
        if taille and not morceau.startswith("`"):
            r.font.size = Pt(taille)
    return p


def titre(texte, niveau):
    doc.add_heading(texte, level=niveau)


def para(texte):
    return remplir(doc.add_paragraph(), texte)


def puces(lignes):
    for l in lignes:
        p = remplir(doc.add_paragraph(style="List Bullet"), l)
        p.paragraph_format.space_after = Pt(3)


def code(texte):
    lignes = texte.strip("\n").split("\n")
    for i, l in enumerate(lignes):
        p = doc.add_paragraph()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "F1F3F4")
        p._p.get_or_add_pPr().append(shd)  # avant spacing et ind, ordre du schéma
        pf = p.paragraph_format
        pf.left_indent = Cm(0.4)
        pf.space_before = Pt(4) if i == 0 else Pt(0)
        pf.space_after = Pt(8) if i == len(lignes) - 1 else Pt(0)
        pf.line_spacing = 1.0
        r = p.add_run(l if l else " ")
        r.font.name = MONO
        polices(r._element, MONO)
        r.font.size = Pt(8.5)
        r.font.color.rgb = CODE


def tableau(entetes, lignes, largeurs):
    t = doc.add_table(rows=1, cols=len(entetes))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for ligne_cells, valeurs, entete in [(t.rows[0].cells, entetes, True)] + [
            (None, v, False) for v in lignes]:
        cells = ligne_cells if ligne_cells is not None else t.add_row().cells
        for c, v, w in zip(cells, valeurs, largeurs):
            c.width = Cm(w)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            remplir(p, f"**{v}**" if entete and v else str(v), taille=9)
            if entete:
                tcpr = c._tc.get_or_add_tcPr()
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:color"), "auto")
                shd.set(qn("w:fill"), "EDEDED")
                tcpr.append(shd)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


# ---------------------------------------------------------------------------------------
titre("Passation : les chiffres de la restitution et les données pour les refaire", 1)
para("Rédigé le 2026-10-12 pour reprendre le travail de Matthieu. Les résultats sont dans "
     "Restitution.pdf. Les deux corpus, leur construction, les six enseignes et le choix des "
     "avis publiés du 7 au 17 août sont décrits dans « Description Corpus Claude », qui n'est "
     "pas repris ici.")
para("Tous les chiffres de ce document ont été recalculés le 2026-10-12 à partir des seuls "
     "fichiers transmis, et retrouvés à l'identique. Seuls ceux des modèles jour par jour "
     "(§ 6.2) sont repris des sorties de l'étude.")
para("Les choix du § 3 ont été arrêtés avec Axel et Romain pendant l'étude. Pour que de "
     "nouveaux chiffres se comparent à ceux de la restitution, il faut les appliquer tels quels.")
para("Questions : Matthieu.")

# ---------------------------------------------------------------------------------------
titre("1. Les fichiers transmis", 2)
tableau(["Fichier", "Contenu", "Lignes"], [
    ["`reviews_doublons_cleaned_all.parquet`", "Corpus 1, la base complète. Copie exacte de la "
     "table BigQuery du même nom, horodatages en UTC.", "4 876 933"],
    ["`03C_7_17_aout.parquet`", "Corpus 2 sans le prénom. Copie de la table BigQuery du même "
     "nom, dates en DATE.", "27 828"],
    ["`businesses.parquet`", "Les fiches : `cid`, `place_id`, `name`, `country`, `industry`, "
     "`bucket`, `review_count_at_build`.", "9 048"],
    ["`features_regression_reviews_doublons_cleaned_all.parquet`", "Corpus 1, plus le prénom et "
     "les variables calculées du § 2.", "4 876 933"],
    ["`features_03C_7_17_aout.parquet`", "Les avis des régressions et des chiffres « nouveaux "
     "avis » de la restitution, mêmes colonnes.", "24 191"],
    ["`features_regression.sql`", "Reconstruit les deux fichiers de features avec DuckDB (§ 2).",
     ""],
    ["`regression_secteur_taille.csv`", "Tous les paramètres du modèle du graphique 3.1, "
     "contrôles compris (§ 6.1).", ""],
], [6.2, 8.4, 2.0])
para("Les deux fichiers de features n'existent qu'en parquet. `review_id` est unique dans "
     "chaque fichier ; `cid` relie un avis à sa fiche.")

# ---------------------------------------------------------------------------------------
titre("2. Les fichiers de features", 2)
para("Ils reprennent les 25 colonnes de `reviews_doublons_cleaned_all`, dans le même ordre, "
     "avec les horodatages passés en DATE (UTC). Suivent le prénom et les variables calculées.")
para("Le prénom suit la règle de `03C_7_17_aout_prenom` (Description Corpus, § 3), appliquée à "
     "toute la base : `prenom` et `prenom_source` y ont le même sens.")
tableau(["Colonne", "Définition"], [
    ["`region`", "`US` si `businesses.country = 'US'`, sinon `Europe`."],
    ["`industry`, `bucket`", "Secteur et taille de la fiche, repris de `businesses`."],
    ["`enseigne`", "La fiche fait partie des six enseignes."],
    ["`fin_suivi`", "Dernier jour où la fiche a été relevée. Fichier complet seulement."],
    ["`fiche_plus_relevee`", "Vrai si `fin_suivi` est avant le 24 août : 3 fiches, 548 avis. "
     "Fichier complet seulement."],
    ["`j_premiere_vue`, `j_reponse`, `j_suppression`", "Jours entre la publication et "
     "`first_seen_at`, `reply_date`, `deleted_detected_at`."],
    ["`jours_de_suivi`", "Jours entre `first_seen_at` et `fin_suivi`. Fichier complet "
     "seulement."],
    ["`n_jour`", "Avis 4-5★ publiés sur la fiche le même jour, avis compris."],
    ["`par_jour`", "Moyenne journalière des avis 4-5★ de la fiche sur les 365 jours avant le "
     "11 août 2026."],
    ["`pic`", "1 si `n_jour >= 3` et `n_jour >= 4 × par_jour`."],
    ["`supprime`", "1 si l'avis a disparu au plus tard 7 jours après sa publication. Rempli "
     "seulement pour les avis publiés du 7 au 17 août."],
    ["`disparu`", "1 si `deleted_detected_at` est renseigné. Fichier complet seulement."],
    ["`photo`", "1 si au moins une photo est jointe (`n_photos > 0`)."],
    ["`reponse_j7`", "1 si le propriétaire a répondu au plus tard au 7e jour."],
    ["`sans_reponse_j7`", "Son contraire."],
    ["`longueur_texte`", "Caractères du texte sans les espaces de début et de fin ; 0 sans "
     "texte."],
    ["`tranche_texte`", "`sans`, `1_40`, `41_80`, `81_130`, `131_240`, `241_plus`."],
    ["`tranche_guide`", "`local_guide_level` en `1_3`, `4_6`, `7_10` ; vide sans niveau."],
], [5.2, 11.4])
para("`features_03C_7_17_aout` contient les lignes du fichier complet qui vérifient "
     "`supprime IS NOT NULL AND star >= 4 AND prenom IS NOT NULL AND NOT fiche_plus_relevee`, "
     "sans les colonnes `fin_suivi`, `fiche_plus_relevee`, `jours_de_suivi` et `disparu`.")

titre("Pourquoi deux variables de suppression", 3)
para("`supprime` ne se mesure que pour un avis publié pendant le relevé. Un avis de 2025 "
     "supprimé à son 3e jour n'est dans aucune table : sur le reste de la base, « supprimé au "
     "plus tard au 7e jour » ne peut pas être observé (Description Corpus, « Pourquoi le corpus "
     "ne prend pas les avis d'avant le 7 août »).")
para("`disparu` vaut pour toute la base : l'avis a disparu pendant le relevé, quel que soit son "
     "âge. Le risque dépend fortement de l'âge de l'avis : `disparu` se lit avec "
     "`j_premiere_vue` (âge au premier passage) et `jours_de_suivi`.")

titre("Reconstruire les fichiers", 3)
para("`features_regression.sql` se lance avec DuckDB, dans un dossier qui contient "
     "`reviews_doublons_cleaned_all.parquet`, `businesses.parquet` et "
     "`reviews_name_enriched.parquet` (l'annotation d'Axel, exportée de la table BigQuery du "
     "même nom) :")
code("duckdb < features_regression.sql")
para("Il faut quelques minutes et environ 6 Go de mémoire. Le résultat est identique aux deux "
     "fichiers transmis (vérifié le 2026-10-12). La requête commence par "
     "`SET TimeZone = 'UTC'` : sans cette ligne, DuckDB passe en heure de Paris et une partie "
     "des avis change de jour.")

# ---------------------------------------------------------------------------------------
titre("3. Les choix de calcul", 2)
para("**Un avis disparu est un avis supprimé.** `deleted_detected_at` renseigné suffit. La base "
     "est déjà dédoublonnée (Description Corpus, § 2) : aucune règle de deux jours d'absence ne "
     "s'applique.")
para("**Avis 4 et 5 étoiles.** Les avis 1 à 3 étoiles restent dans les fichiers ; seul le "
     "graphique A.2 les utilise.")
para("**Deux populations**, décrites au § 4. Un chiffre porte sur l'une ou sur l'autre.")
para("**Fiches plus relevées.** Les 3 fiches qui cessent d'être relevées le 17, le 18 ou le "
     "22 août sont écartées de tous les taux : leurs derniers jours ne sont pas observés. Dans le "
     "corpus 2, cela retire 1 fiche et 2 avis.")
para("**Prénom inconnu.** Dans le corpus 2, les 55 avis 4-5★ absents de l'annotation d'Axel "
     "(`prenom_source = 'introuvable'`) sont écartés.")
para("**Jour de pic** : la fiche reçoit au moins 3 avis 4-5★ dans la journée et au moins 4 fois "
     "sa moyenne journalière des 365 jours avant le 11 août. Il mesure un afflux ponctuel, "
     "distinct de la vitesse habituelle de collecte.")
para("**Réponse.** La référence est « réponse reçue au plus tard au 7e jour », la situation la "
     "plus fréquente (50,1 % des avis du corpus 2 retenus). La variable du modèle est donc "
     "`sans_reponse_j7`.")
para("**Local Guide.** Niveau au dernier relevé. Dans les données, les niveaux 1 à 3 sont des "
     "comptes qui ne sont pas Local Guide, 4 et plus des Local Guides. Les avis sans niveau sont "
     "à part.")
para("**Six enseignes.** Définition dans Description Corpus. Chaque chiffre du § 5 précise s'il "
     "les comprend.")

# ---------------------------------------------------------------------------------------
titre("4. Les deux populations", 2)
para("Les requêtes du § 5 partent de ces deux vues (DuckDB, dans le dossier des fichiers) :")
code("""SET TimeZone = 'UTC';
-- Nouveaux avis
CREATE VIEW n AS SELECT * FROM 'features_03C_7_17_aout.parquet';
-- Stock
CREATE VIEW s AS SELECT * FROM 'features_regression_reviews_doublons_cleaned_all.parquet'
  WHERE star >= 4 AND NOT fiche_plus_relevee;""")
tableau(["Population", "Vue", "Avis", "Supprimés", "Mesure"], [
    ["Nouveaux avis, six enseignes comprises", "`n`", "24 191 (4 705 fiches)", "631",
     "Part supprimée au plus tard au 7e jour (`supprime`)"],
    ["Nouveaux avis, sans les six enseignes", "`n WHERE NOT enseigne`", "22 378", "379",
     "Idem"],
    ["Stock", "`s`", "4 195 342 (8 984 fiches)", "3 520",
     "Part disparue pendant le relevé (`disparu`)"],
], [4.0, 3.4, 3.2, 1.8, 4.2])
para("**Nouveaux avis** : avis 4-5★ du corpus 2 (publiés du 7 au 17 août 2026), après les "
     "exclusions du § 3. Les deux boutiques The Boxer Club n'y ont aucun avis : « sans les six "
     "enseignes » revient à retirer les quatre chaînes.")
para("**Stock** : tous les avis 4-5★ en ligne pendant le relevé, de tous âges. Un écart mesuré "
     "sur le stock compare des avis d'âges différents.")

# ---------------------------------------------------------------------------------------
titre("5. Les chiffres de la restitution", 2)

titre("Fait 1. La première semaine", 3)
para("**641 des 1 272** avis 4-5★ disparus entre 1 et 13 jours après leur publication l'ont été "
     "au 6e ou au 7e jour. Population : stock, avis publiés à partir du 1er juillet 2026.")
code("""SELECT count(*) FILTER (j_suppression IN (6, 7)), count(*)
FROM s
WHERE disparu = 1 AND created_at >= DATE '2026-07-01' AND j_suppression BETWEEN 1 AND 13;""")
para("Les nombres bruts se comparent d'un âge à l'autre : chaque âge de 2 à 13 jours est "
     "observé sur environ 28 000 avis.")
para("**Graphique 1.1** : la version en taux. Avis 4-5★ publiés à partir du 11 août 2026, âges "
     "1 à 12 ; à chaque âge, avis supprimés pour 1 000 avis encore en ligne et observés à cet "
     "âge : 18,6 au 7e jour, 9,9 au 6e, 2,9 au plus aux autres âges.")

titre("Fait 2. Répondre", 3)
para("**Graphique 2.1.** Nouveaux avis sans les six enseignes, vus et encore en ligne à la fin "
     "de leur 2e jour : 17 915 avis.")
tableau(["Réponse au plus tard au 2e jour", "Avis", "Supprimés du 3e au 7e jour",
         "Pour 1 000"], [
    ["Oui", "7 642", "109", "14,3"],
    ["Non", "10 273", "223", "21,7"],
], [5.6, 2.6, 4.8, 2.6])
para("À secteur et taille égaux : ×0,51 [0,31–0,84]. C'est le modèle du § 6.1, avec "
     "`reponse_j2` à la place de `sans_reponse_j7` et la suppression du 3e au 7e jour comme "
     "variable expliquée.")
code("""SELECT coalesce(j_reponse <= 2, FALSE) AS repondu_j2, count(*) AS avis,
       count(*) FILTER (j_suppression BETWEEN 3 AND 7) AS supprimes
FROM n
WHERE NOT enseigne AND j_premiere_vue <= 2 AND NOT coalesce(j_suppression <= 2, FALSE)
GROUP BY 1;""")
para("**Graphique 2.2.** Modèle jour par jour (§ 6.2), sans les six enseignes. Face à un avis "
     "encore sans réponse : réponse le jour même ×0,46 ; le lendemain ×0,52 ; au 2e jour ×0,51 ; "
     "au 3e jour ou après ×0,25. Le taux de réponse habituel de la fiche ne joue pas : face à "
     "0 %, 1 à 25 % ×1,12 ; 25 à 75 % ×0,91 ; 75 % et plus ×1,04.")
para("**Graphique A.3** : le même modèle, six enseignes comprises.")

titre("Fait 3. Le prénom", 3)
para("**×1,05 [0,80–1,38]** à secteur et taille égaux : nouveaux avis, six enseignes comprises, "
     "modèle du § 6.1. C'est la ligne `tous, prenom` de `regression_secteur_taille.csv` et le "
     "graphique 3.1.")
para("**×1,01 [0,85–1,21]** à fiche égale : modèle jour par jour avec un indicateur par fiche, "
     "sur les 223 fiches qui ont au moins une suppression (§ 6.2).")
para("**Le constat par fiche** (les fiches dont les avis positifs avec texte citent souvent un "
     "prénom perdent davantage d'avis, et leur vitesse de collecte ne suffit pas à l'expliquer) "
     "a été mesuré sur toute la base, avis 4-5★ avec texte de tous âges, avec une ancienne règle "
     "de suppression (deux jours d'absence). Il n'a pas été refait avec les fichiers transmis.")

titre("Fait 4. Les services à domicile", 3)
para("Trois mesures, sur des populations différentes.")
para("**a. Nouveaux avis, supprimés au plus tard au 7e jour.**")
tableau(["Groupe", "Supprimés / avis", "%"], [
    ["Services à domicile US, six enseignes comprises", "334 / 3 826", "8,7"],
    ["Services à domicile US, sans les six enseignes", "82 / 2 013", "4,1"],
    ["Idem, sans les autres entreprises antiparasitaires", "44 / 1 656", "2,7"],
    ["Santé US", "70 / 2 081", "3,4"],
    ["Tous secteurs et régions, sans les six enseignes", "379 / 22 378", "1,7"],
], [10.0, 4.0, 2.0])
para("Une « autre entreprise antiparasitaire » est une fiche dont le nom correspond à "
     "`(?i)(pest|exterminat|termite|spidexx|mosquito|rodent|wildlife)` ou "
     "`(?i)abc home ?(&|and) ?commercial`, hors six enseignes.")
code("""SELECT n.industry, n.enseigne,
       regexp_matches(b.name, '(?i)(pest|exterminat|termite|spidexx|mosquito|rodent|wildlife)')
         OR regexp_matches(b.name, '(?i)abc home ?(&|and) ?commercial') AS antiparasitaire,
       count(*), sum(n.supprime)
FROM n JOIN 'businesses.parquet' b USING (cid)
WHERE n.region = 'US' AND n.industry IN ('home_services', 'healthcare')
GROUP BY ALL;""")
para("**b. Stock : avis perdus pour 100 avis publiés, du 11 au 24 août.** Le numérateur compte "
     "les avis 4-5★ disparus pendant le relevé, quelle que soit leur date de publication ; le "
     "dénominateur, les avis 4-5★ publiés du 11 au 24 août. Le ratio mesure le poids des pertes "
     "face au flux entrant : un avis perdu peut être ancien.")
tableau(["", "Toutes fiches", "Sans les six enseignes"], [
    ["Services à domicile", "34,7", "22,2"],
    ["Santé", "12,0", "12,0"],
    ["Sport et bien-être", "10,7", "10,6"],
    ["Tous secteurs", "12,2", "8,7"],
    ["Europe, tous secteurs (de 3,2 en hôtellerie à 8,1 en sport)", "5,4", "5,3"],
    ["États-Unis, services à domicile", "38,9", "26,5"],
    ["États-Unis, sport et bien-être", "15,1", "15,1"],
    ["États-Unis, santé", "13,5", "13,5"],
    ["États-Unis, tous secteurs", "17,1", "11,6"],
], [9.6, 3.2, 3.8])
para("Les deux boutiques The Boxer Club, classées en sport et bien-être, expliquent l'écart de "
     "10,7 à 10,6. Graphique 4.1 : cette mesure par secteur et par taille, toutes fiches. "
     "Graphiques 4.2 et 4.3 : la même, en Europe et aux États-Unis, sans les six enseignes.")
code("""SELECT region, industry,
       round(100.0 * sum(disparu)
             / count(*) FILTER (created_at BETWEEN DATE '2026-08-11' AND DATE '2026-08-24'), 1)
FROM s            -- ajouter WHERE NOT enseigne pour la seconde colonne
GROUP BY GROUPING SETS ((), (industry), (region), (region, industry))
ORDER BY ALL;""")
para("**c. Le sommaire, « 9,8 % contre 1,0 % »** : avis 4-5★ publiés du 11 au 17 août, toutes "
     "fiches, supprimés au plus tard au 7e jour : 9,84 % en services à domicile, 1,04 % en "
     "restauration.")
code("""SELECT industry, round(100 * avg(coalesce(j_suppression <= 7, FALSE)::INT), 2)
FROM s WHERE created_at BETWEEN DATE '2026-08-11' AND DATE '2026-08-17'
GROUP BY 1 ORDER BY 2 DESC;""")

titre("Fait 5. Les jours de pic", 3)
para("Nouveaux avis sans les six enseignes.")
tableau(["Publié", "Avis", "Supprimés", "%"], [
    ["Un jour de pic", "4 974", "115", "2,3"],
    ["Un autre jour", "17 404", "264", "1,5"],
], [5.6, 3.4, 3.4, 2.0])
para("À secteur et taille égaux (modèle du § 6.1 sans les six enseignes) : ×1,78 [1,13–2,80]. "
     "Par région (graphique A.4) : États-Unis ×1,33 [0,78–2,27], Europe ×2,75 [1,32–5,71].")
code("""SELECT pic, count(*), sum(supprime), round(100 * avg(supprime), 2)
FROM n WHERE NOT enseigne GROUP BY 1;""")

titre("Fait 6. États-Unis et Europe", 3)
para("Nouveaux avis supprimés au plus tard au 7e jour, pour 1 000 : sans les six enseignes, "
     "21,3 aux États-Unis contre 11,7 en Europe ; avec elles, 36,5 contre 11,7. Ce sont des "
     "écarts bruts, avant prise en compte du secteur et de la taille.")
code("""SELECT region, round(1000 * avg(supprime), 1),
       round(1000 * avg(supprime) FILTER (NOT enseigne), 1)
FROM n GROUP BY 1;""")

titre("Fait 7. Le texte", 3)
para("**Stock** : pendant le relevé, 0,101 % des avis 4-5★ avec texte ont disparu, contre "
     "0,053 % des avis sans texte. Par quintile de longueur, de 0,08 à 0,12 % (graphique 7.1).")
para("**Nouveaux avis sans les six enseignes**, pour 1 000 : plus de 80 caractères 19,7 (201 sur "
     "10 192) ; sans texte 13,7 (87 sur 6 335). En Europe : 17,2 (77 sur 4 465) contre 4,6 (16 "
     "sur 3 452).")
para("À secteur et taille égaux, la longueur du texte n'a pas d'effet propre : le vérifier en "
     "ajoutant `tranche_texte` (référence `sans`) au modèle du § 6.1. Le modèle ne contient pas "
     "la note de l'avis.")
code("""SELECT region, CASE WHEN longueur_texte = 0 THEN 'sans'
                    WHEN longueur_texte > 80 THEN 'plus de 80' ELSE '1 à 80' END AS texte,
       count(*), sum(supprime), round(1000 * avg(supprime), 1)
FROM n WHERE NOT enseigne
GROUP BY GROUPING SETS ((texte), (region, texte)) ORDER BY ALL;""")

titre("Fait 8. Les photos", 3)
para("×0,66 [0,40–1,08] : modèle du § 6.1, nouveaux avis, six enseignes comprises. L'écart est "
     "trop incertain pour être affirmé.")

titre("Fait 9. Les Local Guides", 3)
para("Stock : part des avis 4-5★ disparus pendant le relevé, selon le niveau de l'auteur.")
tableau(["Niveau", "1", "2", "3", "4", "5", "6", "7", "8", "Sans niveau"], [
    ["% disparus", "0,105", "0,117", "0,122", "0,093", "0,032", "0,014", "0,014", "0,020",
     "2,02"],
], [2.4, 1.4, 1.4, 1.4, 1.4, 1.4, 1.4, 1.4, 1.4, 2.2])
para("Le graphique 9.1 regroupe les niveaux 7 à 10 : les niveaux 8 à 10 ont trop peu de "
     "suppressions pour un taux à part. Il écarte les 12 651 avis sans niveau, dont la barre "
     "écraserait les autres.")
code("""SELECT local_guide_level, count(*), round(100 * avg(disparu), 3)
FROM s GROUP BY 1 ORDER BY 1;""")

titre("Annexes de la restitution", 3)
para("**905 fiches sur 8 997** ont perdu au moins un avis 4-5★ pendant le relevé ; 17 fiches "
     "portent le quart des 3 521 avis 4-5★ disparus. Toutes les fiches sont comptées, enseignes "
     "et fiches plus relevées comprises.")
code("""WITH f AS (
  SELECT cid, sum(disparu) FILTER (star >= 4) AS d
  FROM 'features_regression_reviews_doublons_cleaned_all.parquet' GROUP BY cid),
r AS (
  SELECT d, sum(d) OVER (ORDER BY d DESC ROWS UNBOUNDED PRECEDING) AS cumul,
         sum(d) OVER () AS total, row_number() OVER (ORDER BY d DESC) AS rang
  FROM f)
SELECT (SELECT count(*) FROM f), (SELECT count(*) FROM f WHERE d > 0), max(total),
       min(rang) FILTER (cumul >= 0.25 * total)
FROM r;""")
para("**Graphiques A.1 et A.2** : date de publication croisée avec la date de disparition. A.1, "
     "les quatre chaînes : avis 4-5★ publiés à partir du 22 juillet, les plus anciens regroupés. "
     "A.2, les deux boutiques The Boxer Club : avis 1★ publiés à partir du 31 juillet.")
para("**Les 601 avis disparus puis revenus** sont retirés de la base (Description Corpus, § 2, "
     "étape 3). 69 % sont revenus après un seul jour d'absence.")

# ---------------------------------------------------------------------------------------
titre("6. Les modèles", 2)
titre("6.1 Le modèle de la restitution", 3)
para("C'est le graphique 3.1 ; `regression_secteur_taille.csv` en donne tous les paramètres.")
tableau(["Élément", "Choix"], [
    ["Population", "Nouveaux avis (`features_03C_7_17_aout`)"],
    ["Variable expliquée", "`supprime`"],
    ["Variables affichées", "`prenom`, `photo`, `sans_reponse_j7`, `pic`, toutes en 0/1"],
    ["Contrôles, non affichés", "Secteur (référence : automobile), taille (référence : un "
     "établissement), région (référence : Europe) quand les deux régions sont ensemble"],
    ["Méthode", "Régression logistique, erreurs-types groupées par fiche (`cid`)"],
    ["Hors du modèle", "La note de l'avis"],
    ["Passages", "États-Unis, Europe, les deux"],
], [4.6, 12.0])
para("En formule : `supprime ~ prenom + photo + sans_reponse_j7 + pic + secteur + taille`, plus "
     "`region_US` quand les deux régions sont ensemble.")
tableau(["Passage", "Avis", "Supprimés", "Prénom", "Photo", "Sans réponse au 7e jour",
         "Jour de pic"], [
    ["États-Unis", "14 014", "512", "0,97 [0,73–1,30]", "0,51 [0,26–1,01]", "1,48 [0,95–2,30]",
     "1,86 [1,17–2,97]"],
    ["Europe", "10 177", "119", "1,78 [0,98–3,24]", "1,06 [0,56–2,00]", "2,77 [1,17–6,58]",
     "2,75 [1,32–5,71]"],
    ["Les deux", "24 191", "631", "1,05 [0,80–1,38]", "0,66 [0,40–1,08]", "1,65 [1,14–2,41]",
     "2,00 [1,35–2,95]"],
], [2.2, 1.7, 1.9, 2.7, 2.7, 2.7, 2.7])
para("Rapports de cotes, fourchette à 95 %. Un outil de régression logistique avec erreurs-types "
     "groupées par fiche retrouve ces valeurs à partir de `features_03C_7_17_aout` (vérifié le "
     "2026-10-12).")
para("Variantes utilisées dans la restitution : sans les six enseignes (fait 5, jour de pic "
     "×1,78) ; jalon au 2e jour (fait 2, graphique 2.1).")

titre("6.2 Les modèles jour par jour", 3)
para("Ils produisent le graphique 2.2, le graphique A.3 et le ×1,01 du prénom.")
puces([
    "Une ligne par avis et par jour où l'avis était en ligne et observé, du lendemain de sa "
    "première observation jusqu'à sa suppression ou jusqu'au 24 août. La variable expliquée vaut "
    "1 le jour de la suppression.",
    "Une réponse compte à partir du lendemain de son arrivée : `j_reponse <= jour - 1`.",
    "Graphique 2.2 : sans les six enseignes ; prénom, photo, pic, secteur, taille, région ; le "
    "statut de réponse en cinq modalités (pas encore de réponse, réponse le jour même, le "
    "lendemain, au 2e jour, au 3e jour ou après) ; le taux de réponse de la fiche, soit la part "
    "de ses avis 4-5★ publiés du 7 août 2025 au 6 août 2026 qui ont une réponse (0 %, 1 à 25 %, "
    "25 à 75 %, 75 % et plus ; les fiches de moins de 10 avis à part). Erreurs-types groupées "
    "par fiche.",
    "×1,01 du prénom : prénom, photo, sans réponse, pic, et un indicateur par fiche à la place "
    "du secteur, de la taille et de la région, sur les 223 fiches qui ont au moins une "
    "suppression.",
])
titre("Pourquoi un modèle jour par jour pour la réponse", 3)
para("Un avis supprimé à son 3e jour a eu moins de temps pour recevoir une réponse qu'un avis "
     "resté en ligne. « Répondu au plus tard au 7e jour » compte donc en partie les avis qui ont "
     "survécu. Dans le modèle jour par jour, la réponse n'est comptée qu'à partir du lendemain de "
     "son arrivée, ce qui compare les avis au même moment.")

# ---------------------------------------------------------------------------------------
titre("7. Pièges", 2)
para("**UTC.** Sans `SET TimeZone = 'UTC'`, DuckDB passe en heure de Paris et une partie des "
     "avis change de jour.")
para("**`supprime` et `disparu`.** `supprime` n'existe que pour les avis publiés du 7 au "
     "17 août. Sur le reste de la base, seul `disparu` se mesure.")
para("**Avis modifiés.** La ligne gardée d'un avis modifié porte le `first_seen_at` de la "
     "modification (Description Corpus, § 2) : pour ces avis, `j_premiere_vue` compte à partir de "
     "la version modifiée.")
para("**Passage du 12 août.** Il s'est terminé le 13 à 05:00 UTC : une date de `first_seen_at` "
     "au 13 août peut venir du 2e ou du 3e passage.")
para("**État au dernier passage.** Niveau Local Guide, réponse et nombre d'avis de l'auteur sont "
     "ceux du dernier passage où l'avis a été vu.")
para("**Annotation d'Axel.** Pour la suppression, prendre `supprime` ou `disparu` : la colonne "
     "`deleted` de `reviews_name_enriched` compte 388 fausses suppressions, dues aux 333 "
     "`review_id` en double.")

titre("8. Ce que les données ne montrent pas", 2)
puces([
    "La cause d'une suppression : filtre automatique, signalement du commerçant ou retrait par "
    "l'auteur.",
    "Les avis filtrés avant publication : aucun avis contenant un lien n'a disparu pendant le "
    "relevé, ils sont bloqués en amont.",
    "Les avis supprimés avant le 11 août, et ceux publiés puis supprimés entre deux passages.",
    "Une réponse retirée par le commerçant : l'avis est alors compté sans réponse.",
    "La cause des écarts mesurés : la réponse va avec moins de suppressions, et les données ne "
    "prouvent pas qu'elle les évite.",
])

titre("9. Pistes ouvertes", 2)
puces([
    "Le pic des 6e et 7e jours n'a pas d'explication.",
    "Les quatre chaînes antiparasitaires : leurs textes sont tous différents et on ne voit pas "
    "d'afflux soudain d'avis. Leurs avis citent très souvent un technicien par son prénom.",
    "Aucune recherche systématique d'autres attaques par faux avis n'a été faite.",
    "Le contenu des réponses et les signalements ne sont pas dans les données.",
    "L'effet combiné, par fiche, de la vitesse de collecte, de la part d'avis citant un prénom et "
    "de l'intensité de sollicitation reste à mesurer.",
])

zoom = doc.settings.element.find(qn("w:zoom"))
if zoom is not None and zoom.get(qn("w:percent")) is None:
    zoom.set(qn("w:percent"), "100")
doc.core_properties.title = "Passation : les chiffres de la restitution"
doc.core_properties.author = "Cartelis"
doc.save(SORTIE)
print(f"écrit dans {SORTIE}")
