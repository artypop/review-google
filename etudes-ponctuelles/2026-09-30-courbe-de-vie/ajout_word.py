#!/usr/bin/env python3
"""Ajoute les deux figures de survie à la fin de
livrables/restitution/regression_reviews_doublons_cleaned_180.docx.

    PYTHONPATH=.venv/Lib/site-packages <python uv> etudes-ponctuelles/2026-09-30-courbe-de-vie/ajout_word.py

Le document est complété sur place : il porte des ajouts faits à la main (la
section 6, recopiée d'analyses-complémentaires), qu'une régénération par le
07E effacerait. S'arrête si la section est déjà présente, ou si le document est
ouvert dans Word.
"""
from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

DOSSIER = Path(__file__).resolve().parent
RACINE = DOSSIER.parent.parent
SORTIES = DOSSIER / "sorties"
DOC = RACINE / "livrables" / "restitution" / "regression_reviews_doublons_cleaned_180.docx"
TITRE = "Annexe : survie des avis publiés pendant le relevé"
NAVY, FOND_FAMILLE = "322954", "EDEDF2"


def nombre(x) -> str:
    return f"{int(x):,}".replace(",", " ")


def virgule(x, d=1) -> str:
    return f"{float(x):.{d}f}".replace(".", ",")


def fond(cellule, couleur):
    ombre = OxmlElement("w:shd")
    ombre.set(qn("w:val"), "clear")
    ombre.set(qn("w:color"), "auto")
    ombre.set(qn("w:fill"), couleur)
    cellule._tc.get_or_add_tcPr().append(ombre)


def tableau(doc, entetes, lignes, largeurs):
    """Même mise en forme que les tableaux du 07E."""
    t = doc.add_table(rows=1 + len(lignes), cols=len(entetes))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    bords = OxmlElement("w:tblBorders")
    for cote in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{cote}")
        if cote in ("left", "right", "insideV"):
            e.set(qn("w:val"), "nil")
        else:
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), "4")
            e.set(qn("w:color"), "D9D8DE")
        bords.append(e)
    t._tbl.tblPr.append(bords)
    for i, rangee in enumerate([entetes] + lignes):
        for j, texte in enumerate(rangee):
            c = t.rows[i].cells[j]
            c.width = Cm(largeurs[j])
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.RIGHT
            r = p.add_run(str(texte))
            r.font.size = Pt(9)
            if i == 0:
                r.bold, r.font.color.rgb = True, RGBColor(0xFF, 0xFF, 0xFF)
                fond(c, NAVY)
    for j, colonne in enumerate(t._tbl.tblGrid.findall(qn("w:gridCol"))):
        colonne.set(qn("w:w"), str(int(largeurs[j] * 567)))


def puce(doc, morceaux):
    p = doc.add_paragraph(style="List Bullet")
    for texte, gras in ([(morceaux, False)] if isinstance(morceaux, str) else morceaux):
        p.add_run(texte).bold = gras
    p.paragraph_format.space_after = Pt(3)


def main() -> int:
    doc = Document(DOC)
    if any(p.text == TITRE for p in doc.paragraphs):
        raise SystemExit("La section est déjà dans le document.")
    resume = json.loads((SORTIES / "resume.json").read_text(encoding="utf-8"))
    jours = pd.read_csv(SORTIES / "survie_par_jour.csv")
    sans = jours[jours["jeu"] == "sans_enseignes"].set_index("jour")
    toutes = jours[jours["jeu"] == "toutes_fiches"].set_index("jour")
    dernier = int(toutes.index.max())

    saut = doc.add_paragraph()
    saut.add_run().add_break(WD_BREAK.PAGE)
    doc.add_paragraph(TITRE, style="Heading 1")
    doc.add_paragraph(
        f"Avis de {resume['corpus']} publiés du 11 au 24 août 2026 : "
        f"{nombre(resume['avis_publies'])} avis, {nombre(resume['suppressions'])} "
        "suppressions. Même règle de suppression que la régression. Les graphiques portent "
        "sur tous les avis, enseignes signalées comprises (décision du 2026-10-01).",
        style="NoteTableau")

    for figure in ("taux_survie.png", "suppressions_par_jour.png"):
        doc.add_picture(str(SORTIES / figure), width=Cm(16.6))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph("Ce que montrent les graphiques", style="Heading 2")
    puce(doc, [(f"{virgule(toutes.loc[dernier, 'survie_pct'])} % des avis sont encore en "
                f"ligne à leur {dernier}e jour", True),
               (f" : {virgule(100 - toutes.loc[dernier, 'survie_pct'])} % ont disparu. Sans "
                f"les 6 enseignes, {virgule(sans.loc[dernier, 'survie_pct'])} % sont encore "
                "en ligne.", False)])
    puce(doc, "Les suppressions se concentrent sur deux moments : les trois premiers jours "
              f"({virgule(toutes.loc[1, 'pour_10000'], 0)}, "
              f"{virgule(toutes.loc[2, 'pour_10000'], 0)} et "
              f"{virgule(toutes.loc[3, 'pour_10000'], 0)} pour 10 000 avis), puis les 6e et 7e "
              f"jours ({virgule(toutes.loc[6, 'pour_10000'], 0)} et "
              f"{virgule(toutes.loc[7, 'pour_10000'], 0)}). À partir du 9e jour, "
              f"{virgule(toutes.loc[9:, 'pour_10000'].max(), 0)} pour 10 000 au plus.")
    puce(doc, "Hors enseignes signalées, le pic du 7e jour se répète chaque jour du 18 au "
              "23 août, sur 13 à 29 fiches à chaque fois. Il tient à l'âge de l'avis : "
              "aucune purge datée ne l'explique.")

    # Le tableau ne tient pas sous les deux figures : il ouvre la page suivante.
    doc.add_paragraph("Les chiffres", style="Heading 2").paragraph_format.page_break_before = True
    lignes = [[str(k), nombre(toutes.loc[k, "avis_ayant_atteint_ce_jour"]),
               nombre(toutes.loc[k, "supprimes_ce_jour"]), virgule(toutes.loc[k, "pour_10000"]),
               f"{virgule(toutes.loc[k, 'survie_pct'], 2)} %"] for k in toutes.index]
    tableau(doc, ["Jour", "Avis ayant atteint ce jour", "Supprimés ce jour",
                  "Pour 10 000 avis", "Survie"],
            lignes, [2.0, 4.2, 3.6, 3.4, 3.4])
    doc.add_paragraph("Tous les avis.", style="NoteTableau")

    doc.add_paragraph("Comment ils sont construits", style="Heading 2")
    puce(doc, "Pour chaque jour d'existence, on divise les avis supprimés ce jour-là par les "
              "avis qui avaient atteint ce jour, encore en ligne et suivis par le robot. La "
              "survie au jour N enchaîne ces pertes du 1er au Ne jour : un avis publié le "
              "20 août ne compte que pour ses premiers jours.")
    puce(doc, f"{dernier} jours au plus : un avis publié le 11 août atteint son "
              f"{dernier + 1}e jour le 24, dernier relevé, où deux jours d'absence ne "
              "peuvent plus être constatés. Pour la même raison, les avis publiés les 23 et "
              f"24 août ne comptent pas ({nombre(resume['avis_publies'] - resume['avis_exposables'])}"
              " avis).")
    puce(doc, f"Au 1er jour, seuls {nombre(toutes.loc[1, 'avis_ayant_atteint_ce_jour'])} avis "
              "comptent : ceux relevés le jour même de leur publication. Les autres sont vus "
              "pour la première fois à leur 2e jour.")
    puce(doc, "L'axe du premier graphique commence à 90 %, pour que les écarts soient "
              "lisibles.")

    try:
        doc.save(DOC)
    except PermissionError:
        raise SystemExit(f"{DOC.name} est ouvert dans Word : le fermer, puis relancer.")
    print(f"écrit : {DOC}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
