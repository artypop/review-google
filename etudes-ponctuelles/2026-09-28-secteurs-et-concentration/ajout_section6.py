"""Ajoute la section 6 (tableaux 12 et 13) à livrables/analyses-complémentaires.docx.

    python etudes-ponctuelles/2026-09-28-secteurs-et-concentration/calcul.py --corpus reviews_doublons_cleaned_all
    python etudes-ponctuelles/2026-09-28-secteurs-et-concentration/ajout_section6.py

Demande de Matthieu du 2026-09-30. Copie la mise en forme des tableaux 1 et 2 ;
les chiffres viennent de sorties/doublons-cleaned-all/. S'arrête si le tableau
12 est déjà dans le document.
"""
import copy
import csv
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

RACINE = Path(__file__).resolve().parents[2]
DOC = RACINE / "livrables" / "analyses-complémentaires.docx"
SORTIES = RACINE / "etudes-ponctuelles" / "2026-09-28-secteurs-et-concentration" / "sorties"
CORPUS = "reviews_doublons_cleaned_all"
ORDRE = ["Automobile", "Services à domicile", "Santé", "Sport et bien-être",
         "Restauration", "Voyage", "Hôtellerie"]
NBSP = "\u00a0"


def espace(n):
    return f"{int(n):,}".replace(",", NBSP)


def virgule(x, d):
    return f"{float(x):.{d}f}".replace(".", ",")


def lire(chemin):
    with open(chemin, encoding="utf-8") as f:
        return {r["secteur_fr"]: r for r in csv.DictReader(f)}


def lignes(d):
    out = []
    for s in ORDRE:
        r = d[s]
        out.append([s, espace(r["avis"]), f"{virgule(r['part_corpus_pct'], 1)} %",
                    espace(r["suppressions"]),
                    f"{virgule(float(r['part_supprimee_pct']), 2)} %"])
    avis = sum(int(r["avis"]) for r in d.values())
    suppr = sum(int(r["suppressions"]) for r in d.values())
    out.append(["Total", espace(avis), "100 %", espace(suppr),
                f"{virgule(100 * suppr / avis, 2)} %"])
    return out, avis, suppr


def texte_paragraphe(el, texte):
    """Remplace le texte d'un paragraphe en gardant la mise en forme du premier run."""
    # Signets copiés avec le paragraphe : ils existeraient en double.
    for tag in ("w:bookmarkStart", "w:bookmarkEnd"):
        for b in el.findall(qn(tag)):
            el.remove(b)
    runs = el.findall(qn("w:r"))
    for r in runs[1:]:
        el.remove(r)
    for t in runs[0].findall(qn("w:t"))[1:]:
        runs[0].remove(t)
    t = runs[0].find(qn("w:t"))
    t.text = texte
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")


def remplir_tableau(el, doc, valeurs):
    table = Table(el, doc)
    assert len(table.rows) == len(valeurs) + 1
    for rangee, vals in zip(table.rows[1:], valeurs):
        for cellule, v in zip(rangee.cells, vals):
            texte_paragraphe(cellule.paragraphs[0]._p, v)


def main():
    doc = Document(DOC)
    body = doc.element.body
    if any("Tableau 12." in p.text for p in doc.paragraphs):
        raise SystemExit("La section 6 est déjà présente.")
    els = list(body.iterchildren())

    def trouver(debut):
        for el in els:
            if el.tag == qn("w:p") and Paragraph(el, doc).text.startswith(debut):
                return el
        raise SystemExit(f"introuvable : {debut}")

    sommaire_5 = trouver("5. Textes des avis américains")
    titre_5 = [el for el in els if el.tag == qn("w:p")
               and Paragraph(el, doc).style.name == "Heading 1"
               and Paragraph(el, doc).text.startswith("5.")][0]
    intro_1 = trouver("Toute la base, avis de reviews_doublons_cleaned")
    titre_t1, titre_t2 = trouver("Tableau 1. "), trouver("Tableau 2. ")
    note_t1 = trouver("% de suppression : ")
    note_t2 = trouver("Enseignes signalées : ")
    t1, t2 = doc.tables[0]._tbl, doc.tables[1]._tbl

    tous, n_avis, n_suppr = lignes(lire(SORTIES / "doublons-cleaned-all" / "A1-secteurs.csv"))
    sans, _, _ = lignes(lire(SORTIES / "doublons-cleaned-all" / "A1-secteurs-sans-enseignes.csv"))
    ref = lire(SORTIES / "doublons-cleaned" / "A1-secteurs.csv")
    ref_avis = sum(int(r["avis"]) for r in ref.values())
    ref_suppr = sum(int(r["suppressions"]) for r in ref.values())
    ecart_dom = (int(lire(SORTIES / "doublons-cleaned-all" / "A1-secteurs.csv")
                     ["Services à domicile"]["suppressions"])
                 - int(ref["Services à domicile"]["suppressions"]))

    # Sommaire
    s6 = copy.deepcopy(sommaire_5)
    texte_paragraphe(s6, f"6. Suppressions par secteur, avis de {CORPUS}")
    sommaire_5.addnext(s6)

    # Section 6, en fin de document
    nouveaux = []
    h = copy.deepcopy(titre_5)
    texte_paragraphe(h, f"6. Suppressions par secteur, avis de {CORPUS}")
    nouveaux.append(h)
    intro = copy.deepcopy(intro_1)
    texte_paragraphe(intro,
        f"Toute la base, avis de {CORPUS} : {espace(n_avis)} avis, toutes dates de "
        f"publication, {espace(n_suppr)} suppressions. Cette table reprend les avis de "
        f"reviews_doublons_cleaned (section 1) et en ajoute {espace(n_avis - ref_avis)}, "
        f"qui portent {espace(n_suppr - ref_suppr)} suppressions de plus, dont "
        f"{espace(ecart_dom)} dans les services à domicile.")
    nouveaux.append(intro)
    for titre, tbl, note, valeurs, libelle, texte_note in [
        (titre_t1, t1, note_t1, tous,
         f"Tableau 12. Avis et suppressions par secteur, {CORPUS}",
         "% de suppression : suppressions du secteur divisées par ses avis. Corpus : les "
         f"avis présents dans {CORPUS}. Suppression : même définition que les tableaux 1 "
         "et 2."),
        (titre_t2, t2, note_t2, sans,
         "Tableau 13. Même tableau, sans les six enseignes signalées", None),
    ]:
        p = copy.deepcopy(titre)
        texte_paragraphe(p, libelle)
        t = copy.deepcopy(tbl)
        remplir_tableau(t, doc, valeurs)
        n = copy.deepcopy(note)
        if texte_note:
            texte_paragraphe(n, texte_note)
        nouveaux += [p, t, n]

    fin = body.sectPr
    for el in nouveaux:
        fin.addprevious(el)
    doc.save(DOC)
    print(f"écrit : {DOC}")
    for r in tous + [["--"]] + sans:
        print("  ", r)


if __name__ == "__main__":
    main()
