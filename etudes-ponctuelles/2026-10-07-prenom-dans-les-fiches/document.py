#!/usr/bin/env python3
"""Word de synthèse pour la réponse au mail d'Axel (prénoms et vitesse de collecte).

    PYTHONPATH=".venv/Lib/site-packages;outils" <python uv> etudes-ponctuelles/2026-10-07-prenom-dans-les-fiches/document.py [nom.docx]

Demande de Matthieu du 2026-10-07 : vérifier sa réponse au mail d'Axel et la synthétiser
en 1 à 3 pages. Version 3, à sa demande : moins dense, sans renvois [A]-[D] ; chaque
résultat dit de quelles données il part. Les chiffres sont recopiés des sorties listées
en fin de document.

Refuse d'écrire si le fichier existe déjà (il a pu être retouché à la main).
"""
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

RACINE = Path(__file__).resolve().parents[2]
FINALE = RACINE / "livrables" / "restitution_finale"
SORTIE = FINALE / "2026-10-07-prenoms-vitesse-verification_v3.docx"
GRAPHIQUES = FINALE / "graphiques" / "faits-saignants" / "graphiques"
POLICE = "Roboto"
ENCRE, ENCRE_2, BLEU = RGBColor(0x0B, 0x0B, 0x0B), RGBColor(0x52, 0x51, 0x4E), RGBColor(0x2A, 0x78, 0xD6)
FOND_ENTETE = "E8F1FB"


def police(run, taille=None, gras=None, couleur=None, italique=None):
    run.font.name = POLICE
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), POLICE)
    if taille:
        run.font.size = Pt(taille)
    if gras is not None:
        run.bold = gras
    if italique is not None:
        run.italic = italique
    if couleur is not None:
        run.font.color.rgb = couleur
    return run


def paragraphe(doc, morceaux, taille=10.5, apres=6, style=None, couleur=ENCRE, avant=0):
    """`morceaux` : texte, ou liste de (texte, gras)."""
    p = doc.add_paragraph(style=style)
    if isinstance(morceaux, str):
        morceaux = [(morceaux, False)]
    for texte, gras in morceaux:
        police(p.add_run(texte), taille, gras, couleur)
    f = p.paragraph_format
    f.space_after, f.space_before, f.line_spacing = Pt(apres), Pt(avant), 1.15
    return p


def titre(doc, texte, niveau=1):
    h = doc.add_heading(level=niveau)
    police(h.add_run(texte), 13.5 if niveau == 1 else 11.5, True, ENCRE if niveau == 1 else BLEU)
    h.paragraph_format.space_before = Pt(16 if niveau == 1 else 12)
    h.paragraph_format.space_after = Pt(5)
    h.paragraph_format.keep_with_next = True
    return h


def puce(doc, texte):
    return paragraphe(doc, texte, apres=4, style="List Bullet")


def donnees(doc, texte):
    """La ligne qui dit de quels avis part le résultat."""
    return paragraphe(doc, [("Données : ", True), (texte, False)], taille=9, apres=4,
                      couleur=ENCRE_2, avant=2)


def ombrer(cellule, couleur):
    tc = cellule._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), couleur)
    tc.append(shd)


def tableau(doc, entetes, lignes, largeurs_cm, taille=9.5):
    t = doc.add_table(rows=1 + len(lignes), cols=len(entetes))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, ligne in enumerate([entetes] + lignes):
        for j, contenu in enumerate(ligne):
            c = t.cell(i, j)
            c.width = Cm(largeurs_cm[j])
            c.paragraphs[0].text = ""
            police(c.paragraphs[0].add_run(contenu), taille, i == 0, ENCRE)
            f = c.paragraphs[0].paragraph_format
            f.space_after, f.space_before = Pt(3), Pt(3)
            if i == 0:
                ombrer(c, FOND_ENTETE)
    return t


def image(doc, chemin, legende, largeur_cm=15.5):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(chemin), width=Cm(largeur_cm))
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    lg = paragraphe(doc, legende, taille=8.5, apres=10, couleur=ENCRE_2)
    lg.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in lg.runs:
        r.italic = True


def main() -> int:
    global SORTIE
    if len(sys.argv) > 1:
        SORTIE = FINALE / sys.argv[1]
    if SORTIE.exists():
        raise SystemExit(f"{SORTIE} existe déjà : il a pu être retouché. Choisir un autre nom.")
    doc = Document()
    for s in doc.sections:
        s.page_height, s.page_width = Cm(29.7), Cm(21)
        s.top_margin = s.bottom_margin = Cm(2)
        s.left_margin = s.right_margin = Cm(2.2)
    normal = doc.styles["Normal"]
    normal.font.name = POLICE
    normal.font.size = Pt(10.5)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), POLICE)

    t = doc.add_paragraph()
    police(t.add_run("Prénoms, vitesse de collecte et suppressions"), 18, True, ENCRE)
    t.paragraph_format.space_after = Pt(2)
    paragraphe(doc, "Vérification de la réponse au mail d'Axel – 7 octobre 2026", 10,
               apres=10, couleur=ENCRE_2)
    paragraphe(doc, [("Question d'Axel : ", True),
                     ("« Est-ce qu'on a testé cette idée que la vélocité pure a un impact sur "
                      "le % de suppressions ? »", False)])

    # 1 -------------------------------------------------------------------------
    titre(doc, "1. Un avis qui cite un prénom n'est pas plus supprimé que les autres avis "
               "de sa fiche", 2)
    puce(doc, "Entre avis d'une même fiche, citer un prénom multiplie le risque de suppression "
              "par 1,01. Sans les six enseignes, par 0,96.")
    donnees(doc, "avis 4-5 étoiles publiés du 7 au 17 août 2026 (corpus 03C, 24 191 avis), sur "
                 "les fiches qui ont perdu au moins un avis. Le 1,01 compte les suppressions "
                 "jusqu'au 24 août, enseignes comprises ; le 0,96 jusqu'au 7e jour de l'avis.")
    puce(doc, "Sur les avis de tous âges, à fiche et âge égaux, 759 avis supprimés citent un "
              "prénom quand le hasard en donnerait 721 : un écart de 5 %.")
    donnees(doc, "avis 4-5 étoiles avec texte annotés par Axel, de tous âges, décrits en partie 2.")

    # 2 -------------------------------------------------------------------------
    titre(doc, "2. Les fiches qui citent beaucoup de prénoms perdent 5,7 fois plus d'avis", 2)
    puce(doc, "Quand 50 % et plus des avis d'une fiche citent un prénom, elle perd 148 avis pour "
              "100 000 en deux semaines, contre 26 quand c'est moins de 10 %.")
    puce(doc, "À secteur, taille et région égaux, l'écart reste d'un facteur 3,3.")
    donnees(doc, "avis 4-5 étoiles avec texte annotés par Axel, de tous âges, sans les six "
                 "enseignes, sur les 8 324 fiches qui ont au moins 20 avis annotés : 2 444 287 "
                 "avis, dont 1 676 supprimés du 11 au 24 août 2026. La part de prénoms d'une "
                 "fiche est calculée sur ses avis publiés avant le 7 août.")

    # 3 -------------------------------------------------------------------------
    titre(doc, "3. À vitesse égale, l'écart demeure", 2)
    puce(doc, "Les fiches sont coupées en quatre groupes de même taille selon leur vitesse. "
              "Dans chaque groupe, les fiches à 50 % et plus de prénoms perdent 2,5 à 6,5 fois "
              "plus d'avis que celles à moins de 10 %.")
    puce(doc, "Les fiches à moins de 10 % de prénoms perdent autant d'avis quelle que soit leur "
              "vitesse : de 0,22 à 0,27 pour 1 000.")
    donnees(doc, "les mêmes avis. La vitesse est le nombre d'avis reçus par mois, toutes notes, "
                 "du 7 août 2025 au 6 août 2026.")
    image(doc, GRAPHIQUES / "velocite_prenom_taux.png",
          "Graphique 1 – Avis supprimés pour 1 000, par groupe de vitesse et part de prénoms de "
          "la fiche (avis annotés par Axel, tous âges).")

    # 4 -------------------------------------------------------------------------
    titre(doc, "4. La vitesse seule n'explique presque rien", 2)
    puce(doc, "Quand la vitesse de collecte double, le risque de suppression monte de 28 %.")
    puce(doc, "À âge des avis égal, il ne monte plus que de 12 %.")
    puce(doc, "Sans les 84 fiches qui ont perdu le plus d'avis (1 % des fiches, 62 % des "
              "suppressions), il baisse de 8 %.")
    puce(doc, "Sur des avis neufs, tous du même âge : +8 %, trop incertain pour être affirmé.")
    donnees(doc, "les mêmes avis, comparés à secteur, taille et région égaux. Pour la dernière "
                 "ligne : avis 4-5 étoiles du corpus 03C, sans les six enseignes (22 378 avis, "
                 "560 suppressions).")

    # 5 -------------------------------------------------------------------------
    titre(doc, "5. Sans les fiches les plus touchées, les prénoms gardent un effet : +54 %", 2)
    puce(doc, "À âge, secteur, taille, région et vitesse égaux, les fiches à 50 % et plus de "
              "prénoms ont un risque 54 % plus élevé que celles à moins de 10 % (fourchette : "
              "de +14 % à +110 %).")
    donnees(doc, "les mêmes avis, sans les 84 fiches qui ont perdu le plus d'avis.")
    paragraphe(doc, [("Limites", True)], apres=3, avant=4)
    puce(doc, "Avec toutes les fiches, cet écart (+69 %) est trop incertain pour être affirmé.")
    puce(doc, "Les 84 fiches retirées sont surtout riches en prénoms : 43 sont à 50 % et plus.")
    image(doc, GRAPHIQUES / "velocite_prenom_modeles.png",
          "Graphique 2 – Écart de risque des fiches à 50 % et plus de prénoms, et effet d'un "
          "doublement de la vitesse, selon ce qui est tenu égal. Barres 1 à 4 : avis annotés "
          "par Axel ; barre 5 : avis neufs du corpus 03C. La barre 1 est déjà à secteur, "
          "taille et région égaux.")

    # 6 -------------------------------------------------------------------------
    titre(doc, "6. La règle de Google rejoint l'hypothèse de la sollicitation", 2)
    puce(doc, "Google autorise les commerçants à demander des avis.")
    puce(doc, "Depuis avril 2026, il interdit de demander au personnel de solliciter des avis "
              "qui citent un membre du personnel. Une fiche qui fait citer ses techniciens "
              "tombe sous cette règle.")
    puce(doc, "Nos données ne disent pas comment les avis ont été demandés.")
    donnees(doc, "règles Google Maps sur les contributions, « Prohibited and restricted "
                 "content » (support.google.com/contributionpolicy/answer/7400114) ; ajout "
                 "signalé le 17 avril 2026.")

    # 7 -------------------------------------------------------------------------
    titre(doc, "Ce qu'il faut corriger dans le mail", 1)
    tableau(doc, ["Dans le mail", "À écrire"], [
        ["« plus de 50 % … 2,5 à 6,4 fois plus »", "« 50 % et plus … 2,5 à 6,5 fois plus »"],
        ["« ce lien disparaît »", "« ce lien s'inverse légèrement : −8 % quand la vitesse double »"],
        ["« Rappel : il est interdit de solliciter un avis selon les règles Google Maps »",
         "« Google interdit les contreparties, la sollicitation des seuls clients satisfaits et, "
         "depuis avril 2026, de faire citer un employé dans l'avis »"],
    ], [7.0, 9.6])
    paragraphe(doc, "Les autres chiffres du mail sont justes : ×1,01 au niveau de l'avis, 0,148 % "
                    "contre 0,026 %, 84 fiches sur 8 324 et 62 % des suppressions, +54 %.",
               apres=6, avant=8)

    paragraphe(doc, [("Pour refaire les calculs. ", True),
                     ("Avis annotés par Axel : etudes-ponctuelles/2026-10-06-velocite-prenom/"
                      "calcul.py et, pour le 759 contre 721, 2026-10-07-prenom-dans-les-fiches/"
                      "calcul.py. Corpus 03C : logistic-regression-study/python/"
                      "07G_regression_03C_jour_par_jour.py (×1,01) et 2026-10-07-prenom-dans-les-"
                      "fiches/meme_fiche_03C.py (×0,96).", False)],
               taille=8, apres=0, couleur=ENCRE_2, avant=10)

    SORTIE.parent.mkdir(parents=True, exist_ok=True)
    doc.save(SORTIE)
    print(f"écrit : {SORTIE}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
