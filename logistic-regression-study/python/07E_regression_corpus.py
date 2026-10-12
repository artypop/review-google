#!/usr/bin/env python3
"""
==============================================================================
Script : 07E_regression_corpus.py

Régression logistique sur un corpus d'avis quelconque, en six passages. Écrit le
2026-09-30 à la demande de Matthieu, à partir des caractéristiques du 07B.

    PYTHONPATH=.venv/Lib/site-packages <python uv> logistic-regression-study/python/07E_regression_corpus.py

POUR CHANGER DE CORPUS : modifier `CORPUS`, juste sous les imports, et relancer.
Le script reconstruit alors les caractéristiques (`03C_features_corpus.py`)
puis refait les six passages. Il les reconstruit aussi quand le corpus,
`reviews`, `businesses` ou le 03C sont plus récents que la table.

------------------------------------------------------------------------------
LE MODÈLE
------------------------------------------------------------------------------
Huit caractéristiques, affichées :
  note                  1, 2, 3, 4 étoiles                référence : 5 étoiles
  longueur du texte     1–50, 51–200, 201+ caractères     référence : sans texte
  photo jointe          oui / non
  photos de l'auteur    log(1 + photos publiées)
  palier Local Guide    sans niveau, 5 et plus            référence : 1 à 4
  afflux sur la fiche   log(1 + avis du jour ÷ rythme habituel)
  secteur               sous 1 % des avis du passage : « autres »
                                                          référence : automobile
  taille de l'entreprise                                  référence : mono

Trois contrôles, dans le modèle et jamais affichés dans les tableaux ni les
figures client (décision de Matthieu du 2026-09-30) :
  log_age_premiere_observation   log(1 + jours entre publication et premier
                                 relevé). Sur ce corpus, le risque est divisé
                                 par 40 entre un avis d'une semaine et un avis
                                 de 3 à 6 mois : sans lui, tout ce qui est plus
                                 fréquent chez les avis récents paraîtrait
                                 dangereux.
  fenetre_observation_j          jours où une suppression pouvait commencer et
                                 être confirmée (voir le 03C).
  region_US                      passages « tous » seulement. L'écart
                                 États-Unis / Europe dépend du choix des fiches
                                 et ne se montre pas ; sans lui, il se
                                 reporterait sur le secteur.

Retirés du 07B le 2026-09-30 : log_rc, langue_minoritaire_sur_la_fiche,
vu_tardivement, log_burst. Les deux contrôles d'âge du 07B sont remplacés par
ceux ci-dessus.

Suppression : absent deux jours de suite ou plus (voir le 03C). Les avis dont
la fenêtre d'observation est nulle ne peuvent pas être supprimés au sens de
cette règle : ils sont écartés, et leur nombre est écrit dans les sorties.

------------------------------------------------------------------------------
LES SIX PASSAGES
------------------------------------------------------------------------------
US, Europe, tous ; chacun sur toutes les fiches et sans les 6 enseignes
signalées (4 chaînes antiparasitaires américaines repérées sur le début du
nom, 2 salles de sport espagnoles repérées sur leur `cid`).

Repris du 07B : marges d'incertitude groupées par fiche, garde-fous sur les
colonnes trop rares, qualité mesurée en cinq tours par établissement en
écartant du test les auteurs déjà vus.

------------------------------------------------------------------------------
SORTIES : output-study/<date>-sorties-07E-<nom du corpus>/
------------------------------------------------------------------------------
Pour le client, sans les contrôles :
  07E_tableau_client.csv       risque relatif et fourchette, une colonne par
                               passage
  07E_risques_relatifs.csv     la même chose en format long, pour les notes
  07E_effectifs.csv            avis, suppressions, taux pour 10 000 avis, par
                               valeur de chaque caractéristique et par passage
  07E_figure_<région>.png      une figure par région, avec et sans enseignes
Le document client, reconstruit à chaque passage :
  livrables/restitution/regression_<nom du corpus>.docx
      introduction, règles de lecture, figure et tableau par région, facteurs
      retenus selon la règle de `facteurs()`, limites, effectifs en annexe
Dans technique/ :
  07E_coefficients_<passage>.csv, 07E_summary_<passage>.txt   contrôles compris
  07E_passages.csv             effectifs, avis écartés, qualité de classement
  07E_source.txt               corpus lu, date, moteur
==============================================================================
"""
from __future__ import annotations

import importlib.util
import sys
from datetime import date, datetime
from pathlib import Path

import numpy as np
import pandas as pd

RACINE = Path(__file__).resolve().parents[2]

# =============================================================================
# LE SEUL RÉGLAGE : le corpus d'avis étudié, au schéma de reviews_doublons_cleaned.
CORPUS = RACINE / "data" / "local" / "reviews_doublons_cleaned_180.parquet"
# =============================================================================

try:
    import statsmodels.api as sm
    MOTEUR = "statsmodels"
except ImportError:
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    import glm_numpy
    glm_numpy.installer()
    import statsmodels.api as sm
    MOTEUR = "glm_numpy (statsmodels indisponible sur ce poste)"

import matplotlib  # noqa: E402

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

SORTIES = (Path(__file__).resolve().parent.parent / "output-study"
           / f"{date.today():%Y-%m-%d}-sorties-07E-{CORPUS.stem}")
TECHNIQUE = SORTIES / "technique"
RESTITUTION = RACINE / "livrables" / "restitution"

N_TOURS = 5
GRAINE = 20260914
SEUIL_SECTEUR_RARE = 0.01
MIN_CAS_PAR_COLONNE = 30
MIN_SUPPRESSIONS_PAR_COLONNE = 5

TRANCHES_TEXTE = [-1, 0, 50, 200, np.inf]
NOMS_TRANCHES_TEXTE = ["texte_sans_texte", "texte_1_50", "texte_51_200", "texte_201_plus"]

# Tranches des deux caractéristiques continues, pour les effectifs seulement.
TRANCHES_PHOTOS = ["0 photo", "1 à 10 photos", "11 à 100 photos", "plus de 100 photos"]
TRANCHES_AFFLUX = ["jusqu'à 1 fois le rythme habituel", "1 à 3 fois", "3 à 10 fois",
                   "plus de 10 fois"]

REFERENCES = {
    "etoiles": "etoiles_5",
    "texte": "texte_sans_texte",
    "guide": "guide_1_4",
    "secteur": "secteur_automotive",
    "taille": "taille_mono",
}
CONTINUES = ["has_photo", "log_photos_auteur", "log_ratio_pic_journalier_fiche"]
CONTROLES = ["log_age_premiere_observation", "fenetre_observation_j", "region_US"]

# Libellés du rendu client, dans l'ordre d'affichage : (variable, famille, libellé).
LIBELLES = [
    ("etoiles_1", "Note", "1 étoile"),
    ("etoiles_2", "Note", "2 étoiles"),
    ("etoiles_3", "Note", "3 étoiles"),
    ("etoiles_4", "Note", "4 étoiles"),
    ("etoiles_5", "Note", "5 étoiles (référence)"),
    ("texte_sans_texte", "Longueur du texte", "sans texte (référence)"),
    ("texte_1_50", "Longueur du texte", "1 à 50 caractères"),
    ("texte_51_200", "Longueur du texte", "51 à 200 caractères"),
    ("texte_201_plus", "Longueur du texte", "plus de 200 caractères"),
    ("has_photo", "Photo jointe", "photo jointe à l'avis (contre aucune)"),
    ("log_photos_auteur", "Photos de l'auteur",
     "photos publiées par l'auteur, par multiplication par 2,7"),
    ("guide_sans_niveau", "Profil Local Guide", "compte sans niveau"),
    ("guide_1_4", "Profil Local Guide", "niveau 1 à 4 (référence)"),
    ("guide_5_et_plus", "Profil Local Guide", "niveau 5 et plus"),
    ("log_ratio_pic_journalier_fiche", "Afflux sur la fiche",
     "avis du jour sur la fiche, par multiplication par 2,7"),
    ("secteur_automotive", "Secteur", "automobile (référence)"),
    ("secteur_food_beverage", "Secteur", "restauration"),
    ("secteur_healthcare", "Secteur", "santé"),
    ("secteur_home_services", "Secteur", "services à domicile"),
    ("secteur_hospitality", "Secteur", "hôtellerie"),
    ("secteur_travel", "Secteur", "voyage"),
    ("secteur_wellness_fitness", "Secteur", "sport et bien-être"),
    ("secteur_autres", "Secteur", "autres secteurs"),
    ("taille_mono", "Taille de l'entreprise", "un seul établissement (référence)"),
    ("taille_small", "Taille de l'entreprise", "groupe de 4 à 10 établissements"),
    ("taille_large", "Taille de l'entreprise", "groupe de 20 à 50 établissements"),
]
REFERENCE_DE_FAMILLE = {
    "Note": "5 étoiles", "Longueur du texte": "sans texte",
    "Photo jointe": "sans photo", "Profil Local Guide": "niveau 1 à 4",
    "Secteur": "automobile", "Taille de l'entreprise": "un seul établissement",
}
ORDRE = [v for v, _, _ in LIBELLES]
FAMILLE = {v: f for v, f, _ in LIBELLES}
LIBELLE = {v: l for v, _, l in LIBELLES}

PASSAGES = [(region, jeu) for region in ["US", "Europe", "tous"]
            for jeu in ["toutes_fiches", "sans_enseignes"]]
NOM_REGION = {"US": "États-Unis", "Europe": "Europe", "tous": "États-Unis et Europe"}
NOM_JEU = {"toutes_fiches": "toutes les fiches", "sans_enseignes": "sans les 6 enseignes"}

# Palette de référence du skill dataviz, emplacements 1 et 2 (validés en clair).
COULEUR_JEU = {"sans_enseignes": "#2a78d6", "toutes_fiches": "#eb6834"}
ENCRE, ENCRE_2, GRILLE, FOND = "#0b0b0b", "#52514e", "#e4e3df", "#fcfcfb"


# ---------------------------------------------------------------------------
# Caractéristiques
# ---------------------------------------------------------------------------

def charger_03c():
    chemin = Path(__file__).resolve().parent / "03C_features_corpus.py"
    spec = importlib.util.spec_from_file_location("features_corpus", chemin)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def features_a_jour(f03c) -> Path:
    """La table de caractéristiques du corpus, reconstruite si elle manque ou si
    une de ses sources est plus récente qu'elle."""
    if not CORPUS.exists():
        raise SystemExit(f"Corpus introuvable : {CORPUS}")
    table = f03c.chemin_features(CORPUS)
    if table.exists() and all(table.stat().st_mtime >= s.stat().st_mtime
                              for s in f03c.sources(CORPUS)):
        print(f"[features] à jour : {table.name}")
        return table
    print("[features] construction")
    return f03c.construire(CORPUS)


def preparer(df: pd.DataFrame) -> pd.DataFrame:
    if df["star"].isna().any():
        raise SystemExit(f"{int(df['star'].isna().sum())} avis sans note.")
    df = df.copy()
    # Fenêtre nulle ou négative : l'avis ne peut pas être supprimé au sens de la
    # règle. Écarté de chaque passage, et compté.
    df["exposable"] = df["fenetre_observation_j"] > 0

    df["supprime"] = df["supprime"].astype(int)
    df["cid"] = df["cid"].astype(str)
    df["author_key"] = df["author_key"].astype(str)
    df["etoiles"] = "etoiles_" + df["star"].astype(int).astype(str)
    df["texte"] = pd.cut(df["text_chars"], TRANCHES_TEXTE,
                         labels=NOMS_TRANCHES_TEXTE).astype(str)
    df["guide"] = "guide_" + df["palier_local_guide"].astype(str)
    df["taille"] = "taille_" + df["bucket"].fillna("inconnue").astype(str)
    df["industry"] = df["industry"].fillna("inconnu").astype(str)
    df["has_photo"] = df["has_photo"].astype(int)
    df["log_photos_auteur"] = np.log1p(df["reviewer_photo_count"].clip(lower=0))
    df["log_ratio_pic_journalier_fiche"] = df["log_ratio_pic_journalier_fiche"].fillna(0)
    df["log_age_premiere_observation"] = np.log1p(df["age_a_la_premiere_observation_j"])
    df["fenetre_observation_j"] = df["fenetre_observation_j"].astype(float)
    df["region_US"] = (df["region"] == "US").astype(int)
    df["enseigne"] = df["chaine_antiparasitaire"] | df["salle_espagnole"]

    # Tranches des deux caractéristiques continues, pour les effectifs seulement.
    df["tranche_photos_auteur"] = pd.cut(
        df["reviewer_photo_count"], [-1, 0, 10, 100, np.inf],
        labels=TRANCHES_PHOTOS).astype(str)
    ratio = np.expm1(df["log_ratio_pic_journalier_fiche"])
    df["tranche_afflux"] = pd.cut(
        ratio, [-np.inf, 1, 3, 10, np.inf], labels=TRANCHES_AFFLUX).astype(str)
    return df


def selection(df: pd.DataFrame, region: str, jeu: str) -> pd.Series:
    """Les avis d'un passage, avant l'écart des fenêtres nulles."""
    garde = pd.Series(True, index=df.index)
    if region != "tous":
        garde &= df["region"] == region
    if jeu == "sans_enseignes":
        garde &= ~df["enseigne"]
    return garde


def secteurs(df: pd.DataFrame) -> pd.Series:
    """Règle du 07B, appliquée à chaque passage."""
    parts = df["industry"].value_counts(normalize=True)
    rares = parts[parts < SEUIL_SECTEUR_RARE].index
    return "secteur_" + df["industry"].where(~df["industry"].isin(rares), "autres")


def matrice(df: pd.DataFrame, avec_region: bool) -> pd.DataFrame:
    morceaux = []
    for famille, reference in REFERENCES.items():
        d = pd.get_dummies(df[famille], dtype=float).drop(columns=[reference],
                                                          errors="ignore")
        morceaux.append(d[sorted(d.columns, key=lambda c: ORDRE.index(c)
                                 if c in ORDRE else len(ORDRE))])
    controles = [c for c in CONTROLES if avec_region or c != "region_US"]
    X = pd.concat(morceaux + [df[CONTINUES].astype(float), df[controles].astype(float)],
                  axis=1)
    return sm.add_constant(X, has_constant="add")


def ecarter_colonnes_degenerees(X: pd.DataFrame, y: pd.Series) -> pd.DataFrame:
    """Garde-fou du 07B : une colonne oui/non trop rare, ou dont toutes les
    lignes à 1 ont le même sort, fait diverger l'ajustement."""
    a_jeter = []
    for col in X.columns:
        if col == "const" or X[col].nunique() > 2:
            continue
        masque = X[col] > 0
        n_suppr = int(y[masque].sum())
        if (masque.sum() < MIN_CAS_PAR_COLONNE or n_suppr < MIN_SUPPRESSIONS_PAR_COLONNE
                or n_suppr == int(masque.sum())):
            a_jeter.append(col)
    if a_jeter:
        print(f"  colonnes écartées, trop rares ou sans variation : {', '.join(a_jeter)}")
    return X.drop(columns=a_jeter)


# ---------------------------------------------------------------------------
# Ajustement et qualité
# ---------------------------------------------------------------------------

def ajuster(df: pd.DataFrame, avec_region: bool):
    y = df["supprime"].astype(float)
    X = ecarter_colonnes_degenerees(matrice(df, avec_region), y)
    res = sm.GLM(y, X, family=sm.families.Binomial()).fit(
        cov_type="cluster", cov_kwds={"groups": df["cid"].factorize()[0]})
    return res, X.columns


def aire_sous_courbe(y, p) -> float:
    """Probabilité qu'un avis supprimé reçoive un score plus élevé qu'un avis
    resté en ligne, tiré au hasard."""
    y = np.asarray(y)
    n_pos, n_neg = int(y.sum()), int((1 - y).sum())
    if n_pos == 0 or n_neg == 0:
        return float("nan")
    rangs = pd.Series(np.asarray(p, dtype="float64")).rank(method="average").to_numpy()
    return (rangs[y == 1].sum() - n_pos * (n_pos + 1) / 2) / (n_pos * n_neg)


def cinq_tours(df: pd.DataFrame, avec_region: bool) -> float:
    """Qualité de classement sur des établissements jamais vus, comme le 07B :
    fiches réparties en cinq groupes, chacun testé une fois par un modèle
    ajusté sur les quatre autres, auteurs déjà vus écartés du test."""
    tirage = np.random.default_rng(GRAINE)
    fiches = np.sort(df["cid"].unique())
    tirage.shuffle(fiches)
    groupe = df["cid"].map(pd.Series(np.arange(len(fiches)) % N_TOURS, index=fiches))
    ys, scores = [], []
    for tour in range(N_TOURS):
        est_teste = (groupe == tour).to_numpy()
        ajustement = df[~est_teste]
        vus = df["author_key"].isin(set(ajustement["author_key"])).to_numpy()
        test = df[est_teste & ~vus]
        res, colonnes = ajuster(ajustement, avec_region)
        X_test = matrice(test, avec_region).reindex(columns=colonnes, fill_value=0.0)
        scores.append(np.asarray(res.predict(X_test)))
        ys.append(test["supprime"].to_numpy())
    return aire_sous_courbe(np.concatenate(ys), np.concatenate(scores))


# ---------------------------------------------------------------------------
# Un passage
# ---------------------------------------------------------------------------

def effectifs(df: pd.DataFrame, passage: str) -> pd.DataFrame:
    familles = [("etoiles", "Note"), ("texte", "Longueur du texte"),
                ("has_photo", "Photo jointe"), ("tranche_photos_auteur", "Photos de l'auteur"),
                ("guide", "Profil Local Guide"), ("tranche_afflux", "Afflux sur la fiche"),
                ("secteur", "Secteur"), ("taille", "Taille de l'entreprise")]
    lignes = []
    for colonne, famille in familles:
        g = df.groupby(colonne, observed=True)["supprime"].agg(["sum", "size"])
        for valeur, row in g.iterrows():
            if colonne == "has_photo":
                libelle = "photo jointe" if valeur == 1 else "sans photo"
            else:
                libelle = LIBELLE.get(valeur, valeur).replace(" (référence)", "")
            lignes.append({"passage": passage, "famille": famille, "valeur": libelle,
                           "avis": int(row["size"]), "suppressions": int(row["sum"]),
                           "taux_pour_10000_avis": round(row["sum"] / row["size"] * 1e4, 1)})
    return pd.DataFrame(lignes)


def un_passage(df: pd.DataFrame, passage: str) -> tuple[pd.DataFrame, dict, pd.DataFrame]:
    df = df.assign(secteur=secteurs(df))
    avec_region = df["region"].nunique() > 1
    res, _ = ajuster(df, avec_region)

    ic = res.conf_int()
    tableau = pd.DataFrame({
        "coefficient": res.params, "std_err": res.bse,
        "risque_relatif": np.exp(res.params),
        "borne_basse": np.exp(ic[0]), "borne_haute": np.exp(ic[1]),
        "p_value": res.pvalues,
    })
    tableau["role"] = ["controle" if v in CONTROLES or v == "const" else "resultat"
                       for v in tableau.index]
    tableau.index.name = "variable"
    tableau.to_csv(TECHNIQUE / f"07E_coefficients_{passage}.csv")
    (TECHNIQUE / f"07E_summary_{passage}.txt").write_text(
        f"Passage : {passage}\nCorpus : {CORPUS.name}\n"
        f"{len(df)} avis, {int(df['supprime'].sum())} suppressions, "
        f"{df['cid'].nunique()} fiches\n"
        "log_age_premiere_observation, fenetre_observation_j et region_US sont des\n"
        "contrôles : ils ne s'affichent pas dans le rendu client.\n\n"
        + str(res.summary()), encoding="utf-8")

    print(f"\n[{passage}] {nombre(len(df))} avis, {nombre(df['supprime'].sum())} "
          f"suppressions, {nombre(df['cid'].nunique())} fiches")
    auc = cinq_tours(df, avec_region)
    print(f"  qualité de classement sur des fiches jamais vues : {auc:.3f}")

    # Rendu client : caractéristiques seulement, références comprises.
    resultats = tableau[tableau["role"] == "resultat"]
    lignes = []
    for variable in ORDRE:
        if variable in REFERENCES.values():
            if variable.startswith("secteur_") and variable not in set(df["secteur"]):
                continue
            rr, basse, haute = 1.0, np.nan, np.nan
        elif variable in resultats.index:
            rr, basse, haute = resultats.loc[variable, ["risque_relatif", "borne_basse",
                                                        "borne_haute"]]
        else:
            continue
        lignes.append({"passage": passage, "famille": FAMILLE[variable],
                       "variable": variable, "caracteristique": LIBELLE[variable],
                       "risque_relatif": rr, "borne_basse": basse, "borne_haute": haute})
    bilan = {"passage": passage, "avis": len(df),
             "suppressions": int(df["supprime"].sum()), "fiches": df["cid"].nunique(),
             "taux_pour_10000_avis": round(df["supprime"].mean() * 1e4, 1),
             "qualite_classement_auc": round(auc, 3)}
    return pd.DataFrame(lignes), bilan, effectifs(df, passage)


# ---------------------------------------------------------------------------
# Rendu
# ---------------------------------------------------------------------------

def fmt(x: float) -> str:
    return f"{x:.2f}".replace(".", ",")


def nombre(x) -> str:
    return f"{int(x):,}".replace(",", " ")


def tableau_client(long: pd.DataFrame) -> pd.DataFrame:
    long = long.copy()
    long["valeur"] = [
        "référence" if np.isnan(b) else f"×{fmt(r)} ({fmt(b)} – {fmt(h)})"
        for r, b, h in zip(long["risque_relatif"], long["borne_basse"], long["borne_haute"])]
    large = long.pivot_table(index=["famille", "variable", "caracteristique"],
                             columns="passage", values="valeur", aggfunc="first")
    large = large.reindex(columns=[f"{r}_{j}" for r, j in PASSAGES])
    large.columns = [f"{NOM_REGION[r]}, {NOM_JEU[j]}" for r, j in PASSAGES]
    large = large.reset_index()
    large["ordre"] = large["variable"].map(ORDRE.index)
    return (large.sort_values("ordre").drop(columns=["ordre", "variable"])
            .fillna("—"))


def figure_region(long: pd.DataFrame, bilans: dict, region: str) -> None:
    """Risques relatifs d'une région, avec et sans les 6 enseignes. Échelle
    multiplicative : ×2 et ×0,5 sont à la même distance de ×1."""
    sous = long[long["passage"].str.startswith(f"{region}_")
                & long["borne_basse"].notna()]
    variables = [v for v in ORDRE if v in set(sous["variable"])]
    if not variables:
        return

    # Une ligne par caractéristique, un titre de famille au-dessus de chaque groupe.
    rangs, etiquettes, titres = {}, [], []
    y = 0
    famille_courante = None
    for v in variables:
        if FAMILLE[v] != famille_courante:
            famille_courante = FAMILLE[v]
            titres.append((y, famille_courante))
            y += 1
        rangs[v] = y
        etiquettes.append((y, LIBELLE[v]))
        y += 1

    bas_x, haut_x = 0.05, 50.0
    fig, ax = plt.subplots(figsize=(8.5, 0.32 * y + 2.2), dpi=150)
    fig.patch.set_facecolor(FOND)
    ax.set_facecolor(FOND)
    ax.axvline(1, color=ENCRE_2, lw=1, zorder=1)
    for jeu, decalage in [("sans_enseignes", -0.17), ("toutes_fiches", 0.17)]:
        d = sous[sous["passage"] == f"{region}_{jeu}"]
        for _, r in d.iterrows():
            yy = rangs[r["variable"]] + decalage
            b, h = max(r["borne_basse"], bas_x), min(r["borne_haute"], haut_x)
            ax.plot([b, h], [yy, yy], color=COULEUR_JEU[jeu], lw=2,
                    solid_capstyle="round", zorder=2)
            ax.plot(min(max(r["risque_relatif"], bas_x), haut_x), yy, "o",
                    color=COULEUR_JEU[jeu], markersize=5.5,
                    markeredgecolor=FOND, markeredgewidth=1, zorder=3)
        b = bilans[f"{region}_{jeu}"]
        ax.plot([], [], "o-", color=COULEUR_JEU[jeu], lw=2, markersize=5.5,
                label=f"{NOM_JEU[jeu]} : {nombre(b['avis'])} avis, "
                      f"{nombre(b['suppressions'])} suppressions")

    ax.set_xscale("log")
    ax.set_xlim(bas_x, haut_x)
    graduations = [0.1, 0.2, 0.5, 1, 2, 5, 10, 20]
    ax.set_xticks(graduations)
    ax.set_xticklabels([f"×{g:g}".replace(".", ",") for g in graduations],
                       color=ENCRE_2, fontsize=8)
    ax.xaxis.set_minor_locator(matplotlib.ticker.NullLocator())
    ax.set_yticks([p for p, _ in etiquettes])
    ax.set_yticklabels([t for _, t in etiquettes], color=ENCRE, fontsize=8)
    for p, t in titres:
        if t in REFERENCE_DE_FAMILLE:
            t = f"{t} (référence : {REFERENCE_DE_FAMILLE[t]})"
        ax.annotate(t, xy=(0, p), xycoords=("axes fraction", "data"), xytext=(-6, 0),
                    textcoords="offset points", ha="right", va="center",
                    fontsize=8.5, fontweight="bold", color=ENCRE)
    ax.set_ylim(y - 0.5, -0.8)
    ax.tick_params(axis="y", length=0, pad=6)
    ax.tick_params(axis="x", length=0)
    ax.grid(axis="x", color=GRILLE, lw=0.8, zorder=0)
    for cote in ["top", "right", "left"]:
        ax.spines[cote].set_visible(False)
    ax.spines["bottom"].set_color(GRILLE)
    ax.set_xlabel("risque de suppression comparé à la référence de la famille\n"
                  "(×1 = même risque ; le trait donne la fourchette)",
                  color=ENCRE_2, fontsize=8)
    # En-tête en pouces depuis le haut, pour qu'il ne dépende pas du nombre de
    # lignes de la figure.
    hauteur = fig.get_figheight()
    fig.text(0.01, 1 - 0.22 / hauteur,
             f"{NOM_REGION[region]} : ce qui augmente ou réduit le risque qu'un avis "
             f"soit supprimé", va="top", fontsize=11, color=ENCRE)
    fig.text(0.01, 1 - 0.50 / hauteur,
             f"Corpus {CORPUS.stem}. À âge de l'avis, durée d'observation"
             + (" et région" if region == "tous" else "") + " égaux.",
             va="top", fontsize=8, color=ENCRE_2)
    poignees, textes = ax.get_legend_handles_labels()
    fig.legend(poignees, textes, loc="upper left", bbox_to_anchor=(0.005, 1 - 0.68 / hauteur),
               ncol=1, frameon=False, fontsize=8, labelcolor=ENCRE, handlelength=1.8)
    fig.tight_layout(rect=(0, 0, 1, 1 - 1.25 / hauteur))
    fig.savefig(SORTIES / f"07E_figure_{region}.png", facecolor=FOND)
    plt.close(fig)


# ---------------------------------------------------------------------------
# Restitution client : livrables/restitution/regression_<corpus>.docx
#
# Même mise en forme que les notes déjà livrées (etudes-ponctuelles/
# 2026-09-30-secteurs-panel-03B/note.py) : Calibri, en-têtes bleu marine.
# ---------------------------------------------------------------------------

NAVY_HEX, GRIS_HEX, FOND_FAMILLE_HEX = "322954", "6B6A66", "EDEDF2"
POLICE = "Calibri"
LARGEUR_UTILE_CM = 16.6
HAUTEUR_FIGURE_MAX_CM = 22.5


def taux(x: float) -> str:
    """Taux pour 10 000 avis, une décimale au plus : « 155,1 », « 30 »."""
    return f"{x:.1f}".replace(".", ",").removesuffix(",0")


def ecart_age(garde: pd.DataFrame) -> str:
    """Phrase sur l'effet de l'âge, mesurée sur le corpus lu : taux de suppression
    des avis vus dans leur première semaine contre ceux vus à 91 jours ou plus."""
    age = garde["age_a_la_premiere_observation_j"]
    jeunes, vieux = garde.loc[age <= 7, "supprime"], garde.loc[age >= 91, "supprime"]
    if len(jeunes) < 1000 or len(vieux) < 1000 or vieux.mean() == 0:
        return ""
    return (f" Sur ce corpus, un avis vu dans sa première semaine disparaît "
            f"{jeunes.mean() / vieux.mean():.0f} fois plus souvent qu'un avis vu à "
            f"trois mois ou plus ({taux(jeunes.mean() * 1e4)} contre "
            f"{taux(vieux.mean() * 1e4)} pour 10 000 avis).")


# Tournures de la conclusion.
COMPARE_A = {
    "Note": "5 étoiles", "Longueur du texte": "un avis sans texte",
    "Photo jointe": "un avis sans photo", "Profil Local Guide": "un compte de niveau 1 à 4",
    "Secteur": "l'automobile", "Taille de l'entreprise": "un seul établissement",
}
PAR_MULTIPLICATION = {
    "log_photos_auteur": "à chaque multiplication par 2,7 du nombre de photos publiées "
                         "par l'auteur",
    "log_ratio_pic_journalier_fiche": "à chaque multiplication par 2,7 du nombre d'avis "
                                      "reçus par la fiche le jour du dépôt",
}


def nom_court(v: str) -> str:
    return FAMILLE[v] if v in PAR_MULTIPLICATION else f"{FAMILLE[v]}, {LIBELLE[v]}"


def net(ligne) -> bool:
    """L'écart est net quand la fourchette ne contient pas ×1."""
    return ligne["borne_basse"] > 1 or ligne["borne_haute"] < 1


def valeur(ligne) -> str:
    return (f"×{fmt(ligne['risque_relatif'])} "
            f"({fmt(ligne['borne_basse'])} – {fmt(ligne['borne_haute'])})")


def facteurs(long: pd.DataFrame) -> dict:
    """Classe chaque caractéristique pour la conclusion, selon une règle fixe.

    Retenue : écart net sur États-Unis et Europe sans les 6 enseignes, et même
    sens aux États-Unis seuls, en Europe seule et quand on garde les enseignes.
    À ne pas citer seule : écart net, mais un de ces trois sens s'inverse.
    Opposée selon la région : aucun écart net sur l'ensemble, mais un écart net
    dans chaque région, en sens contraires.
    """
    est = long[long["borne_basse"].notna()].set_index(["passage", "variable"])

    def ligne(passage, v):
        return est.loc[(passage, v)] if (passage, v) in est.index else None

    classes = {"hausse": [], "baisse": [], "fragiles": [], "opposes": []}
    for v in ORDRE:
        t = ligne("tous_sans_enseignes", v)
        us, eu = ligne("US_sans_enseignes", v), ligne("Europe_sans_enseignes", v)
        toutes = ligne("tous_toutes_fiches", v)
        if t is not None and net(t):
            sens = t["risque_relatif"] > 1
            raisons = [texte for texte, x in [
                ("sens inverse aux États-Unis", us), ("sens inverse en Europe", eu),
                ("sens inverse quand on garde les 6 enseignes", toutes)]
                if x is not None and (x["risque_relatif"] > 1) != sens]
            if raisons:
                classes["fragiles"].append((v, t, raisons))
            else:
                classes["hausse" if sens else "baisse"].append((v, t))
        elif (us is not None and eu is not None and net(us) and net(eu)
              and (us["risque_relatif"] > 1) != (eu["risque_relatif"] > 1)):
            classes["opposes"].append((v, us, eu))
    citees = {FAMILLE[x[0]] for c in classes.values() for x in c}
    classes["sans_effet"] = [f for f in dict.fromkeys(FAMILLE[v] for v in ORDRE)
                             if f not in citees and f in set(long["famille"])]
    return classes


def ecrire_word(df: pd.DataFrame, long: pd.DataFrame, bilans: dict,
                effs: pd.DataFrame) -> Path:
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    from PIL import Image

    navy, gris, encre = (RGBColor.from_string(NAVY_HEX), RGBColor.from_string(GRIS_HEX),
                         RGBColor(0x1F, 0x1F, 0x1F))

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(1.8)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name, normal.font.size, normal.font.color.rgb = POLICE, Pt(10.5), encre
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), POLICE)
    normal.paragraph_format.space_after = Pt(6)
    for nom, taille, avant, apres in (("Title", 22, 0, 4), ("Heading 1", 15, 16, 6),
                                      ("Heading 2", 11.5, 10, 4)):
        s = styles[nom]
        s.font.name, s.font.size, s.font.bold, s.font.italic = POLICE, Pt(taille), True, False
        s.font.color.rgb = navy
        polices = s.element.get_or_add_rPr().find(qn("w:rFonts"))
        if polices is None:
            polices = OxmlElement("w:rFonts")
            s.element.get_or_add_rPr().append(polices)
        for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            polices.set(qn(a), POLICE)
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            polices.attrib.pop(qn(a), None)
        s.paragraph_format.space_before, s.paragraph_format.space_after = Pt(avant), Pt(apres)
        s.paragraph_format.keep_with_next = True
        for bordure in s.element.get_or_add_pPr().findall(qn("w:pBdr")):
            s.element.get_or_add_pPr().remove(bordure)
    note = styles.add_style("NoteTableau", 1)
    note.base_style = normal
    note.font.size, note.font.color.rgb = Pt(9), gris
    note.paragraph_format.space_after = Pt(10)

    def fond(cellule, couleur):
        ombre = OxmlElement("w:shd")
        ombre.set(qn("w:val"), "clear")
        ombre.set(qn("w:color"), "auto")
        ombre.set(qn("w:fill"), couleur)
        cellule._tc.get_or_add_tcPr().append(ombre)

    def puce(morceaux):
        """Une puce ; `morceaux` est un texte ou une liste de (texte, gras)."""
        p = doc.add_paragraph(style="List Bullet")
        for texte, gras in ([(morceaux, False)] if isinstance(morceaux, str) else morceaux):
            p.add_run(texte).bold = gras
        p.paragraph_format.space_after = Pt(3)
        return p

    def tableau(entetes, lignes, largeurs):
        """`lignes` : listes de cellules, ou ("famille", titre) pour un intertitre.
        Une cellule est un texte, ou (texte, gras)."""
        t = doc.add_table(rows=1 + len(lignes), cols=len(entetes))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        bords = OxmlElement("w:tblBorders")
        for cote in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{cote}")
            if cote in ("left", "right", "insideV"):
                e.set(qn("w:val"), "nil")
            else:
                e.set(qn("w:val"), "single")
                e.set(qn("w:sz"), "4")
                e.set(qn("w:color"), "D9D8DE")
            bords.append(e)
        t._tbl.tblPr.append(bords)
        for i, rangee in enumerate([entetes] + lignes):
            cellules = t.rows[i].cells
            if isinstance(rangee, tuple) and rangee[0] == "famille":
                fusion = cellules[0].merge(cellules[-1])
                p = fusion.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                r = p.add_run(rangee[1])
                r.bold, r.font.size, r.font.color.rgb = True, Pt(9), navy
                fond(fusion, FOND_FAMILLE_HEX)
                continue
            for j, contenu in enumerate(rangee):
                c = cellules[j]
                c.width = Cm(largeurs[j])
                c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                p = c.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.RIGHT
                texte, gras = contenu if isinstance(contenu, tuple) else (contenu, False)
                r = p.add_run(str(texte))
                r.font.size = Pt(9)
                r.bold = gras
                if i == 0:
                    r.bold, r.font.color.rgb = True, RGBColor(0xFF, 0xFF, 0xFF)
                    fond(c, NAVY_HEX)
        for j, colonne in enumerate(t._tbl.tblGrid.findall(qn("w:gridCol"))):
            colonne.set(qn("w:w"), str(int(largeurs[j] * 567)))
        # Les en-têtes se répètent en haut de chaque page.
        entete = t.rows[0]._tr.get_or_add_trPr()
        repete = OxmlElement("w:tblHeader")
        repete.set(qn("w:val"), "true")
        entete.append(repete)
        return t

    def image(chemin):
        with Image.open(chemin) as im:
            largeur_px, hauteur_px = im.size
        largeur = min(LARGEUR_UTILE_CM, HAUTEUR_FIGURE_MAX_CM * largeur_px / hauteur_px)
        doc.add_picture(str(chemin), width=Cm(largeur))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    def nouvelle_page():
        doc.add_section(WD_SECTION.NEW_PAGE)

    # --- chiffres du texte, tous lus sur le corpus et les résultats ---------
    garde = df[df["exposable"]]
    tous, sans = bilans["tous_toutes_fiches"], bilans["tous_sans_enseignes"]
    suppr_enseignes = int(garde.loc[garde["enseigne"], "supprime"].sum())
    non_comptes = int((df["disparu_au_dernier_releve"] & (df["supprime"] == 0)).sum())
    aucs = [b["qualite_classement_auc"] for b in bilans.values()]
    estimes = long[long["borne_basse"].notna()]
    exemple = estimes[(estimes["passage"] == "tous_sans_enseignes")
                      & (estimes["variable"] == "etoiles_1")]

    # --- titre et introduction ------------------------------------------------
    doc.add_paragraph("Ce qui fait qu'un avis Google est supprimé", style="Title")
    doc.add_paragraph(
        f"Régression logistique sur les avis publiés du "
        f"{garde['created_at_day'].min():%d/%m/%Y} au {garde['created_at_day'].max():%d/%m/%Y}. "
        f"Document produit le {date.today():%d/%m/%Y}.", style="NoteTableau")

    doc.add_paragraph("Ce que nous avons fait", style="Heading 1")
    puce([(f"{nombre(tous['fiches'])} fiches Google", True),
          (", aux États-Unis et en Europe, ont été relevées une fois par jour du 11 au "
           "24 août 2026.", False)])
    puce([(f"{nombre(tous['avis'])} avis", True),
          (f" publiés entre le {garde['created_at_day'].min():%d/%m/%Y} et le "
           f"{garde['created_at_day'].max():%d/%m/%Y} ont été suivis pendant ces deux "
           "semaines. Un avis est compté supprimé quand il n'est plus visible deux jours "
           f"de suite ou plus : {nombre(tous['suppressions'])} avis sont dans ce cas.", False)])
    puce("Pour chaque avis, huit caractéristiques ont été relevées : la note, la longueur "
         "du texte, la présence d'une photo, le nombre de photos publiées par l'auteur, son "
         "niveau Local Guide, l'afflux d'avis sur la fiche le jour du dépôt, le secteur et "
         "la taille de l'entreprise.")
    puce("Un calcul unique, une régression logistique, mesure l'effet de chacune sur le "
         "risque de suppression, les sept autres étant égales par ailleurs. Il tient compte "
         "aussi de l'âge de l'avis, de sa durée d'observation et de la région de la fiche, "
         "sans afficher leur effet.")
    puce([("Six enseignes attaquées pendant la période", True),
          (" (quatre chaînes américaines de traitement antiparasitaire, deux salles de sport "
           f"espagnoles) portent {nombre(suppr_enseignes)} des "
           f"{nombre(tous['suppressions'])} suppressions. Chaque résultat est donné sans elles "
           "et avec elles. La version sans les 6 enseignes est celle à citer "
           f"({nombre(sans['avis'])} avis, {nombre(sans['suppressions'])} suppressions).",
           False)])

    doc.add_paragraph("Comment lire les graphiques et les tableaux", style="Heading 1")
    puce("Chaque ligne compare un groupe d'avis à la référence de sa famille, indiquée "
         "dans le titre de la famille : 1 étoile comparé à 5 étoiles, un compte sans "
         "niveau Local Guide comparé à un compte de niveau 1 à 4, etc.")
    if not exemple.empty:
        r = exemple.iloc[0]["risque_relatif"]
        puce([("Le chiffre dit combien de fois le groupe est plus supprimé que sa "
               "référence. ", False),
              (f"×{fmt(r)} pour 1 étoile", True),
              (f" : quand 10 avis 5 étoiles sont supprimés, {10 * r:.0f} avis 1 étoile le "
               "sont, à nombre d'avis égal. ×1 veut dire autant, ×0,5 deux fois moins.",
               False)])
    else:
        puce("Le chiffre dit combien de fois le groupe est plus supprimé que sa référence : "
             "×2 deux fois plus, ×1 autant, ×0,5 deux fois moins.")
    puce("Sur les graphiques, le point donne ce chiffre et le trait sa fourchette : la "
         "valeur réelle se situe très probablement entre ses deux bouts. Un trait qui coupe "
         "la ligne verticale ×1 signale un écart non établi.")
    puce("L'échelle horizontale place ×2 et ×0,5 à la même distance de ×1 : un effet qui "
         "double le risque et un effet qui le divise par deux se voient autant.")
    puce([("En bleu, les fiches sans les 6 enseignes ; en orange, toutes les fiches. ", False),
          ("Dans les tableaux, les écarts établis sont en gras.", True)])
    puce("« Par multiplication par 2,7 » : effet d'un passage de 10 à 27 photos publiées "
         "par l'auteur, ou de 100 à 270. Même lecture pour l'afflux d'avis sur la fiche.")

    # --- une section par région ------------------------------------------------
    for region in ["tous", "US", "Europe"]:
        nouvelle_page()
        doc.add_paragraph(NOM_REGION[region], style="Heading 1")
        figure = SORTIES / f"07E_figure_{region}.png"
        if figure.exists():
            image(figure)
        nouvelle_page()
        doc.add_paragraph(f"{NOM_REGION[region]} : les chiffres", style="Heading 2")
        par_jeu = {jeu: estimes[estimes["passage"] == f"{region}_{jeu}"].set_index("variable")
                   for jeu in ["sans_enseignes", "toutes_fiches"]}
        lignes, famille_courante = [], None
        for v in ORDRE:
            if not any(v in t.index for t in par_jeu.values()):
                continue
            if FAMILLE[v] != famille_courante:
                famille_courante = FAMILLE[v]
                reference = REFERENCE_DE_FAMILLE.get(famille_courante)
                lignes.append(("famille", famille_courante
                               + (f" (référence : {reference})" if reference else "")))
            cellules = [LIBELLE[v]]
            for t in par_jeu.values():
                cellules.append((valeur(t.loc[v]), net(t.loc[v])) if v in t.index else "—")
            lignes.append(cellules)
        b_sans, b_toutes = bilans[f"{region}_sans_enseignes"], bilans[f"{region}_toutes_fiches"]
        tableau(["Caractéristique", "Sans les 6 enseignes", "Toutes les fiches"], lignes,
                [7.6, 4.5, 4.5])
        doc.add_paragraph(
            f"Sans les 6 enseignes : {nombre(b_sans['avis'])} avis, "
            f"{nombre(b_sans['suppressions'])} suppressions. Toutes les fiches : "
            f"{nombre(b_toutes['avis'])} avis, {nombre(b_toutes['suppressions'])} "
            "suppressions. En gras, les écarts établis : la fourchette ne contient pas ×1.",
            style="NoteTableau")

    # --- conclusion -----------------------------------------------------------
    nouvelle_page()
    classes = facteurs(long)
    doc.add_paragraph("Les facteurs retenus", style="Heading 1")
    doc.add_paragraph(
        "Règle : un facteur est retenu quand son écart est établi sur l'ensemble des fiches "
        "sans les 6 enseignes, et qu'il va dans le même sens aux États-Unis, en Europe et "
        "quand on garde les 6 enseignes.", style="NoteTableau")

    def liste_facteurs(titre, elements):
        if not elements:
            return
        doc.add_paragraph(titre, style="Heading 2")
        par_famille = {}
        for v, t in elements:
            par_famille.setdefault(FAMILLE[v], []).append((v, t))
        for famille, items in par_famille.items():
            morceaux = [(f"{famille} : ", True)]
            for k, (v, t) in enumerate(items):
                chiffre = [(f"×{fmt(t['risque_relatif'])}", True),
                           (f" ({fmt(t['borne_basse'])} – {fmt(t['borne_haute'])})", False)]
                if v in PAR_MULTIPLICATION:
                    morceaux += chiffre + [(f" {PAR_MULTIPLICATION[v]}", False)]
                else:
                    morceaux += ([(("" if k == 0 else " ; ") + f"{LIBELLE[v]} ", False)]
                                 + chiffre)
            if famille in COMPARE_A:
                morceaux.append((f", par rapport à {COMPARE_A[famille]}", False))
            morceaux.append((".", False))
            puce(morceaux)

    liste_facteurs("Ce qui augmente le risque de suppression", classes["hausse"])
    liste_facteurs("Ce qui le réduit", classes["baisse"])
    if classes["fragiles"] or classes["opposes"]:
        doc.add_paragraph("Écarts à ne pas citer seuls", style="Heading 2")
        for v, t, raisons in classes["fragiles"]:
            puce([(f"{nom_court(v)} : ", True),
                  (f"{valeur(t)} sans les 6 enseignes, mais {', '.join(raisons)}.", False)])
        for v, us, eu in classes["opposes"]:
            puce([(f"{nom_court(v)} : ", True),
                  (f"{valeur(us)} aux États-Unis et {valeur(eu)} en Europe, sans écart "
                   "établi sur l'ensemble.", False)])
    if classes["sans_effet"]:
        doc.add_paragraph("Sans effet établi", style="Heading 2")
        puce(", ".join(classes["sans_effet"]).capitalize() + ".")

    doc.add_paragraph("Limites", style="Heading 2")
    puce("Un avis supprimé avant le 11 août n'a jamais été vu. Pour les avis anciens, "
         "seules les suppressions survenues pendant les deux semaines de relevé sont "
         "mesurées.")
    puce(("L'âge de l'avis pèse lourd." + ecart_age(garde)).strip() + " Le calcul le "
         "neutralise : les effets affichés comparent des avis du même âge.")
    puce("L'écart entre fiches américaines et européennes n'est pas affiché. Il dépend "
         "des fiches choisies pour l'étude et ne dit rien des deux marchés.")
    puce(f"{nombre(tous['avis_ecartes_fenetre_nulle'])} avis vus pour la première fois lors "
         "des deux derniers relevés sont écartés, faute de pouvoir constater deux jours "
         f"d'absence. Pour la même raison, {nombre(non_comptes)} avis disparus au dernier "
         "relevé ne sont pas comptés comme supprimés.")
    puce(f"Sur des fiches que le calcul n'a jamais vues, un avis supprimé reçoit un score "
         f"plus élevé qu'un avis resté en ligne dans {min(aucs) * 100:.0f} à "
         f"{max(aucs) * 100:.0f} % des cas selon le passage. L'âge de l'avis compte dans "
         "ce score.")

    # --- annexe : effectifs ---------------------------------------------------
    nouvelle_page()
    doc.add_paragraph("Annexe : les effectifs", style="Heading 1")
    doc.add_paragraph("Avis et suppressions observés, avant tout calcul d'effet.",
                      style="NoteTableau")
    tableau(["Passage", "Avis", "Suppressions", "Fiches", "Pour 10 000 avis"],
            [[f"{NOM_REGION[r]}, {NOM_JEU[j]}", nombre(bilans[f'{r}_{j}']['avis']),
              nombre(bilans[f"{r}_{j}"]["suppressions"]), nombre(bilans[f"{r}_{j}"]["fiches"]),
              taux(bilans[f"{r}_{j}"]["taux_pour_10000_avis"])] for r, j in PASSAGES],
            [7.1, 2.5, 2.5, 2.0, 2.5])

    doc.add_paragraph(f"Par caractéristique, {NOM_REGION['tous']}", style="Heading 2")
    fusion = (effs[effs["passage"] == "tous_toutes_fiches"]
              .merge(effs[effs["passage"] == "tous_sans_enseignes"], on=["famille", "valeur"],
                     how="left", suffixes=("_toutes", "_sans"), sort=False))
    ordre_valeurs = ([LIBELLE[v].replace(" (référence)", "") for v in ORDRE]
                     + ["sans photo", "photo jointe"] + TRANCHES_PHOTOS + TRANCHES_AFFLUX)
    ordre_familles = list(dict.fromkeys(fusion["famille"]))
    fusion = fusion.assign(
        _f=fusion["famille"].map(ordre_familles.index),
        _v=fusion["valeur"].map(lambda x: ordre_valeurs.index(x)
                                if x in ordre_valeurs else len(ordre_valeurs)),
    ).sort_values(["_f", "_v"], kind="stable")
    lignes, famille_courante = [], None
    for _, r in fusion.iterrows():
        if r["famille"] != famille_courante:
            famille_courante = r["famille"]
            lignes.append(("famille", famille_courante))
        sans_cols = (["—", "—"] if pd.isna(r["avis_sans"]) else
                     [nombre(r["avis_sans"]), taux(r["taux_pour_10000_avis_sans"])])
        lignes.append([r["valeur"]] + sans_cols
                      + [nombre(r["avis_toutes"]), taux(r["taux_pour_10000_avis_toutes"])])
    tableau(["Caractéristique", "Avis, sans les 6 enseignes", "Pour 10 000 avis",
             "Avis, toutes les fiches", "Pour 10 000 avis"], lignes,
            [5.8, 3.0, 2.4, 3.0, 2.4])
    doc.add_paragraph("« Pour 10 000 avis » : nombre d'avis supprimés sur 10 000 avis du "
                      "groupe, sans correction de l'âge ni des autres caractéristiques.",
                      style="NoteTableau")

    RESTITUTION.mkdir(parents=True, exist_ok=True)
    sortie = RESTITUTION / f"regression_{CORPUS.stem}.docx"
    doc.save(sortie)
    return sortie


# ---------------------------------------------------------------------------
# Programme
# ---------------------------------------------------------------------------

def main() -> int:
    TECHNIQUE.mkdir(parents=True, exist_ok=True)
    print(f"[corpus] {CORPUS}\n[moteur] {MOTEUR}")
    table = features_a_jour(charger_03c())
    df = preparer(pd.read_parquet(table))
    ecartes = int((~df["exposable"]).sum())
    print(f"[préparation] {nombre(len(df))} avis, dont {nombre(ecartes)} écartés "
          f"(fenêtre d'observation nulle : ils ne peuvent pas être supprimés)")
    (TECHNIQUE / "07E_source.txt").write_text(
        f"Corpus : {CORPUS}\nCaractéristiques : {table}\n"
        f"Lancé le : {datetime.now():%Y-%m-%d %H:%M}\nMoteur : {MOTEUR}\n"
        f"Avis du corpus : {len(df)}\n"
        f"Avis écartés, fenêtre d'observation nulle : {ecartes}\n"
        f"Suppression : absent deux jours de suite ou plus (03C_features_corpus.py)\n",
        encoding="utf-8")

    longs, bilans, effs = [], {}, []
    for region, jeu in PASSAGES:
        garde = selection(df, region, jeu)
        passage = f"{region}_{jeu}"
        long, bilan, eff = un_passage(df[garde & df["exposable"]], passage)
        bilan["avis_ecartes_fenetre_nulle"] = int((garde & ~df["exposable"]).sum())
        longs.append(long)
        bilans[passage] = bilan
        effs.append(eff)

    long = pd.concat(longs, ignore_index=True)
    long.to_csv(SORTIES / "07E_risques_relatifs.csv", index=False)
    tableau_client(long).to_csv(SORTIES / "07E_tableau_client.csv", index=False,
                                encoding="utf-8-sig")
    pd.concat(effs, ignore_index=True).to_csv(SORTIES / "07E_effectifs.csv", index=False,
                                              encoding="utf-8-sig")
    pd.DataFrame(bilans.values()).to_csv(TECHNIQUE / "07E_passages.csv", index=False)
    for region in ["US", "Europe", "tous"]:
        figure_region(long, bilans, region)
    restitution = ecrire_word(df, long, bilans, pd.concat(effs, ignore_index=True))

    print(f"\nécrit dans {SORTIES}\nrestitution : {restitution}")
    print(pd.DataFrame(bilans.values()).to_string(index=False))
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception as erreur:
        print(f"\nARRÊT : {erreur}", file=sys.stderr)
        raise
