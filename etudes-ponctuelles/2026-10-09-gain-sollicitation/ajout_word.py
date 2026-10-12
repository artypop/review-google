#!/usr/bin/env python3
"""Ajoute le gain d'avis (solde net du 7 au 17 août) au Word de l'analyse exploratoire.

    PYTHONPATH=".venv/Lib/site-packages;outils" <python uv> etudes-ponctuelles/2026-10-09-gain-sollicitation/ajout_word.py

Demande de Matthieu du 2026-10-12. Le Word existant
(livrables/restitution_finale_v2/analyse_exploratoire/fiches_touchees_nouveaux_avis.docx)
est modifié sur place, sans être régénéré, pour garder d'éventuelles retouches :
  - une partie « 5. Le gain d'avis » est insérée avant la conclusion, qui devient la 6 ;
  - les phrases qui disaient le gain « à chiffrer » sont mises à jour (sautées, avec un
    message, si elles ont été retouchées) ;
  - la ligne « Pour refaire les calculs » cite aussi les scripts du gain.
Les chiffres sont lus dans les sorties de solde_par_jour.py. Refuse de tourner deux fois.
"""
import importlib.util
import sys
from pathlib import Path

import pandas as pd

RACINE = Path(__file__).resolve().parents[2]
_spec = importlib.util.spec_from_file_location(
    "mise_en_forme", RACINE / "etudes-ponctuelles" / "2026-10-07-prenom-dans-les-fiches" / "document.py")
mf = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(mf)
from docx import Document  # noqa: E402

WORD = RACINE / "livrables" / "restitution_finale_v2" / "analyse_exploratoire" / "fiches_touchees_nouveaux_avis.docx"
GAIN = RACINE / "livrables" / "restitution_finale_v2" / "gain_sollicitation"
TITRE_PARTIE = "5. Le gain d'avis : la sollicitation rapporte-t-elle ?"


def nombre(x: float, nd: int = 0) -> str:
    return f"{x:,.{nd}f}".replace(",", " ").replace(".", ",")


def signe(x: float, nd: int = 1) -> str:
    return f"{x:+.{nd}f}".replace(".", ",").replace("-", "−")


def remplacer(doc, ancien_debut: str, nouveau: str) -> None:
    """Remplace le texte du paragraphe qui commence par `ancien_debut`, mise en forme du 1er morceau."""
    for p in doc.paragraphs:
        if p.text.startswith(ancien_debut):
            for r in p.runs[1:]:
                r._element.getparent().remove(r._element)
            p.runs[0].text = nouveau
            return
    print(f"  non trouvé (retouché ?), laissé tel quel : « {ancien_debut[:60]}… »")


def chiffres() -> dict:
    c = pd.read_csv(GAIN / "solde_net_commerces.csv", dtype={"cid": str})
    g = c.groupby("groupe")
    s = g.agg(commerces=("cid", "size"), enseignes=("enseigne", "sum"), stock=("stock_depart", "sum"),
              publies=("publies", "sum"), nouveaux=("nouveaux_disparus", "sum"),
              nouveaux_tard=("nouveaux_disparus_apres_17", "sum"), anciens=("anciens_disparus", "sum"),
              solde17=("solde_cumule", "sum"), solde_complet=("solde_suivi_complet", "sum"),
              negatifs=("solde_cumule", lambda x: int((x < 0).sum())))  # au 17 août, comme les graphiques
    s["solde17_pct"] = 100 * s["solde17"] / s["stock"]
    s["solde_complet_pct"] = 100 * s["solde_complet"] / s["stock"]
    s["pertes_pour_100"] = 100 * (s["nouveaux"] + s["nouveaux_tard"] + s["anciens"]) / s["publies"]
    s["stock_moyen"] = s["stock"] / s["commerces"]
    s["publies_pct_stock"] = 100 * s["publies"] / s["stock"]
    return {"s": s, "total": len(c), "seuil_vitesse": c["avis_par_mois"].quantile(0.75)}


def main() -> int:
    doc = Document(WORD)
    if any(p.text.startswith(TITRE_PARTIE) for p in doc.paragraphs):
        raise SystemExit("La partie sur le gain d'avis est déjà dans le Word.")
    conclusion = next((p for p in doc.paragraphs if p.text.startswith("5. Ce qu'on peut répondre à Axel")), None)
    refaire = next((p for p in doc.paragraphs if p.text.startswith("Pour refaire les calculs")), None)
    if conclusion is None:
        raise SystemExit("Titre « 5. Ce qu'on peut répondre à Axel » introuvable : le Word a été retouché.")

    k = chiffres()
    s, pv, au = k["s"], k["s"].loc["prenoms_vitesse"], k["s"].loc["autres"]
    corps = doc.element.body
    avant = len(corps) - 1  # le dernier enfant est sectPr (mise en page) : les ajouts viennent avant lui

    # La nouvelle partie, écrite en fin de document puis déplacée avant la conclusion.
    mf.titre(doc, TITRE_PARTIE, 1)
    mf.paragraphe(doc, [("La question de Romain. ", True),
                        ("Les commerces qui sollicitent beaucoup d'avis finissent-ils avec plus d'avis "
                         "en ligne, une fois les suppressions retirées ? On compare deux groupes de "
                         "commerces jour après jour, du 7 au 17 août.", False)])
    mf.puce(doc, [("Les commerces : ", True),
                  (f"ceux qui ont au moins 15 avis 4-5 étoiles publiés du 7 au 17 août (corpus 03C), "
                   f"relevés jusqu'au 24 août : {k['total']} commerces.", False)])
    mf.puce(doc, [("Le groupe « prénoms et vitesse élevés » : ", True),
                  (f"au moins 50 % de leurs avis citent un prénom, et ils reçoivent au moins "
                   f"{nombre(k['seuil_vitesse'])} avis par mois (le quart des commerces les plus "
                   f"rapides). Ce sont les signes de sollicitation repérés par Axel : "
                   f"{int(pv['commerces'])} commerces. ", False),
                  ("Les autres : ", True), (f"{int(au['commerces'])} commerces.", False)])
    mf.puce(doc, [("Le solde net : ", True),
                  ("chaque jour, les avis publiés moins les avis disparus (nouveaux comme anciens), "
                   "cumulés depuis le 7 août, en pourcentage des avis en ligne au 7 août. Un solde "
                   "de +1 % veut dire que le commerce a 1 % d'avis de plus qu'au départ.", False)])
    mf.donnees(doc, "corpus 03C pour les avis publiés du 7 au 17 août 2026, base des avis pour les "
                    "anciens avis ; avis 4-5 étoiles. Prénoms : avis publiés avant le 7 août, annotés "
                    "par Axel ; vitesse : avis reçus par mois du 7 août 2025 au 6 août 2026.")

    mf.paragraphe(doc, [("Ce que montrent les graphiques. ", True),
                        ("Un graphique par groupe, à la même échelle. La ligne pleine est le solde "
                         "de l'ensemble du groupe ; la ligne en tirets, celui du commerce médian ; la "
                         "bande bleue contient la moitié des commerces. La zone grisée (7 au 10 août) "
                         "précède le début du suivi : on y voit les avis publiés, mais pas encore les "
                         "disparitions, qui sont notées en bloc le 12 août.", False)], avant=4)
    mf.image(doc, GAIN / "solde_net_par_jour_prenoms_vitesse.png",
             "Graphique 3 – Solde net d'avis du 7 au 17 août, commerces à prénoms et vitesse élevés.")
    mf.image(doc, GAIN / "solde_net_par_jour_autres.png",
             "Graphique 4 – Solde net d'avis du 7 au 17 août, autres commerces.")
    mf.tableau(doc, ["", "Prénoms et vitesse élevés", "Autres"], [
        ["Commerces", nombre(pv["commerces"]), nombre(au["commerces"])],
        ["Solde net au 17 août, en % des avis en ligne au départ",
         f"{signe(pv['solde17_pct'])} %", f"{signe(au['solde17_pct'])} %"],
        ["Même solde, en suivant 7 jours les avis publiés en fin de période",
         f"{signe(pv['solde_complet_pct'], 2)} %", f"{signe(au['solde_complet_pct'], 2)} %"],
        ["Avis perdus pour 100 avis publiés (nouveaux et anciens)",
         nombre(pv["pertes_pour_100"]), nombre(au["pertes_pour_100"])],
        ["Commerces en solde négatif au 17 août",
         f"{int(pv['negatifs'])} sur {int(pv['commerces'])}", f"{int(au['negatifs'])} sur {int(au['commerces'])}"],
        ["Avis en ligne au départ, en moyenne par commerce", nombre(pv["stock_moyen"]), nombre(au["stock_moyen"])],
    ], [8.6, 4.0, 3.4])
    mf.paragraphe(doc, "", apres=2)
    mf.paragraphe(doc, [("Ce qu'on en tire :", True)], apres=3)
    mf.puce(doc, "La sollicitation rapporte : dans les deux groupes, les commerces ont plus d'avis au "
                 "17 août qu'au 7 août, et presque aucun ne finit en solde négatif.")
    mf.puce(doc, f"Elle coûte plus cher aux commerces à prénoms et vitesse élevés : ils perdent "
                 f"{nombre(pv['pertes_pour_100'])} avis pour 100 publiés, contre "
                 f"{nombre(au['pertes_pour_100'])} pour les autres. Leur courbe s'aplatit à partir du "
                 "15 août, quand les disparitions s'accélèrent.")
    mf.puce(doc, f"L'écart de solde ({signe(pv['solde17_pct'])} % contre {signe(au['solde17_pct'])} %) "
                 f"tient aussi à la taille : ces commerces ont deux fois plus d'avis en ligne au départ, "
                 f"et leurs avis de la période ne représentent que {nombre(pv['publies_pct_stock'], 1)} % "
                 f"de leur stock, contre {nombre(au['publies_pct_stock'], 1)} % pour les autres. Le "
                 "chiffre « pour 100 avis publiés » est le plus parlant.")
    mf.paragraphe(doc, [("À savoir : ", True),
                        (f"{int(pv['enseignes'])} des {int(pv['commerces'])} commerces à prénoms et "
                         "vitesse élevés sont des fiches des quatre chaînes antiparasitaires ; elles "
                         "pèsent lourd dans ce groupe. Les deux groupes sont définis par des signes de "
                         "sollicitation, pas par une sollicitation observée.", False)],
                  taille=9.5, avant=4)

    nouveaux = [el for el in list(corps)[avant:] if el.tag.endswith("}p") or el.tag.endswith("}tbl")]
    for el in nouveaux:
        conclusion._element.addprevious(el)

    # La conclusion devient la partie 6, et tient compte du gain chiffré.
    for r in conclusion.runs:
        r.text = r.text.replace("5. Ce qu'on peut répondre à Axel", "6. Ce qu'on peut répondre à Axel")
    remplacer(doc, "Pour Axel : solliciter n'est pas un problème en soi",
              "Pour Axel : solliciter n'est pas un problème en soi, et cela rapporte des avis ; "
              "solliciter en rafale, auprès de comptes peu établis, coûte plus d'avis supprimés.")
    remplacer(doc, "Ce document ne chiffre pas ce que la sollicitation rapporte",
              "Le gain d'avis (partie 5) repose sur une période de 11 jours et sur des signes de "
              "sollicitation (prénoms, vitesse), pas sur une sollicitation observée.")
    question = next((p for p in doc.paragraphs if p.text.startswith("Question : Axel s'est-il emballé")
                     and "Avant de chiffrer" in p.text and len(p.runs) == 2), None)
    if question is None:
        print("  paragraphe « Question » retouché, laissé tel quel")
    else:
        question.runs[1].text = ("Axel s'est-il emballé sur la sollicitation d'avis ? On regarde quelles "
                                 "fiches perdent beaucoup de leurs nouveaux avis et ce qui les distingue, "
                                 "puis ce que la sollicitation rapporte en solde d'avis.")
    puce_axel = next((p for p in doc.paragraphs
                      if p.text.startswith("Solliciter des avis ne fait pas perdre d'avis en soi")), None)
    if puce_axel is not None:
        np_ = puce_axel.insert_paragraph_before(style="List Bullet")
        mf.police(np_.add_run(f"La sollicitation rapporte : du 7 au 17 août, les commerces à signes de "
                              f"sollicitation gagnent {signe(pv['solde17_pct'])} % d'avis, les autres "
                              f"{signe(au['solde17_pct'])} %. Mais les premiers perdent "
                              f"{nombre(pv['pertes_pour_100'])} avis pour 100 publiés, contre "
                              f"{nombre(au['pertes_pour_100'])}."), 10.5, False, mf.ENCRE)
        f = np_.paragraph_format
        f.space_after, f.line_spacing = mf.Pt(4), 1.15
    if refaire is not None:
        mf.police(refaire.add_run(" Gain d'avis : etudes-ponctuelles/2026-10-09-gain-sollicitation/"
                                  "solde_par_jour.py, puis ajout_word.py."), 8, False, mf.ENCRE_2)

    doc.save(WORD)
    print(f"écrit : {WORD}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
