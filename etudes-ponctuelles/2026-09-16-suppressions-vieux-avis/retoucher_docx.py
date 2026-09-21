#!/usr/bin/env python3
"""Retouche le document Word en place, sans le régénérer.

    python etudes-ponctuelles/2026-09-16-suppressions-vieux-avis/retoucher_docx.py

`rapport.py` reconstruit le document entier et écrase donc les corrections
faites à la main dans Word. Ce script-ci ne touche qu'aux éléments listés dans
REMPLACEMENTS et IMAGES, et laisse le reste intact.

Ce qu'il corrige, le 2026-09-16 : les figures 1 et 4 et les phrases qui les
commentent passent du dénominateur en journées d'exposition au dénominateur en
avis. « Sur 10 000 avis de cet âge, combien ont disparu » remplace
« suppressions pour un million de journées en ligne ».

Chaque remplacement est recherché par un fragment de texte. Un fragment absent
arrête le script plutôt que de modifier le mauvais paragraphe : le document
ayant pu être réécrit à la main, aucune position n'est fiable.
"""

from __future__ import annotations

import shutil
import sys
from pathlib import Path

from docx import Document

DOSSIER = Path(__file__).resolve().parent
SORTIES = DOSSIER / "sorties"
FIGURES = SORTIES / "figures"
CIBLE = SORTIES / "2026-09-16-suppressions-avis-google.docx"

# (fragment à retrouver, texte de remplacement complet du paragraphe)
REMPLACEMENTS = [
    (
        "fois plus souvent qu'un avis de plus de trois ans",
        "Sur 10 000 avis âgés de quatre à sept jours, 198 ont disparu pendant les quatorze "
        "jours de suivi. Sur 10 000 avis de plus de trois ans, 2.",
    ),
    (
        "Comparer des nombres bruts de suppressions ne dit rien",
        "Le nombre brut de suppressions ne suffit pas à comparer les âges : le parc contient "
        "2,87 millions d'avis de plus de trois ans contre 40 000 avis âgés de quatre à sept "
        "jours. Chaque tranche est donc rapportée au nombre d'avis qui ont eu cet âge pendant "
        "le suivi.",
    ),
    (
        "Figure 1 —",
        "Figure 1 — Sur 10 000 avis ayant cet âge pendant le suivi, combien ont disparu.",
    ),
    (
        "Le risque culmine entre quatre et sept jours",
        "Le risque culmine entre quatre et sept jours, à 198 disparitions pour 10 000 avis, "
        "puis s'effondre : 27 entre un et trois mois, 2 au-delà de trois ans. Une réserve "
        "tient à la durée d'observation : un avis de plus de trois ans reste dans sa tranche "
        "pendant les 12,9 jours du suivi, un avis de quatre à sept jours n'y reste que "
        "3,2 jours avant d'en sortir. À risque identique, la tranche large ramasse donc "
        "davantage de suppressions.",
    ),
    (
        "Figure 4 —",
        "Figure 4 — Sur 10 000 avis de cette note et de cet âge, combien ont disparu pendant "
        "le suivi. Les deux échelles diffèrent.",
    ),
]

# (fragment de la légende qui suit l'image, fichier PNG à mettre à la place)
IMAGES = [
    ("Figure 1 —", "figure1-age.png"),
    ("Figure 4 —", "figure4-note.png"),
]


def paragraphe_contenant(doc, fragment):
    trouves = [p for p in doc.paragraphs if fragment in p.text]
    if len(trouves) != 1:
        raise SystemExit(
            f"« {fragment} » : {len(trouves)} paragraphes trouvés, il en faut exactement un. "
            f"Le document a changé, rien n'a été écrit.")
    return trouves[0]


def remplacer_texte(paragraphe, texte):
    """Réécrit le paragraphe en gardant la mise en forme de son premier fragment."""
    for run in paragraphe.runs[1:]:
        run._element.getparent().remove(run._element)
    if paragraphe.runs:
        paragraphe.runs[0].text = texte
    else:
        paragraphe.add_run(texte)


def remplacer_image(doc, legende_fragment, fichier):
    """Remplace l'image du paragraphe qui précède immédiatement la légende."""
    legende = paragraphe_contenant(doc, legende_fragment)
    precedent = legende._p.getprevious()
    if precedent is None:
        raise SystemExit(f"Aucun paragraphe avant « {legende_fragment} ».")
    from docx.oxml.ns import qn
    blips = precedent.findall(".//" + qn("a:blip"))
    if len(blips) != 1:
        raise SystemExit(f"{len(blips)} image(s) avant « {legende_fragment} », il en faut une.")
    rid = blips[0].get(qn("r:embed"))
    part = doc.part.related_parts[rid]
    chemin = FIGURES / fichier
    part._blob = chemin.read_bytes()
    return fichier


def main() -> None:
    if not CIBLE.exists():
        raise SystemExit(f"{CIBLE} est introuvable.")
    verrou = CIBLE.parent / ("~$" + CIBLE.name[2:])
    if verrou.exists():
        raise SystemExit("Le document est ouvert dans Word. Le fermer, puis relancer.")

    secours = CIBLE.with_suffix(".avant-retouche.docx")
    shutil.copy2(CIBLE, secours)

    doc = Document(str(CIBLE))

    # Les images d'abord : elles se repèrent par la légende, qui va changer.
    for fragment, fichier in IMAGES:
        print("image remplacée :", remplacer_image(doc, fragment, fichier))

    for fragment, texte in REMPLACEMENTS:
        remplacer_texte(paragraphe_contenant(doc, fragment), texte)
        print("texte remplacé  :", fragment)

    doc.save(str(CIBLE))
    print(f"\n{CIBLE.name} mis à jour. Version précédente : {secours.name}")


if __name__ == "__main__":
    sys.exit(main())
